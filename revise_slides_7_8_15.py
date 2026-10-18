from __future__ import annotations

import shutil
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "课题1-推进会-任务2-0302(1)_修改版.pptx"
BACKUP = ROOT / "课题1-推进会-任务2-0302(1)_用户删图后备份.pptx"
OUTPUT = ROOT / "课题1-推进会-任务2-0302(1)_第7-8页修改版.pptx"

TEAL = RGBColor(0x00, 0x78, 0x73)
DEEP_TEAL = RGBColor(0x00, 0x69, 0x68)
BLUE = RGBColor(0x1F, 0x66, 0xB1)
GREEN = RGBColor(0x4F, 0x91, 0x35)
ORANGE = RGBColor(0xE9, 0x70, 0x0A)
RED = RGBColor(0xD9, 0x2D, 0x2D)
PURPLE = RGBColor(0x6D, 0x49, 0xA8)
BLACK = RGBColor(0x16, 0x20, 0x24)
GRAY = RGBColor(0x5B, 0x63, 0x69)
LIGHT_LINE = RGBColor(0xC9, 0xD5, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def remove_effects(shape):
    sp_pr = shape._element.spPr
    for tag in ("a:effectLst", "a:effectDag", "a:scene3d", "a:sp3d"):
        node = sp_pr.find(qn(tag))
        if node is not None:
            sp_pr.remove(node)


def set_east_asian_font(run, typeface="微软雅黑"):
    run.font.name = typeface
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", typeface)


def style_run(run, size, *, bold=False, color=BLACK, name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_east_asian_font(run, name)


def set_shape_text(
    shape,
    text,
    *,
    size,
    bold=False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margins=(0.05, 0.05, 0.05, 0.05),
    line_spacing=1.0,
):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    lines = text.split("\n")
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        style_run(run, size, bold=bold, color=color)


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size,
    bold=False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margins=(0.03, 0.03, 0.02, 0.02),
    line_spacing=1.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_text(
        shape,
        text,
        size=size,
        bold=bold,
        color=color,
        align=align,
        valign=valign,
        margins=margins,
        line_spacing=line_spacing,
    )
    return shape


def add_box(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill,
    line,
    radius=True,
    line_width=1.1,
):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    remove_effects(shape)
    return shape


def add_header_tag(slide, x, y, w, text, color):
    box = add_box(
        slide,
        x,
        y,
        w,
        0.38,
        fill=color,
        line=color,
        radius=True,
        line_width=0.8,
    )
    set_shape_text(
        box,
        text,
        size=16,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.02, 0.02, 0.01, 0.01),
    )
    return box


def add_summary_band(slide, label, text):
    tag = add_box(
        slide,
        0.36,
        1.50,
        1.00,
        0.82,
        fill=TEAL,
        line=TEAL,
        radius=False,
        line_width=0.8,
    )
    set_shape_text(
        tag,
        label,
        size=17,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.06, 0.06, 0.03, 0.03),
    )
    band = add_box(
        slide,
        1.42,
        1.50,
        11.48,
        0.82,
        fill=RGBColor(0xF4, 0xF7, 0xF7),
        line=LIGHT_LINE,
        radius=False,
        line_width=0.8,
    )
    set_shape_text(
        band,
        text,
        size=18,
        color=BLACK,
        align=PP_ALIGN.LEFT,
        margins=(0.18, 0.14, 0.06, 0.04),
        line_spacing=1.05,
    )
    return tag, band


def add_picture_contain(slide, blob, x, y, w, h):
    from PIL import Image

    image = Image.open(BytesIO(blob))
    image_ratio = image.width / image.height
    box_ratio = w / h
    if image_ratio > box_ratio:
        target_w = w
        target_h = w / image_ratio
        target_x = x
        target_y = y + (h - target_h) / 2
    else:
        target_h = h
        target_w = h * image_ratio
        target_x = x + (w - target_w) / 2
        target_y = y
    return slide.shapes.add_picture(
        BytesIO(blob),
        Inches(target_x),
        Inches(target_y),
        width=Inches(target_w),
        height=Inches(target_h),
    )


def add_bullet_lines(slide, x, y, w, lines, color):
    for index, line in enumerate(lines):
        dot = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Inches(x),
            Inches(y + index * 0.28 + 0.085),
            Inches(0.09),
            Inches(0.09),
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()
        add_text(
            slide,
            x + 0.16,
            y + index * 0.28,
            w - 0.16,
            0.26,
            line,
            size=13.5,
            color=BLACK,
            valign=MSO_ANCHOR.MIDDLE,
        )


def add_cross_layer_graph(slide, x, y, w, h):
    colors = (BLUE, GREEN, ORANGE)
    labels = ("主网", "配网", "微网")
    ys = (y + 0.28, y + 0.93, y + 1.58)
    for row, (row_y, color, label) in enumerate(zip(ys, colors, labels)):
        add_text(
            slide,
            x,
            row_y - 0.12,
            0.55,
            0.28,
            label,
            size=13,
            bold=True,
            color=color,
            align=PP_ALIGN.CENTER,
        )
        line = slide.shapes.add_connector(
            1,
            Inches(x + 0.70),
            Inches(row_y),
            Inches(x + w - 0.20),
            Inches(row_y),
        )
        line.line.color.rgb = color
        line.line.width = Pt(1.7)
        for node_index in range(4):
            node_x = x + 0.80 + node_index * ((w - 1.20) / 3)
            node = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                Inches(node_x - 0.07),
                Inches(row_y - 0.07),
                Inches(0.14),
                Inches(0.14),
            )
            node.fill.solid()
            node.fill.fore_color.rgb = color
            node.line.color.rgb = WHITE
            node.line.width = Pt(0.6)
    for boundary_y, boundary_label in ((y + 0.60, "主配接口"), (y + 1.25, "PCC")):
        stem = slide.shapes.add_connector(
            1,
            Inches(x + w * 0.62),
            Inches(boundary_y - 0.22),
            Inches(x + w * 0.62),
            Inches(boundary_y + 0.22),
        )
        stem.line.color.rgb = RED
        stem.line.width = Pt(1.2)
        ring = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Inches(x + w * 0.62 - 0.10),
            Inches(boundary_y - 0.10),
            Inches(0.20),
            Inches(0.20),
        )
        ring.fill.background()
        ring.line.color.rgb = RED
        ring.line.width = Pt(1.4)
        add_text(
            slide,
            x + w * 0.62 + 0.16,
            boundary_y - 0.12,
            0.82,
            0.24,
            boundary_label,
            size=11.5,
            bold=True,
            color=RED,
        )


def update_slide_7(prs):
    slide = prs.slides[6]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if len(pictures) < 2:
        raise RuntimeError("Slide 7 no longer contains the two reusable charts")
    duck_blob = min(pictures, key=lambda shape: shape.left).image.blob
    wave_blob = max(pictures, key=lambda shape: shape.top).image.blob

    for shape in list(slide.shapes)[4:]:
        remove_shape(shape)

    section = slide.shapes[3]
    set_shape_text(
        section,
        "任务难题",
        size=20,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.02, 0.02, 0.01, 0.01),
    )

    add_summary_band(
        slide,
        "核心\n难题",
        "分布式光伏、风电、储能与可控负荷的时空波动，与主网—配网—微网模型异构及边界耦合叠加，使潮流状态变化快、相互影响强。传统整体求解难以同时兼顾模型一致性、区域自治性和计算效率。",
    )

    panel_xs = (0.38, 4.64, 8.90)
    panel_w = 4.04
    panel_colors = (BLUE, GREEN, ORANGE)
    panel_titles = ("多源波动", "跨层异构", "边界耦合")
    for x, color, title in zip(panel_xs, panel_colors, panel_titles):
        panel = add_box(
            slide,
            x,
            2.55,
            panel_w,
            3.87,
            fill=WHITE,
            line=color,
            radius=True,
            line_width=1.2,
        )
        add_header_tag(slide, x + 0.12, 2.64, panel_w - 0.24, title, color)

    add_picture_contain(slide, duck_blob, 0.55, 3.10, 3.70, 2.18)
    add_bullet_lines(
        slide,
        0.67,
        5.42,
        3.55,
        ("净负荷低谷加深、反向爬坡增强", "不同时间断面的潮流状态快速变化"),
        BLUE,
    )

    add_cross_layer_graph(slide, 4.84, 3.02, 3.66, 2.10)
    add_bullet_lines(
        slide,
        4.93,
        5.42,
        3.42,
        ("相制式与设备模型存在差异", "参数基准、运行时标与边界变量难对齐"),
        GREEN,
    )

    add_picture_contain(slide, wave_blob, 9.10, 3.08, 3.64, 1.72)
    add_bullet_lines(
        slide,
        9.12,
        5.06,
        3.55,
        ("潮流方向随运行状态改变", "电压与功率波动经接口跨区域传播", "局部误差可能在协调迭代中累积"),
        ORANGE,
    )

    chain = add_box(
        slide,
        0.38,
        6.60,
        12.56,
        0.58,
        fill=RGBColor(0xF3, 0xF7, 0xF7),
        line=LIGHT_LINE,
        radius=True,
        line_width=0.9,
    )
    set_shape_text(
        chain,
        "源荷不确定性  →  模型与拓扑异构  →  边界耦合传播  →  收敛性与计算效率下降\n关键矛盾：局部子问题可以独立求解，但跨区边界状态必须保持一致",
        size=14,
        bold=True,
        color=DEEP_TEAL,
        align=PP_ALIGN.CENTER,
        margins=(0.04, 0.04, 0.03, 0.03),
        line_spacing=0.9,
    )


def add_stage(slide, number, x, title, subtitle, color, fill):
    box = add_box(
        slide,
        x,
        2.56,
        1.55,
        1.48,
        fill=fill,
        line=color,
        radius=True,
        line_width=1.15,
    )
    number_circle = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x + 0.10),
        Inches(2.68),
        Inches(0.32),
        Inches(0.32),
    )
    number_circle.fill.solid()
    number_circle.fill.fore_color.rgb = color
    number_circle.line.fill.background()
    set_shape_text(
        number_circle,
        str(number),
        size=12,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0, 0, 0, 0),
    )
    add_text(
        slide,
        x + 0.12,
        3.03,
        1.31,
        0.34,
        title,
        size=14.5,
        bold=True,
        color=color,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        x + 0.12,
        3.42,
        1.31,
        0.48,
        subtitle,
        size=11.2,
        color=BLACK,
        align=PP_ALIGN.CENTER,
        line_spacing=0.95,
    )
    return box


def add_innovation_panel(slide, x, title, body, color, fill, glyph):
    panel = add_box(
        slide,
        x,
        4.40,
        4.00,
        1.73,
        fill=fill,
        line=color,
        radius=True,
        line_width=1.15,
    )
    icon = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x + 0.18),
        Inches(4.60),
        Inches(0.56),
        Inches(0.56),
    )
    icon.fill.solid()
    icon.fill.fore_color.rgb = color
    icon.line.fill.background()
    set_shape_text(
        icon,
        glyph,
        size=18,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0, 0, 0, 0),
    )
    add_text(
        slide,
        x + 0.86,
        4.53,
        2.93,
        0.42,
        title,
        size=17,
        bold=True,
        color=color,
    )
    add_text(
        slide,
        x + 0.86,
        4.99,
        2.90,
        0.90,
        body,
        size=13.2,
        color=BLACK,
        valign=MSO_ANCHOR.TOP,
        line_spacing=1.0,
    )
    return panel


def update_slide_8(prs):
    slide = prs.slides[7]
    for shape in list(slide.shapes)[4:]:
        remove_shape(shape)

    section = slide.shapes[3]
    set_shape_text(
        section,
        "技术创新",
        size=20,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.02, 0.02, 0.01, 0.01),
    )

    add_summary_band(
        slide,
        "核心\n思路",
        "将全网强耦合潮流问题转化为拓扑完整的局部潮流子问题，以边界电压、主配交换功率和微网等值注入维持跨区一致性，通过层内并行与层间协调提升求解效率。",
    )

    stages = (
        ("统一建模", "语义、基准\n与时标对齐", BLUE, RGBColor(0xEE, 0xF4, 0xFC)),
        ("边界辨识", "电压分层、\n变压器与PCC", TEAL, RGBColor(0xEC, 0xF7, 0xF6)),
        ("关联度量", "阻抗、潮流\n与源荷注入", DEEP_TEAL, RGBColor(0xEC, 0xF5, 0xF5)),
        ("智能分区", "锚点生长、\n多级图划分", GREEN, RGBColor(0xF0, 0xF7, 0xED)),
        ("解耦建模", "节点复制、\n接口等值", ORANGE, RGBColor(0xFE, 0xF2, 0xE9)),
        ("协调迭代", "松弛更新、\n联合残差", TEAL, RGBColor(0xEC, 0xF7, 0xF6)),
        ("并行求解", "层内并行、\n层间协调", PURPLE, RGBColor(0xF3, 0xEF, 0xFA)),
    )
    stage_xs = [0.36 + index * 1.82 for index in range(7)]
    for index, (x, stage) in enumerate(zip(stage_xs, stages), start=1):
        title, subtitle, color, fill = stage
        add_stage(slide, index, x, title, subtitle, color, fill)
        if index < 7:
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                Inches(x + 1.57),
                Inches(3.16),
                Inches(0.22),
                Inches(0.20),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(0x7B, 0x8E, 0x95)
            arrow.line.fill.background()

    add_innovation_panel(
        slide,
        0.36,
        "模型层：形成统一计算对象",
        "消除设备语义、参数基准和运行时标差异\n为边界识别、网络分区和潮流求解提供一致输入",
        BLUE,
        RGBColor(0xF2, 0xF6, 0xFB),
        "模",
    )
    add_innovation_panel(
        slide,
        4.66,
        "结构层：形成可协调子区",
        "接口锚点约束边界位置，电气耦合指导区域划分\n保证区内强耦合、区间弱耦合和子区拓扑连通",
        GREEN,
        RGBColor(0xF3, 0xF8, 0xF0),
        "区",
    )
    add_innovation_panel(
        slide,
        8.96,
        "求解层：实现协同收敛",
        "边界电压、交换功率与等值注入双向传递\n层内并行计算，层间按变量依赖关系顺序协调",
        ORANGE,
        RGBColor(0xFE, 0xF5, 0xED),
        "解",
    )

    output = add_box(
        slide,
        0.36,
        6.38,
        12.60,
        0.70,
        fill=RGBColor(0xF1, 0xF7, 0xF6),
        line=TEAL,
        radius=True,
        line_width=1.0,
    )
    set_shape_text(
        output,
        "统一计算对象  →  可协调分区  →  拓扑完整的局部潮流子问题  →  边界一致的协调解  →  可扩展并行求解\n形成“建模—分区—解耦—协调—并行”的完整技术闭环",
        size=14.5,
        bold=True,
        color=DEEP_TEAL,
        align=PP_ALIGN.CENTER,
        margins=(0.04, 0.04, 0.03, 0.03),
        line_spacing=0.92,
    )


def fix_slide_15_borders(prs):
    slide = prs.slides[14]
    # The saved deck has five pairs of outer rounded cards and shallow header strips.
    card_pairs = (
        (5, 6, BLUE),
        (10, 11, GREEN),
        (15, 16, ORANGE),
        (58, 59, BLUE),
        (63, 64, TEAL),
    )
    resolved_pairs = [
        (slide.shapes[outer_index], slide.shapes[header_index], color)
        for outer_index, header_index, color in card_pairs
    ]
    for outer, header, color in resolved_pairs:

        # Keep the card itself as a transparent foreground outline so its arcs
        # cannot be hidden by the shallow title strip.
        outer.fill.background()
        outer.line.color.rgb = color
        outer.line.width = Pt(1.45)
        remove_effects(outer)

        header.left = outer.left + Inches(0.10)
        header.top = outer.top + Inches(0.055)
        header.width = outer.width - Inches(0.20)
        header.height = Inches(0.255)
        header.line.fill.background()
        remove_effects(header)

        sp_tree = slide.shapes._spTree
        outer_element = outer._element
        sp_tree.remove(outer_element)
        sp_tree.append(outer_element)


def restore_untouched_slide_parts(source, output):
    changed_slides = {7, 8, 15}
    members = set()
    for slide_number in range(1, 22):
        if slide_number in changed_slides:
            continue
        members.add(f"ppt/slides/slide{slide_number}.xml")
        members.add(f"ppt/slides/_rels/slide{slide_number}.xml.rels")

    replacement = output.with_suffix(".restored.pptx")
    with ZipFile(source) as source_zip, ZipFile(output) as output_zip, ZipFile(
        replacement, "w", ZIP_DEFLATED
    ) as replacement_zip:
        source_names = set(source_zip.namelist())
        for item in output_zip.infolist():
            data = output_zip.read(item.filename)
            if item.filename in members and item.filename in source_names:
                data = source_zip.read(item.filename)
            replacement_zip.writestr(item, data)
    replacement.replace(output)


def verify_only_requested_slides_changed(source, output):
    with ZipFile(source) as source_zip, ZipFile(output) as output_zip:
        for slide_number in range(1, 22):
            if slide_number in {7, 8, 15}:
                continue
            member = f"ppt/slides/slide{slide_number}.xml"
            if source_zip.read(member) != output_zip.read(member):
                raise RuntimeError(f"Slide {slide_number} changed unexpectedly")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    prs = Presentation(SOURCE)
    if len(prs.slides) != 21:
        raise RuntimeError(f"Expected 21 slides, found {len(prs.slides)}")

    update_slide_7(prs)
    update_slide_8(prs)
    fix_slide_15_borders(prs)
    prs.save(OUTPUT)
    restore_untouched_slide_parts(SOURCE, OUTPUT)
    verify_only_requested_slides_changed(SOURCE, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
