from __future__ import annotations

import shutil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "课题1-推进会-任务2-0302(1).pptx"
BACKUP = ROOT / "课题1-推进会-任务2-0302(1)_原始备份.pptx"
OUTPUT = ROOT / "课题1-推进会-任务2-0302(1)_修改版.pptx"

TEAL = RGBColor(0x00, 0x77, 0x73)
BLUE = RGBColor(0x1F, 0x66, 0xB1)
GREEN = RGBColor(0x4F, 0x91, 0x35)
ORANGE = RGBColor(0xE9, 0x70, 0x0A)
RED = RGBColor(0xD9, 0x2D, 0x2D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def first_run(text_frame):
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            return run
    return None


def set_east_asian_font(run, typeface="微软雅黑"):
    run.font.name = typeface
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(
            r_pr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea"
        )
    ea.set("typeface", typeface)


def replace_text_in_place(shape, text: str):
    """Replace a single text block while retaining its geometry and paragraph style."""
    if not shape.has_text_frame:
        raise ValueError(f"Shape {shape.name} has no text frame")
    runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        shape.text_frame.paragraphs[0].add_run().text = text


def set_shape_text(
    shape,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = BLACK,
    align=PP_ALIGN.LEFT,
):
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_east_asian_font(run)


def set_cell_text(
    cell,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = BLACK,
    align=PP_ALIGN.LEFT,
):
    text_frame = cell.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 0.95
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        set_east_asian_font(run)


def set_cell_rich_text(
    cell,
    lead: str,
    body: str,
    *,
    size: float,
    color: RGBColor = BLACK,
):
    text_frame = cell.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    for text, bold in ((lead, True), (body, False)):
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        set_east_asian_font(run)


def set_mixed_count_line(shape, prefix: str, count: str, suffix: str):
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for value, color, bold in (
        (prefix, BLACK, False),
        (count, RGBColor(0xC0, 0x00, 0x00), True),
        (suffix, BLACK, False),
    ):
        run = paragraph.add_run()
        run.text = value
        run.font.size = Pt(18)
        run.font.bold = bold
        run.font.color.rgb = color
        set_east_asian_font(run)


def update_slide_9(prs: Presentation):
    slide = prs.slides[8]
    table_shape = next(shape for shape in slide.shapes if getattr(shape, "has_table", False))
    table = table_shape.table

    headers = ["研究内容", "预期成果", "取得的成果及创新特色", "成果证明", "完成情况"]
    for col, text in enumerate(headers):
        set_cell_text(table.cell(0, col), text, size=16.5, bold=True, align=PP_ALIGN.CENTER)

    rows = [
        (
            "1. 异构模型统一建模与边界辨识",
            "形成主网、配网和微网设备语义映射、参数基准统一及主配接口/PCC自动识别方法。",
            "建立“模型标准化—设备映射—电压分层—跨域变压器识别—PCC映射”流程，形成分区与潮流求解共享的统一计算对象。",
            "方法流程与原型功能已形成",
            "方法研究与原型实现已完成",
        ),
        (
            "2. 主配微层级智能分区",
            "提出兼顾接口约束、电气耦合、区域连通与规模均衡的主配微分区方法。",
            "形成接口锚点区域生长与配网加权图多级缩约方法，经粗图初分、逐级投影、局部优化和连通修复获得可协调子区。",
            "分区算法与阶段验证材料已形成",
            "核心算法已完成",
        ),
        (
            "3. 边界等值与协调迭代",
            "形成边界节点复制、接口等值及边界电压—交换功率协调收敛方法。",
            "构建主网、配网和微网局部子问题，交替传递边界电压、主配交换功率与微网等值注入，并以松弛更新和功率—电压联合残差判定收敛。",
            "发明专利成果1项（申请文本已形成）",
            "协调机制已完成",
        ),
        (
            "4. 分层并行潮流与工程验证",
            "形成分层并行潮流求解模块，完成实际电网数据上的精度、收敛性与效率验证。",
            "完成主网子区、配网实例与微网实例的独立任务封装，支持层内并行、层间顺序协调和动态进程配置；正在结合石家庄数据开展综合验证。",
            "论文成果2篇（均已发表）",
            "模块已形成，工程验证推进中",
        ),
    ]

    for row_index, values in enumerate(rows, start=1):
        for col_index, text in enumerate(values):
            if col_index == 3:
                size = 11.3
            elif col_index in (1, 2):
                size = 12.5
            else:
                size = 13.0
            align = PP_ALIGN.LEFT if col_index != 4 else PP_ALIGN.CENTER
            set_cell_text(
                table.cell(row_index, col_index),
                text,
                size=size,
                bold=False,
                align=align,
            )


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def set_solid_line(shape, color: RGBColor, width: float):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def set_card_style(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    set_solid_line(shape, color, 1.35)


def set_header_style(shape, color: RGBColor, tint: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = tint
    shape.line.fill.background()


def update_slide_15(prs: Presentation):
    slide = prs.slides[14]

    # Remove the stray reviewer remnant above the slide title.
    for shape in list(slide.shapes):
        if shape.shape_id == 11 and getattr(shape, "text", "").strip() == "1.":
            remove_shape(shape)

    by_id = {shape.shape_id: shape for shape in slide.shapes}
    card_specs = [
        (14, 15, BLUE, RGBColor(0xEE, 0xF4, 0xFC)),
        (19, 20, GREEN, RGBColor(0xF0, 0xF7, 0xED)),
        (24, 25, ORANGE, RGBColor(0xFE, 0xF2, 0xE9)),
        (68, 69, BLUE, RGBColor(0xEE, 0xF4, 0xFC)),
        (73, 74, TEAL, RGBColor(0xEC, 0xF7, 0xF6)),
    ]
    for card_id, header_id, color, tint in card_specs:
        set_card_style(by_id[card_id], color)
        set_header_style(by_id[header_id], color, tint)

    # Make the return path visually enter Step 1 instead of ending in empty space.
    for shape_id in (94, 95, 96, 97):
        shape = by_id.get(shape_id)
        if shape is not None:
            remove_shape(shape)

    return_vertical = slide.shapes.add_connector(
        1, Inches(0.30), Inches(3.82), Inches(0.30), Inches(6.79)
    )
    set_solid_line(return_vertical, RED, 1.35)
    return_bottom = slide.shapes.add_connector(
        1, Inches(0.30), Inches(6.79), Inches(0.78), Inches(6.79)
    )
    set_solid_line(return_bottom, RED, 1.35)
    return_up = slide.shapes.add_connector(
        1, Inches(0.78), Inches(6.79), Inches(0.78), Inches(6.43)
    )
    set_solid_line(return_up, RED, 1.35)
    arrow = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
        Inches(0.29),
        Inches(3.71),
        Inches(0.29),
        Inches(0.22),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()

    # Clarify the decision branch without moving its existing box.
    by_id = {shape.shape_id: shape for shape in slide.shapes}
    if 92 in by_id:
        set_shape_text(
            by_id[92],
            "否则：返回 1",
            size=6.0,
            bold=True,
            color=RED,
            align=PP_ALIGN.CENTER,
        )


def update_slide_17(prs: Presentation):
    slide = prs.slides[16]
    patent_summary = next(shape for shape in slide.shapes if getattr(shape, "text", "").startswith("已受理/授权"))
    paper_summary = next(shape for shape in slide.shapes if getattr(shape, "text", "").startswith("已投稿"))
    set_mixed_count_line(patent_summary, "已形成发明专利成果 ", "1", " 项。")
    set_mixed_count_line(paper_summary, "已发表论文 ", "2", " 篇。")

    tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
    patent_table, paper_table = tables

    patent_values = [
        ["序号", "专利名称", "成果状态", "完成时间"],
        ["1", "一种面向数据异构的元素级顺序解耦联邦学习方法", "申请文本已形成", "2026.08"],
    ]
    for row_index, row in enumerate(patent_values):
        for col_index, text in enumerate(row):
            set_cell_text(
                patent_table.cell(row_index, col_index),
                text,
                size=17 if row_index == 0 else 15,
                bold=row_index == 0,
                align=PP_ALIGN.CENTER if row_index == 0 or col_index != 1 else PP_ALIGN.LEFT,
            )

    paper_values = [
        ["序号", "论文名称", "期刊", "发表时间"],
        [
            "1\n2",
            "SeqFedEDT: Element-Wise Decoupled Training\nSeqFedRPC: Regularized Parameter Clustering",
            "IEEE TMC\nTETT",
            "2026.05\n2026.07",
        ],
    ]
    for row_index, row in enumerate(paper_values):
        for col_index, text in enumerate(row):
            if row_index == 0:
                size = 17
            elif col_index == 1:
                size = 13.0
            else:
                size = 14.5
            set_cell_text(
                paper_table.cell(row_index, col_index),
                text,
                size=size,
                bold=row_index == 0,
                align=PP_ALIGN.CENTER if row_index == 0 or col_index != 1 else PP_ALIGN.LEFT,
            )


def update_slide_19(prs: Presentation):
    slide = prs.slides[18]
    table = next(shape.table for shape in slide.shapes if getattr(shape, "has_table", False))
    set_cell_text(
        table.cell(0, 1), "2026年9月", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    set_cell_text(
        table.cell(0, 2), "2026年12月", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER
    )
    set_cell_rich_text(
        table.cell(1, 1),
        "2.2.1. 完善主配微层级分区与边界协调算法。",
        "面向石家庄实际电网数据，完成模型转换、主配接口识别与分区结果校核；针对弱耦合边界和不平衡分区，优化锚点选种、多级缩约及连通修复策略，形成可重复的分区解耦流程。",
        size=16,
    )
    set_cell_rich_text(
        table.cell(1, 2),
        "2.2.2. 完成主配微分层并行潮流与多时间断面验证。",
        "完善边界电压—交换功率—微网等值注入的协调迭代与残差判据，研究任务调度、热启动和异常回退机制；基于石家庄数据开展精度、收敛性与计算效率对比，形成算法模块及验证材料。",
        size=16,
    )


def clean_slide_16(prs: Presentation):
    slide = prs.slides[15]
    for shape in list(slide.shapes):
        if shape.shape_id == 11 and getattr(shape, "text", "").strip() == "1.":
            remove_shape(shape)


def update_slide_20(prs: Presentation):
    slide = prs.slides[19]
    replacements = {
        "在已有层级分区解耦方法的基础上": (
            "面向石家庄实际电网数据，完善“电压等级分层—跨域变压器识别—PCC映射—接口锚点选种—"
            "电气耦合分区”的完整流程；研究边界节点复制与接口等值模型，校核子区连通性、规模均衡性"
            "和边界可协调性，形成可解释、可复现的主配微层级分区结果。"
        ),
        "在分区解耦框架下": (
            "完善主网子区、配网实例与微网实例的独立任务封装及分层并行调度，建立按计算资源动态配置"
            "工作进程、层内同步汇总和层间顺序协调机制；结合热启动、超时回退与负载均衡，开展串并行"
            "对比，评价求解精度、收敛稳定性和计算效率。"
        ),
        "针对主网、配网和微网": (
            "统一主网正序、配网三相和微网局部潮流的数据契约与状态量映射，完善边界电压、交换功率及"
            "微网等值注入的双向传递；研究松弛参数自适应与功率—电压联合残差判据，解决异构求解器"
            "协同过程中的振荡、失配和局部不收敛问题。"
        ),
    }
    replaced = 0
    for shape in iter_shapes(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        current = shape.text.strip()
        for prefix, replacement in replacements.items():
            if current.startswith(prefix):
                replace_text_in_place(shape, replacement)
                replaced += 1
                break
    if replaced != 3:
        raise RuntimeError(f"Expected 3 slide 20 text replacements, made {replaced}")


def enlarged_size(size: float) -> float:
    if size < 3.0:
        return size
    if size < 4.5:
        return min(4.8, size + 1.0)
    if size < 6.0:
        return size + 1.2
    if size < 8.0:
        return size + 1.5
    if size < 10.0:
        return size + 1.3
    if size < 12.0:
        return size + 1.0
    if size < 16.0:
        return size + 0.8
    return size


def enlarge_technical_slides(prs: Presentation):
    # Slides 10-16 are the seven method pages. Geometry remains unchanged.
    for slide_number in range(10, 17):
        slide = prs.slides[slide_number - 1]
        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None:
                            run.font.size = Pt(enlarged_size(run.font.size.pt))
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if run.font.size is not None:
                                    run.font.size = Pt(enlarged_size(run.font.size.pt))


def zip_member_bytes(path: Path, member: str) -> bytes:
    with ZipFile(path) as archive:
        return archive.read(member)


def restore_untouched_slide_xml(source: Path, output: Path):
    members_to_restore = {
        "ppt/slides/slide7.xml",
        "ppt/slides/slide8.xml",
        "ppt/slides/_rels/slide7.xml.rels",
        "ppt/slides/_rels/slide8.xml.rels",
    }
    replacement = output.with_suffix(".restored.pptx")
    with ZipFile(source) as source_zip, ZipFile(output) as output_zip, ZipFile(
        replacement, "w", ZIP_DEFLATED
    ) as replacement_zip:
        source_names = set(source_zip.namelist())
        for item in output_zip.infolist():
            data = output_zip.read(item.filename)
            if item.filename in members_to_restore and item.filename in source_names:
                data = source_zip.read(item.filename)
            replacement_zip.writestr(item, data)
    replacement.replace(output)


def verify_untouched_slides(source: Path, output: Path):
    # Exact XML equality provides a stronger guard than comparing extracted text.
    for slide_number in (7, 8):
        member = f"ppt/slides/slide{slide_number}.xml"
        if zip_member_bytes(source, member) != zip_member_bytes(output, member):
            raise RuntimeError(f"Slide {slide_number} changed unexpectedly")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    prs = Presentation(SOURCE)
    if len(prs.slides) != 21:
        raise RuntimeError(f"Expected 21 slides, found {len(prs.slides)}")

    update_slide_9(prs)
    update_slide_15(prs)
    update_slide_17(prs)
    update_slide_19(prs)
    update_slide_20(prs)
    clean_slide_16(prs)
    enlarge_technical_slides(prs)
    prs.save(OUTPUT)
    restore_untouched_slide_xml(SOURCE, OUTPUT)
    verify_untouched_slides(SOURCE, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
