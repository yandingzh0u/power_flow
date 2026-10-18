from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("主配微分区解耦潮流计算技术报告.docx")


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld)
    run._r.append(instr)
    run._r.append(fld2)


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.4)
    sec.header_distance = Cm(1.4)
    sec.footer_distance = Cm(1.2)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.first_line_indent = Pt(21)
    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(6)
    header = sec.header.paragraphs[0]
    header.text = "主配微分区解耦与并行潮流计算技术报告"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in header.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(sec.footer.paragraphs[0])


def p(doc, text):
    para = doc.add_paragraph(text)
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style=None)
        para.paragraph_format.first_line_indent = Pt(0)
        para.add_run("（%s）" % item[0]).bold = True
        para.add_run(item[1])


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], str(val))
    return t


sections = [
("摘要", [
"本报告面向当前目录下“主配微分区解耦与并行潮流计算”项目代码和河北示范区数据集，形成一份可用于课题汇报、软件验收和后续工程联调的技术报告。项目以 Python 为主要实现语言，围绕主网、配网、微网多层级耦合潮流计算，构建了数据转换、拓扑识别、区域划分、边界等值、并行求解和 HTTP API 服务的一体化计算模块。",
"从代码结构看，gmss/core.py 承载核心算法，包括 MATPOWER/BPA 数据解析、pandapower 主网建模、OpenDSS 配网/微网模型调用、主配边界识别、ADMM/硬分区潮流求解、METIS 配网图划分、微网等值注入迭代和结果汇总；gmss/service.py 与 gmss/http_api.py 提供外部调用适配和标准 JSON 接口；tools/convert_demo_feeders.py、tools/convert_bpa_to_pandapower.py 用于河北示范区数据转换。项目同时保留了功能概述、API 输入输出说明、年中科技报告模板和示范区原始/转换数据，为技术报告撰写提供了较完整的工程依据。",
"本报告结合项目代码、既有说明文档和公开文献，对主配微耦合潮流的建模逻辑、分区解耦算法、并行求解流程、接口设计、数据基础、测试结果和工程化改进方向进行系统阐述。文献方面主要引用了主配协同潮流与分布式计算、pandapower、MATPOWER、OpenDSS、METIS 图划分和 Texas A&M 合成电网测试集等资料，以支撑项目技术路线的合理性和可扩展性。"
]),
("1 项目背景与技术定位", [
"随着分布式光伏、储能、电动汽车、柔性负荷和微网等资源快速接入，配电网逐渐从被动受电网络转变为具有局部自治、双向潮流和可调节能力的主动网络。传统调度潮流计算通常以主网为核心，配网侧被等值为负荷或少量注入；传统配网潮流又往往在馈线或局部区域内部独立完成，难以准确反映主配边界电压、配微 PCC 注入以及多层级资源响应之间的交互。对于主配微协同运行场景，单纯整网一次性拼接计算会面临模型规模大、数据来源异构、接口职责不清、并行能力不足和工程联调困难等问题。",
"本项目的技术定位是提供一个确定性的主配微分区解耦与并行潮流计算模块。它不是替代 D6.0、配电自动化系统或微网 EMS 的基础平台，而是在既有模型、断面和量测数据之上，补充主配微接口映射、区域自动划分、边界解耦、配网/微网并行潮流和标准 API 输出能力。其计算结果可作为安全校核、主配协同调控、配网重过载治理、电压无功优化和微网接入评估的基础状态量。",
"项目当前已经形成两类应用对象：一类是 ACTIVSg 等合成主网与 IEEE13/OpenDSS 配网组合构成的大规模算法验证算例，另一类是河北示范区 BPA DAT 主配网断面和 10kV 馈线 XLS 负荷曲线构成的工程示范数据。前者用于验证规模扩展、主配多接口耦合与大规模分区并行能力；后者用于支撑课题汇报、API 联调和工程场景说明。"
]),
("2 本地项目代码与资料梳理", [
"当前目录下的代码和资料可以分为六类。第一类是计算核心代码，位于 gmss 包内；第二类是启动和调用入口，包括 start_api.py、call_api.py 和 scripts/run_gmss.py；第三类是数据转换工具，位于 tools 目录；第四类是算例数据，位于 data 目录；第五类是既有 Word/PPT 技术材料；第六类是缓存与参考数据，包括转换后的 pandapower JSON、OpenDSS 文件、微网缓存和分区 YAML。",
"gmss/core.py 是项目的技术主体。文件中定义了 ADMMConfig、PartitionSpec 等基础结构，提供 wrap_angle_deg、vm_va_to_complex、complex_to_vm_va 等电压相角工具，封装 runpp_safe 以适配 pandapower/numba 环境，并实现从 MATPOWER .m 到 pandapower 网络、从 OpenDSS 配网到等值注入、从主配边界电压到分区局部潮流的完整求解链条。代码中保留了硬分区和软边界两套思路，默认采用更稳定的 hard partition；软边界通过等效阻抗参数 RHO_SOFT 支持进一步试验。",
"gmss/service.py 负责将外部紧凑请求转换为内部求解配置，并将内部结果压缩为面向调用方的标准响应。该层设计体现了工程服务化思路：调用方不需要逐条传入线路、变压器、母线、微网内部设备等明细，只需要传入任务编号、数据集、断面、时刻和求解选项。gmss/http_api.py 基于 Python 标准库 ThreadingHTTPServer 实现健康检查、契约查询、请求规范化和计算运行四类 HTTP 接口，便于在无复杂 Web 框架依赖的情况下快速联调。",
"tools/convert_demo_feeders.py 用于解析河北示范区 10kV 馈线 XLS 负荷曲线，提取线路节点、设备类型、长度、配变容量和 96 点有/无功曲线，并输出 demo_area_feeders.json。tools/convert_bpa_to_pandapower.py 用于将 BPA DAT 断面转换为 pandapower 模型，并保留基础结构检查、卡片数量、母线/支路/发电机数量和基准潮流收敛结果。"
]),
("3 数据基础与算例规模", [
"河北示范区数据集包含 13 点和 20 点两个运行断面。转换摘要显示，13 点断面包含 2121 个 pandapower 母线、2036 条线路、73 台变压器、18 个阻抗支路、698 个负荷、121 个 sgen、49 个 gen 和 1 个外部电网；基础 pandapower 潮流收敛，迭代 4 次，电压幅值范围约为 0.9423 至 1.0354 pu。20 点断面包含 2122 个母线、2038 条线路、73 台变压器、18 个阻抗支路、664 个负荷、6 个 sgen、48 个 gen 和 1 个外部电网；基础潮流同样收敛，迭代 4 次，电压幅值范围约为 0.9378 至 1.0345 pu。",
"从年中科技报告和 API 说明可见，模块依据电压等级和变压器连接关系识别主配边界。在河北示范区算例中，35 kV 及以上母线归入主网集合，低于 35 kV 母线归入配网集合；若变压器两端分别属于主网和配网，则被识别为主配边界支路。既有材料给出的识别结果为两个断面均包含 67 个主网母线、70 台主配边界变压器和 49 个配网侧边界母线，说明边界识别在不同运行断面下具有结构稳定性。",
"大规模合成验证数据以 ACTIVSg25k 主网为基础，配接 500 个 IEEE13/OpenDSS 配网等值模型，形成 case_ACTIVSg25k_plus_500D_ieee.json 和对应映射 YAML。INDEX_synth_multi.json 中记录了 500 个主配接口、每个配网侧 4.16 kV 母线和 25 MVA 主配变压器容量等信息。该数据主要用于检验算法在大规模主配接口场景下的扩展能力。"
]),
("4 理论模型：主配微耦合潮流", [
"主配微系统可抽象为三层网络：主网层 T、配网层 D 和微网层 M。主网层包含高压母线、发电机、主网线路、主变高压侧和平衡节点；配网层包含 10 kV/低压母线、配电线路、配变、馈线负荷和分布式资源；微网层包含 PCC、微网内部线路、负荷、光伏、储能、柴油机或可控逆变器等设备。三层系统通过少量物理接口耦合，而不是在所有节点之间任意交互。",
"主配接口由主配变压器确定，其耦合变量主要包括边界电压幅值 Vb、边界电压相角 θb、边界有功 Pb 和边界无功 Qb。主网向配网提供边界电压，配网向主网反馈等值功率。配微接口由 PCC 母线确定，微网在给定 PCC 电压下独立求解，返回对配网等效的有功/无功注入。以函数形式表示，主网潮流可写作 F_T(x_T, S_TD)=0，配网第 k 个分区可写作 F_D^k(x_D^k, V_b^k, S_MG^k)=0，微网第 m 个系统可写作 F_M^m(x_M^m, V_pcc^m)=0。",
"该建模方式的核心优势在于边界变量清晰、模型职责分离、数据接口可控。主网求解器不需要掌握每条低压馈线的内部细节，配网求解器也不需要重复求解上级大电网；微网可以以内部详细模型或 PCC 等值模型方式接入。对于工程系统而言，这种分层模型更符合省调、地调、配电自动化和微网 EMS 的数据权限与运行边界。"
]),
("5 主配边界识别与拓扑分层", [
"项目将 pandapower 网络转换为图模型 G=(V,E)，其中母线为节点，线路、变压器和阻抗支路为边。节点属性包含电压等级、负荷、发电、名称和原始数据映射；边属性包含设备类型、阻抗、容量、开断状态和基础潮流信息。主配分层首先依据电压等级形成主网集合 V_T 和配网集合 V_D，然后遍历变压器支路识别跨层连接。",
"自动边界识别比人工维护接口表更适合多断面、多模型版本场景。主配网模型一旦因检修、拓扑调整或数据转换更新而发生变化，人工接口清单容易滞后；而基于电压等级和变压器两端属性的识别方法能够直接随模型重算。对于河北示范区算例，两个断面在负荷和潮流上有差异，但主配边界数量和主网母线数量保持一致，体现了结构识别对运行断面的鲁棒性。",
"在实际工程应用中，单纯电压阈值仍需结合变电站资产台账和调度分区规则进行校核。例如部分 35 kV 网络可能在不同地区承担主网、配网或过渡层角色；少数专线、用户站或新能源升压站也可能导致电压等级与调度归属不完全一致。因此，建议后续在当前自动识别规则基础上增加白名单/黑名单、变电站归属、设备类型和运维单位等辅助条件。"
]),
("6 配网 METIS 智能分区方法", [
"主配边界切开后，配网子图通常会形成多个物理连通分量。若仅按连通分量求解，部分大分量仍可能造成计算负载集中；若盲目按节点编号切分，又会破坏电气拓扑连通性并增加边界耦合误差。因此项目在物理解耦基础上引入图划分算法，对较大配网连通分量进行二次分区。",
"METIS 的多层 k 路图划分思想是先通过图粗化降低规模，在粗图上求初始划分，再逐层展开并局部优化。Karypis 和 Kumar 的研究表明，多层 k 路划分能够在线性量级时间内处理大规模不规则图，并在边割质量与计算时间之间取得较好平衡[8]。电网拓扑具有稀疏、局部连通、支路参数差异明显等特点，适合转化为加权图后使用该类算法进行区域划分。",
"项目中的节点权重反映局部计算负荷，可由母线负荷规模、接入元件数量或预估方程规模构造；边权重反映电气耦合强度，可综合线路阻抗、变压器容量和基准潮流。目标函数可概括为：在各分区节点权重基本均衡的约束下，最小化跨分区边权重之和。边权越高表示越不希望被切开，从而使分区边界尽量落在弱耦合位置。",
"分区后还必须进行电力系统语义校验。项目对 METIS 返回结果执行连通性检查，对不连通片段进行修复，对极小分区进行合并，避免出现无法独立求解或工程意义较弱的碎片区域。这一处理对工程落地很重要，因为图划分算法本身只优化图指标，不天然保证潮流求解所需的边界母线、参考电压、拓扑连通和设备完整性。"
]),
("7 分区解耦潮流求解流程", [
"项目支持主网直接求解、主网 ADMM/分区求解、配网分区求解和微网独立求解等多种组合。河北示范区默认配置为主网 pandapower_direct、主网分区 none、配网 partitioned、配网分区 metis、配网局部并行 true、微网并行 true。这个配置符合示范区主网规模相对较小、配网节点多且适合并行拆分的特点。",
"一次典型计算流程如下。第一步，读取或构建指定断面的 pandapower 主配网模型，并按 time_point 回写 10 kV 馈线负荷曲线。第二步，识别主配边界和配网侧边界母线，切分主网与配网。第三步，对主网进行潮流求解，得到边界电压幅值和相角。第四步，对配网连通分量进行 METIS 分区，将边界电压写入局部网络，使用进程池并行求解各分区潮流。第五步，在启用微网时，以 PCC 电压为输入并行求解微网，得到等值 P/Q 注入，回写配网后按设定次数进行配微耦合迭代。第六步，汇总全网状态、分区结果、边界交互量、收敛标志和误差指标，形成 API 响应。",
"硬分区模式的工程含义是：边界母线电压作为局部子问题的固定边界条件，分区内部状态独立求解，分区间通过边界电压和等值注入交换信息。该方法计算稳定、接口清晰，适合作为示范区软件服务默认路径。软边界或 ADMM 模式则更强调分区间一致性迭代，适合进一步研究高精度弱耦合场景，但需要更细致的惩罚参数、收敛准则和异常处理策略。"
]),
("8 微网等值接入与配微耦合", [
"河北示范区当前未提供真实微网内部台账，因此项目采用可替换的等值微网模型验证主配微接口和并行求解流程。代码中 build_tr66_microgrid、build_tr66_list_for_D、solve_MG_cluster_at_voltage_and_get_sum 等函数体现了微网模型构建、批量接入和 PCC 等值求解逻辑。调用方可以通过 enable_microgrid、demo_mg_count、mg_bus_count、mg_parallel、mg_coupling_iterations 等参数控制微网接入数量、内部规模、并行方式和配微迭代次数。",
"微网等值的基本思想是在 PCC 电压给定时独立求解微网内部潮流，得到 PCC 处对外表现为注入或负荷的 P/Q。对于配网而言，微网可等效为一个电压相关注入单元；对于微网而言，上级配网仅表现为 PCC 电压边界。该方式符合微网 EMS 与配网调度之间常见的信息交换边界，也为后续替换真实微网模型保留了接口稳定性。",
"后续如果获取真实微网台账，可在不改变 API 和主流程的前提下替换内部模型，包括光伏逆变器容量、储能 SOC 约束、柴油机出力上下限、负荷曲线、无功控制策略和孤网/并网切换逻辑。若进一步用于控制策略校核，还应引入逆变器 Volt/Var、PQ 限幅、储能充放电效率、低电压穿越和保护约束等模型。"
]),
("9 软件架构与 API 设计", [
"项目采用“核心算法层、服务适配层、HTTP 接口层、数据/缓存层”的轻量化架构。核心算法层集中在 gmss/core.py，提供可被脚本、Notebook 或服务直接调用的计算函数。服务适配层集中在 gmss/service.py，负责将外部紧凑请求归一化为内部 dataset_config 和 solver_config，并将内部结果整理为面向调用方的 summary/result/artifacts。HTTP 接口层集中在 gmss/http_api.py，通过 /api/v1/health、/api/v1/contract、/api/v1/gmss/normalize 和 /api/v1/gmss/run 暴露服务能力。数据/缓存层则位于 data 目录，保存原始数据、转换结果、模型缓存、微网缓存和分区结果。",
"API 输入字段设计强调任务追溯和配置可控。task_id 用于追踪一次计算任务，model_id 用于标识模型，dataset_id 指定 demo_area 或合成数据集，scenario_id 指定 13 点或 20 点断面，time_point 指定 96 点负荷曲线时刻。求解参数包括 t_solver、t_partition_mode、d_solver、d_partition_mode、d_n_parts、d_target_buses_per_part、d_local_parallel、mg_parallel、mg_coupling_iterations、alpha、beta 和 max_iter 等。",
"API 输出分为三层。第一层是 status、task_id、model_id、scenario_id 和 schema_version，用于服务调用状态与版本追踪。第二层是 summary，包含数据集、母线数量、线路数量、主配边界数量、配网分区数量、微网数量、分区是否收敛、METIS 分区统计和求解配置。第三层是 result 与 artifacts，保存最后状态量、潮流结果、边界交互量和结果文件路径。该输出结构便于外部系统既快速读取摘要，也能按需下载详细结果。"
]),
("10 关键技术实现分析", [
"数据解析方面，项目同时处理 MATPOWER .m、BPA DAT、pandapower JSON、OpenDSS DSS、XLS 馈线曲线和 YAML/JSON 映射文件。MATPOWER 解析函数 parse_matpower_m_case 通过正则提取 baseMVA、bus、gen、branch 和 gencost 等矩阵，并对 branch 规范列进行清理；BPA 转换工具则面向河北示范区 DAT 断面，生成 pandapower 模型和转换摘要。多格式支持使项目能够同时服务学术算例、合成算例和工程示范数据。",
"潮流求解方面，项目以 pandapower 作为主网和统一模型的主要求解工具，以 OpenDSS/dss-python 作为配网和微网建模求解的重要支撑。pandapower 基于 pandas 数据结构和改进的 Newton-Raphson 潮流实现，适合自动化静态和准静态分析[4][5]；OpenDSS 面向配电系统仿真和 DER 接入研究，具有灵活的脚本和外部程序接口[6][7]。两者结合可以兼顾主网平衡潮流和配网模型表达能力。",
"并行方面，代码使用 concurrent.futures.ProcessPoolExecutor，并通过 make_process_pool 兼容 fork 上下文；_collect_results 函数统一串行和并行结果收集逻辑；微网和配网分区均可独立提交进程池。由于潮流计算属于 CPU 密集型且各分区之间数据交互较少，进程级并行比线程级并行更容易获得实际加速，同时避免 Python GIL 对计算密集任务的限制。",
"稳定性方面，runpp_safe 根据 numba 是否可用自动设置 pandapower 参数，减少环境差异导致的告警；边界 sgen 注入通过统一前缀清理，避免多次迭代时重复叠加；角度处理函数使用 wrap_angle_deg 和 signed_angle_diff_deg 避免相角跨越 ±180 度时的误差突变；微网缓存版本号 MG_CACHE_VERSION 为模型缓存演进预留了版本控制入口。"
]),
("11 验证结果与工程可运行性", [
"已有转换摘要表明，河北示范区 13 点和 20 点基础主配网 pandapower 潮流均可收敛，且收敛迭代次数均为 4 次。这说明原始 BPA DAT 数据转换后的网络模型在设备参数、基准容量、母线类型、发电机/外部电网设置和线路变压器连接方面具备基本一致性。13 点断面的最低电压约 0.9423 pu，20 点断面的最低电压约 0.9378 pu，两个断面电压范围均处于可用于进一步分区校核的合理区间。",
"功能文档显示，示范区两个断面均识别到 70 台主配边界变压器和 49 个配网侧边界母线。该结果对后续并行分区具有关键意义：边界母线数量决定主配电压传递和等值功率反馈维度，边界变压器数量决定主配接口设备校核范围。分区算法需要保证这些边界不被遗漏，并在结果中保留接口映射关系。",
"从软件可运行性看，项目已经具备本地函数调用和 HTTP 调用两种方式。调用方可以先使用 /api/v1/gmss/normalize 检查请求规范化结果，再使用 /api/v1/gmss/run 执行计算。服务返回的 summary 字段覆盖节点规模、边界规模、微网规模、分区状态和算法配置，适合被调度辅助决策软件或前端页面直接展示。"
]),
("12 与公开文献和工具体系的对应关系", [
"主配协同潮流方向已有较多研究。Tang 等提出的异步分布式全局潮流方法强调在 TSO/DSO 通信延迟存在时，通过部分同步机制避免主网等待最慢配网，从而缩短同等精度要求下的协调分析时间[1]。本项目当前实现以工程稳定性为优先，默认采用同步边界交换和硬分区求解；但进程池并行、边界电压/功率交换和微网独立求解的设计，与异步或半异步扩展具有兼容性。",
"综合主配网潮流方面，2021 年 Electric Power Systems Research 上关于 integrated T&D fast-decoupled power flow 的研究指出，分布式电源、电动汽车和智能电网技术增强了主配相互影响，需要统一或协调的潮流分析方法评价不同电压等级之间的相互关系[2]。本项目没有选择将所有层级完全统一为一个巨型方程组，而是采用边界解耦和分区并行，适合工程服务化和多主体数据边界约束。",
"优化与调控方面，Kargarian 等综述了分布式/去中心化 DC OPF 的多类分解协调算法，包括 ATC、OCD、ADMM、APP、consensus+innovations 和 proximal message passing[3]。虽然本项目主要解决确定性潮流而非最优潮流，但 ADMM 配置、分区变量一致性和边界功率交换为后续扩展到主配协同 OPF、无功优化和柔性资源调控提供了方法基础。",
"工具链方面，MATPOWER 提供开放的稳态潮流和最优潮流研究工具[9]，ACTIVSg 合成电网提供可公开使用的大规模测试系统[10]，pandapower 提供 Python 化、表格化和可扩展的电力系统分析工具[4][5]，OpenDSS 则面向配电系统和 DER 接入仿真[6][7]。本项目将这些工具和数据思想组合到一个面向主配微分区解耦的服务模块中，技术路线具有较好的可解释性。"
]),
("13 工程应用场景", [
"第一类场景是主配协同调控前的状态校核。省调或地调在形成调峰、调频、负荷计划跟踪、断面潮流控制和配网设备潮流控制策略前，需要确认当前或预测断面下主配边界电压、线路负载、变压器负载和配网电压是否满足约束。本模块可以根据断面和时刻快速生成基础潮流状态，为后续优化模型提供初值和安全约束判断。",
"第二类场景是配网重过载和低电压治理。通过对配网分区并行潮流求解，可以定位电压薄弱区域、重载线路、重载变压器和受微网注入影响明显的 PCC。若叠加多个时刻的 96 点负荷曲线，还可形成日内电压/负载时序曲线，为无功补偿、分布式资源调节和网络重构提供基础计算结果。",
"第三类场景是微网接入影响评估。微网对配网的影响不仅取决于额定容量，还与 PCC 电压、内部负荷、新能源出力和储能策略有关。本模块通过 PCC 电压驱动微网独立求解，再将等值 P/Q 回写配网，可以在较低数据耦合成本下评估多个微网并发接入对配网潮流和主配边界的影响。",
"第四类场景是多工况批量仿真。分区并行能力使模块适合处理多断面、多时刻、多微网数量和多分区参数组合的计算任务。对于调度辅助决策软件而言，这种批量计算能力可以支撑方案对比、灵敏度分析和异常工况筛选。"
]),
("14 技术先进性与创新点", [
"项目的第一项技术特点是自动主配边界识别。相比人工维护接口表，基于拓扑图、电压等级和变压器跨层连接的识别方法能够随模型变化自动更新边界，降低多断面工程数据维护成本。",
"第二项技术特点是面向电网语义的 METIS 分区。项目不是直接使用普通图划分结果，而是结合节点负荷、边阻抗、变压器容量、潮流强度和连通性修复，使分区结果更符合潮流计算需要。",
"第三项技术特点是主配微三层边界解耦。主网、配网和微网通过边界电压与等值功率交换信息，既保持物理含义清晰，又避免整网拼接导致的数据权限、计算规模和软件接口复杂度问题。",
"第四项技术特点是服务化 API 封装。通过紧凑请求和摘要化响应，外部系统不需要理解内部数据缓存、模型转换和分区细节，即可调用计算能力并追踪结果文件。",
"第五项技术特点是兼容学术算例与工程示范数据。项目既支持 ACTIVSg/MATPOWER/IEEE13 等公开测试系统，也支持河北示范区 BPA DAT 和 10 kV 馈线 XLS 曲线，便于同时开展算法研究、课题汇报和工程联调。"
]),
("15 当前不足与风险分析", [
"第一，微网模型仍以等值验证模型为主。由于示范区尚未提供真实微网内部台账，当前微网结果主要用于验证接口和并行流程，不能直接代表真实微网运行特性。后续需要接入真实 PCC、设备容量、控制策略、负荷曲线和量测数据。",
"第二，电压等级阈值识别主配边界在多数示范区场景下有效，但在复杂地区电网中可能需要结合调度归属、站线变关系和资产台账进行补充。建议在自动识别基础上增加人工校核和规则配置能力。",
"第三，硬分区方法稳定且高效，但分区间强耦合场景下可能需要更多边界一致性迭代或软边界模型。若后续用于高精度校核，应建立整网潮流基准对比，量化分区误差随分区数量、边界位置和负荷水平变化的规律。",
"第四，当前 HTTP 服务基于标准库实现，适合本地联调和轻量部署；若进入生产环境，需要补充鉴权、任务队列、超时控制、并发限流、日志追踪、异常码规范、结果归档和容器化部署。",
"第五，部分数据转换依赖 LibreOffice、pandapower、dss-python、numba、networkx 等环境组件。建议形成 requirements.txt、环境检查脚本和一键自检样例，降低跨机器部署风险。"
]),
("16 后续优化建议", [
"算法方面，建议建立“整网潮流-分区潮流-ADMM/软边界潮流”的分层对比测试体系。对于河北示范区 13 点和 20 点断面，可固定相同负荷和设备状态，比较母线电压最大误差、边界 P/Q 误差、线路负载误差和计算耗时；对于 ACTIVSg25k+500D 大算例，可重点比较并行加速比、分区规模均衡性和边界数量对收敛的影响。",
"工程方面，建议将核心结果写入标准化 artifacts，包括分区 YAML、边界 CSV、母线电压 CSV、线路负载 CSV、变压器负载 CSV、微网 PCC 注入 CSV 和 summary JSON。这样外部系统既可通过 API 获取摘要，也可批量读取明细文件进行可视化和二次分析。",
"接口方面，建议为 /api/v1/gmss/run 增加异步任务模式。对于大规模多微网或多时刻计算，POST 请求可立即返回 task_id，后台执行完成后通过 /api/v1/gmss/result/{task_id} 查询结果，避免 HTTP 长连接超时。",
"数据方面，建议对河北示范区馈线曲线匹配建立质量报告，包括匹配行数、未匹配母线、重复匹配、负荷总量变化、异常负荷值和时刻覆盖情况。当前既有说明提到馈线负荷记录匹配数量，可进一步固化为每次运行的自动质检输出。",
"文档方面，建议在本报告基础上拆分形成三份交付材料：技术报告用于专家评审，API 说明用于联调，部署运维手册用于工程安装。三者共享同一套术语、字段和算法流程，减少后续沟通成本。"
]),
("17 结论", [
"当前项目已经具备较完整的主配微分区解耦与并行潮流计算原型能力。代码层面实现了从数据转换、主配边界识别、配网智能分区、微网等值接入、并行潮流求解到 API 输出的闭环；数据层面同时支持河北示范区 13 点/20 点工程断面和 ACTIVSg25k+500D 大规模合成验证；文档层面已有功能概述、API 输入输出说明和科技报告初稿作为基础。",
"从技术路线看，项目采用边界电压与等值功率交换的分层耦合模型，使用 pandapower 和 OpenDSS 作为底层潮流工具，使用 METIS 图划分思想实现配网区域自动划分，并通过进程池实现配网/微网并行求解。该路线与主配协同潮流、分布式计算和工程服务化需求相匹配，具有较好的可解释性和可扩展性。",
"后续工作的重点不应是重写框架，而是围绕真实微网数据接入、分区误差量化、异步任务服务、结果文件标准化、部署环境固化和生产级接口治理进行增强。完成这些工作后，本模块可以更稳地作为主配微协同调控辅助决策软件中的基础潮流计算服务。"
])
]

refs = [
"[1] K. Tang, S. Dong, Y. Liu, L. Wang, Y. Song. Asynchronous distributed global power flow method for transmission-distribution coordinated analysis considering communication conditions. Electric Power Systems Research, 2020, 182:106256. DOI: 10.1016/j.epsr.2020.106256.",
"[2] Fast-decoupled power flow method for integrated analysis of transmission and distribution systems. Electric Power Systems Research, 2021, 196:107215. DOI: 10.1016/j.epsr.2021.107215.",
"[3] A. Kargarian, J. Mohammadi, J. Guo, S. Chakrabarti, M. Barati, G. Hug, S. Kar, R. Baldick. Toward Distributed/Decentralized DC Optimal Power Flow Implementation in Future Electric Power Systems. IEEE Transactions on Smart Grid, 2018, 9(4):2574-2594. DOI: 10.1109/TSG.2016.2614904.",
"[4] L. Thurner, A. Scheidler, F. Schafer, J. H. Menke, J. Dollichon, F. Meier, S. Meinecke, M. Braun. pandapower - An Open-Source Python Tool for Convenient Modeling, Analysis, and Optimization of Electric Power Systems. IEEE Transactions on Power Systems, 2018, 33(6):6510-6521. DOI: 10.1109/TPWRS.2018.2829021.",
"[5] pandapower Documentation. Balanced AC Power Flow and project documentation. https://pandapower.readthedocs.io/ ; https://www.pandapower.org/about/.",
"[6] EPRI. OpenDSS Repository and Introduction to OpenDSS. https://opendss.epri.com/OpenDSSRepository.html ; https://opendss.epri.com/IntroductiontoOpenDSS.html.",
"[7] EPRI. OpenDSS and State Estimation. https://opendss.epri.com/OpenDSSandStateEstimation.html.",
"[8] G. Karypis, V. Kumar. Multilevel k-way Partitioning Scheme for Irregular Graphs. Journal of Parallel and Distributed Computing, 1998, 48:96-129.",
"[9] R. D. Zimmerman, C. E. Murillo-Sanchez, R. J. Thomas. MATPOWER: Steady-State Operations, Planning, and Analysis Tools for Power Systems Research and Education. IEEE Transactions on Power Systems, 2011, 26(1):12-19. DOI: 10.1109/TPWRS.2010.2051168.",
"[10] Texas A&M University Electric Grid Test Case Repository, ACTIVSg synthetic grid cases. https://electricgrids.engr.tamu.edu/ ; https://electricgrids.engr.tamu.edu/electric-grid-test-cases/activsg2000/.",
"[11] 项目本地资料：《主配微分区解耦与并行潮流API输入输出说明-河北示范区》《主配微国重项目课题科技报告模板-年中会-北京邮电大学填写版》《主配微电网分区解耦潮流计算软件功能概述》.",
]

doc = Document()
style_doc(doc)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.first_line_indent = Pt(0)
r = title.add_run("主配微分区解耦与并行潮流计算技术报告")
r.bold = True
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(22)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.first_line_indent = Pt(0)
rr = sub.add_run("面向河北示范区主配网数据与工程化 API 服务")
rr.font.name = "宋体"
rr._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
rr.font.size = Pt(14)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.first_line_indent = Pt(0)
meta.add_run("编制日期：2026年7月20日\\n项目目录：/home/y/Desktop/power_flow")
doc.add_page_break()

doc.add_heading("目录", level=1)
for i, (h, _) in enumerate(sections):
    p(doc, h)
doc.add_page_break()

for heading, paras in sections:
    doc.add_heading(heading, level=1 if heading[0].isdigit() or heading == "摘要" else 2)
    for text in paras:
        p(doc, text)
    if heading == "3 数据基础与算例规模":
        table(doc, ["断面", "母线", "线路", "变压器", "负荷", "发电/外部电网", "潮流结果"],
              [["13点", "2121", "2036", "73", "698", "49 gen + 121 sgen + 1 ext_grid", "收敛，4次迭代，Vm 0.9423-1.0354 pu"],
               ["20点", "2122", "2038", "73", "664", "48 gen + 6 sgen + 1 ext_grid", "收敛，4次迭代，Vm 0.9378-1.0345 pu"],
               ["ACTIVSg25k+500D", "主网25k级+500个配网接口", "合成主网及IEEE13配网", "500个主配变压器", "合成/等值", "按接口等值", "用于大规模扩展验证"]])
    if heading == "9 软件架构与 API 设计":
        table(doc, ["接口", "方法", "作用"],
              [["/api/v1/health", "GET", "健康检查"],
               ["/api/v1/contract", "GET", "返回服务契约、示例请求和响应字段"],
               ["/api/v1/gmss/normalize", "POST", "规范化请求但不执行计算"],
               ["/api/v1/gmss/run", "POST", "执行主配微分区解耦与并行潮流计算"]])
    if heading != "17 结论":
        doc.add_page_break()

doc.add_heading("参考文献", level=1)
for ref in refs:
    para = doc.add_paragraph(ref)
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.space_after = Pt(4)

doc.add_heading("附录A：建议的报告配图清单", level=1)
for item in [
    "主配微三层耦合架构图：主网T、配网D、微网M及边界变量 V/P/Q。",
    "计算流程图：数据读取、边界识别、分区划分、并行潮流、微网迭代、结果输出。",
    "软件架构图：核心算法层、服务适配层、HTTP接口层、数据缓存层。",
    "河北示范区断面数据统计图：13点与20点母线、线路、负荷和电压范围对比。",
    "METIS配网分区示意图：节点权重、边权重、割边和连通性修复。"
]:
    para = doc.add_paragraph(item)
    para.paragraph_format.first_line_indent = Pt(0)

doc.save(OUT)
print(OUT.resolve())
