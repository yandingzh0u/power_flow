from __future__ import annotations

import json
import shutil
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import PercentFormatter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "课题1-推进会-任务2-0302(1)_第7-8页修改版.pptx"
BACKUP = ROOT / "课题1-推进会-任务2-0302(1)_补算例前备份.pptx"
OUTPUT = ROOT / "课题1-推进会-任务2-0302(1)_补充算例版.pptx"
ASSET_DIR = ROOT / "_case_study_assets"

SUMMARY_JSON = ROOT / "data/demo_area/converted/demo_area_20_summary.json"
PARTITION_JSON = ROOT / "data/demo_area/converted/demo_area_20_distribution_partitions.json"
NETWORK_JSON = ROOT / "data/demo_area/converted/demo_area_20_pandapower.json"

TEAL = RGBColor(0x00, 0x78, 0x73)
DEEP_TEAL = RGBColor(0x00, 0x66, 0x66)
BLUE = RGBColor(0x1F, 0x66, 0xB1)
GREEN = RGBColor(0x4F, 0x91, 0x35)
ORANGE = RGBColor(0xE9, 0x70, 0x0A)
PURPLE = RGBColor(0x6D, 0x49, 0xA8)
RED = RGBColor(0xD9, 0x2D, 0x2D)
BLACK = RGBColor(0x16, 0x20, 0x24)
GRAY = RGBColor(0x5F, 0x69, 0x70)
MID_GRAY = RGBColor(0x8A, 0x96, 0x9C)
LIGHT_LINE = RGBColor(0xC9, 0xD5, 0xD9)
PALE = RGBColor(0xF4, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_east_asian_font(run, typeface="微软雅黑"):
    run.font.name = typeface
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", typeface)


def style_run(run, size, *, bold=False, color=BLACK, font="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    set_east_asian_font(run, font)


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def clear_slide(slide):
    for shape in list(slide.shapes):
        remove_shape(shape)


def set_text(
    shape,
    text,
    *,
    size,
    bold=False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margins=(0.05, 0.05, 0.03, 0.03),
    line_spacing=1.0,
    font="微软雅黑",
):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for index, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        style_run(run, size, bold=bold, color=color, font=font)


def add_text(slide, x, y, w, h, text, **kwargs):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(shape, text, **kwargs)
    return shape


def add_box(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill=WHITE,
    line=LIGHT_LINE,
    rounded=False,
    line_width=0.9,
):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if rounded
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_line(slide, x1, y1, x2, y2, color, width=1.2, dash=None):
    line = slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_common_header(slide, page_number):
    marker = add_box(slide, 0, 0, 0.16, 0.76, fill=TEAL, line=TEAL)
    marker.line.fill.background()
    add_line(slide, 0, 0.76, 12.42, 0.76, TEAL, 1.45)
    add_text(
        slide,
        0.35,
        0.13,
        5.9,
        0.46,
        "二 、任务进展情况",
        size=26,
        bold=True,
        color=RGBColor(0, 0, 0),
        font="黑体",
    )
    add_text(
        slide,
        12.75,
        7.19,
        0.36,
        0.18,
        str(page_number),
        size=9.5,
        color=RGBColor(0x88, 0x88, 0x88),
        align=PP_ALIGN.RIGHT,
        margins=(0, 0, 0, 0),
    )


def add_slide_title(slide, sequence, title):
    tag = add_box(slide, 0.30, 0.94, 1.68, 0.42, fill=TEAL, line=TEAL)
    set_text(
        tag,
        f"算例验证 {sequence}",
        size=18,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.02, 0.02, 0.01, 0.01),
    )
    add_text(
        slide,
        2.17,
        0.94,
        10.70,
        0.42,
        title,
        size=21.5,
        bold=True,
        color=DEEP_TEAL,
        margins=(0, 0, 0, 0),
    )


def add_summary_band(slide, label, text):
    label_box = add_box(slide, 0.35, 1.48, 1.03, 0.64, fill=TEAL, line=TEAL)
    set_text(
        label_box,
        label,
        size=15.5,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.03, 0.03, 0.02, 0.02),
    )
    body = add_box(slide, 1.43, 1.48, 11.50, 0.64, fill=PALE, line=LIGHT_LINE)
    set_text(
        body,
        text,
        size=15.2,
        color=BLACK,
        margins=(0.16, 0.12, 0.04, 0.03),
        line_spacing=1.0,
    )


def add_panel_header(slide, x, y, w, text, color):
    header = add_box(slide, x, y, w, 0.38, fill=color, line=color)
    set_text(
        header,
        text,
        size=15.5,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        margins=(0.02, 0.02, 0.01, 0.01),
    )


def add_node(slide, x, y, color, radius=0.12, *, ring=False):
    node = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x - radius),
        Inches(y - radius),
        Inches(radius * 2),
        Inches(radius * 2),
    )
    if ring:
        node.fill.solid()
        node.fill.fore_color.rgb = WHITE
    else:
        node.fill.solid()
        node.fill.fore_color.rgb = color
    node.line.color.rgb = color
    node.line.width = Pt(1.2)
    return node


def add_transformer(slide, x, y):
    add_line(slide, x, y - 0.28, x, y + 0.28, MID_GRAY, 1.0)
    top = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x - 0.12),
        Inches(y - 0.12),
        Inches(0.24),
        Inches(0.24),
    )
    top.fill.background()
    top.line.color.rgb = MID_GRAY
    top.line.width = Pt(1.0)
    bottom = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x - 0.12),
        Inches(y + 0.03),
        Inches(0.24),
        Inches(0.24),
    )
    bottom.fill.background()
    bottom.line.color.rgb = MID_GRAY
    bottom.line.width = Pt(1.0)


def add_topology_schematic(slide):
    x, y, w, h = 0.35, 2.33, 7.22, 3.70
    add_box(slide, x, y, w, h, fill=WHITE, line=LIGHT_LINE, line_width=0.9)
    add_panel_header(slide, x, y, w, "实际电网模型与分区结果（拓扑脱敏示意）", BLUE)

    # Main-grid backbone.
    main_y = 3.10
    main_xs = [0.92, 1.80, 2.72, 3.70, 4.68, 5.62, 6.72]
    for left, right in zip(main_xs, main_xs[1:]):
        add_line(slide, left, main_y, right, main_y, BLUE, 1.55)
    add_line(slide, 2.72, main_y, 2.28, 2.78, BLUE, 1.15)
    add_line(slide, 4.68, main_y, 5.12, 2.78, BLUE, 1.15)
    for nx in main_xs:
        add_node(slide, nx, main_y, BLUE, 0.075)
    add_node(slide, 2.28, 2.78, BLUE, 0.075)
    add_node(slide, 5.12, 2.78, BLUE, 0.075)
    add_text(
        slide,
        0.57,
        2.69,
        0.78,
        0.24,
        "主网区域",
        size=12,
        bold=True,
        color=BLUE,
        margins=(0, 0, 0, 0),
    )

    transformer_xs = [1.80, 3.70, 5.62]
    region_colors = [GREEN, TEAL, ORANGE]
    region_fills = [
        RGBColor(0xF1, 0xF7, 0xEE),
        RGBColor(0xEE, 0xF7, 0xF6),
        RGBColor(0xFE, 0xF3, 0xEA),
    ]
    region_labels = ["配网子区 A", "配网子区 B", "配网子区 C"]
    region_xs = [0.66, 2.85, 5.02]
    for tx in transformer_xs:
        add_line(slide, tx, main_y, tx, 3.43, MID_GRAY, 1.0)
        add_transformer(slide, tx, 3.55)
        add_node(slide, tx, 3.86, RED, 0.095, ring=True)
    add_text(
        slide,
        6.02,
        3.42,
        1.25,
        0.26,
        "主配边界",
        size=11.5,
        bold=True,
        color=RED,
        margins=(0, 0, 0, 0),
    )

    for index, (rx, color, fill, label) in enumerate(
        zip(region_xs, region_colors, region_fills, region_labels)
    ):
        region = add_box(
            slide,
            rx,
            4.00,
            1.86,
            1.40,
            fill=fill,
            line=color,
            rounded=True,
            line_width=1.05,
        )
        set_text(
            region,
            "",
            size=1,
            margins=(0, 0, 0, 0),
        )
        add_text(
            slide,
            rx + 0.12,
            4.08,
            1.62,
            0.25,
            label,
            size=11.8,
            bold=True,
            color=color,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )
        source_x = transformer_xs[index]
        center_x = rx + 0.93
        add_line(slide, source_x, 3.95, center_x, 4.45, color, 1.2)
        branch_points = [
            (center_x, 4.45),
            (rx + 0.38, 4.72),
            (rx + 0.93, 4.80),
            (rx + 1.48, 4.70),
            (rx + 0.60, 5.12),
            (rx + 1.28, 5.10),
        ]
        edges = [(0, 1), (0, 2), (0, 3), (1, 4), (2, 4), (2, 5), (3, 5)]
        for a, b in edges:
            add_line(
                slide,
                branch_points[a][0],
                branch_points[a][1],
                branch_points[b][0],
                branch_points[b][1],
                color,
                0.9,
            )
        for px, py in branch_points:
            add_node(slide, px, py, color, 0.055)

    # PCC and local-resource layer.
    for px, color in ((1.59, GREEN), (3.78, TEAL), (5.95, ORANGE)):
        add_line(slide, px, 5.10, px, 5.54, color, 1.0)
        add_node(slide, px, 5.55, RED, 0.08, ring=True)
    add_text(
        slide,
        0.55,
        5.48,
        0.62,
        0.24,
        "PCC",
        size=11,
        bold=True,
        color=RED,
        margins=(0, 0, 0, 0),
    )
    add_text(
        slide,
        1.82,
        5.45,
        3.82,
        0.24,
        "微网等值注入 / 分布式资源",
        size=11.3,
        color=ORANGE,
        align=PP_ALIGN.CENTER,
        margins=(0, 0, 0, 0),
    )
    add_text(
        slide,
        0.54,
        5.73,
        6.82,
        0.20,
        "注：拓扑关系与站线名称已作脱敏处理，图中颜色表示不同局部潮流子区。",
        size=9.8,
        color=GRAY,
        align=PP_ALIGN.CENTER,
        margins=(0, 0, 0, 0),
    )


def add_labeled_bullet(slide, x, y, w, label, body, color, size=12.8):
    dot = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x),
        Inches(y + 0.08),
        Inches(0.09),
        Inches(0.09),
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = color
    dot.line.fill.background()
    box = slide.shapes.add_textbox(
        Inches(x + 0.17), Inches(y), Inches(w - 0.17), Inches(0.31)
    )
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    r1 = p.add_run()
    r1.text = label
    style_run(r1, size, bold=True, color=color)
    r2 = p.add_run()
    r2.text = body
    style_run(r2, size, color=BLACK)


def add_case_settings(slide):
    x, y, w = 7.80, 2.33, 5.13
    add_box(slide, x, y, w, 1.62, fill=WHITE, line=LIGHT_LINE)
    add_panel_header(slide, x, y, w, "算例设置与验证基准", TEAL)
    bullets = [
        ("算例对象：", "石家庄实际电网模型与典型运行断面"),
        ("验证基准：", "同一断面下的整体交流潮流解"),
        ("模型校核：", "设备端点、电压等级、运行状态与拓扑连通"),
        ("边界对象：", "主配变压器、馈线接口与微网 PCC"),
    ]
    for i, (label, body) in enumerate(bullets):
        add_labeled_bullet(slide, x + 0.20, y + 0.48 + i * 0.27, w - 0.37, label, body, TEAL)

    y2 = 4.10
    add_box(slide, x, y2, w, 1.93, fill=WHITE, line=LIGHT_LINE)
    add_panel_header(slide, x, y2, w, "分区实施流程与输出", GREEN)
    steps = [
        ("1", "统一设备语义、参数基准与运行时标"),
        ("2", "自动辨识跨层接口并锁定边界锚点"),
        ("3", "构建阻抗—潮流—源荷注入加权图"),
        ("4", "执行多级图划分与接口约束优化"),
        ("5", "连通修复后生成拓扑完整的局部模型"),
    ]
    for i, (number, text) in enumerate(steps):
        sy = y2 + 0.48 + i * 0.28
        circle = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Inches(x + 0.20),
            Inches(sy + 0.01),
            Inches(0.22),
            Inches(0.22),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = GREEN
        circle.line.fill.background()
        set_text(
            circle,
            number,
            size=10.5,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )
        add_text(
            slide,
            x + 0.52,
            sy - 0.01,
            w - 0.70,
            0.25,
            text,
            size=12.4,
            color=BLACK,
            margins=(0, 0, 0, 0),
        )


def add_status_cards(slide):
    cards = [
        ("结构校核", "通过", BLUE, RGBColor(0xF0, 0xF5, 0xFB)),
        ("边界辨识", "自动完成", TEAL, RGBColor(0xEE, 0xF7, 0xF6)),
        ("分区执行", "并行开启", GREEN, RGBColor(0xF2, 0xF7, 0xEF)),
        ("局部潮流", "全部收敛", ORANGE, RGBColor(0xFE, 0xF3, 0xEA)),
    ]
    for i, (label, value, color, fill) in enumerate(cards):
        x = 0.35 + i * 3.16
        card = add_box(
            slide,
            x,
            6.19,
            2.96,
            0.73,
            fill=fill,
            line=color,
            rounded=True,
            line_width=1.0,
        )
        set_text(card, "", size=1, margins=(0, 0, 0, 0))
        add_text(
            slide,
            x + 0.18,
            6.31,
            1.25,
            0.22,
            label,
            size=12.2,
            bold=True,
            color=GRAY,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )
        add_text(
            slide,
            x + 1.40,
            6.24,
            1.35,
            0.35,
            value,
            size=17,
            bold=True,
            color=color,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )
        add_line(slide, x + 1.36, 6.30, x + 1.36, 6.80, LIGHT_LINE, 0.8)


def extract_frame(network, key):
    encoded = network["_object"][key]["_object"]
    return json.loads(encoded)


def format_sci_unicode(value):
    if value == 0:
        return "0"
    exponent = int(np.floor(np.log10(abs(value))))
    coefficient = value / (10**exponent)
    superscript = str(exponent).translate(
        str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")
    )
    return f"{coefficient:.2f}×10{superscript}"


def make_charts():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    network = json.loads(NETWORK_JSON.read_text(encoding="utf-8"))
    res_bus = extract_frame(network, "res_bus")
    vm_index = res_bus["columns"].index("vm_pu")
    vm_values = np.asarray([row[vm_index] for row in res_bus["data"]], dtype=float)

    partition_data = json.loads(PARTITION_JSON.read_text(encoding="utf-8"))
    vm_errors = np.sort(
        np.asarray(
            [part["max_vm_delta_vs_full_pu"] for part in partition_data["partitions"]],
            dtype=float,
        )
    )
    va_errors = np.sort(
        np.asarray(
            [part["max_va_delta_vs_full_deg"] for part in partition_data["partitions"]],
            dtype=float,
        )
    )

    plt.rcParams.update(
        {
            "font.sans-serif": ["Noto Sans CJK SC", "Microsoft YaHei", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "mathtext.fontset": "dejavusans",
        }
    )

    voltage_path = ASSET_DIR / "shijiazhuang_voltage_distribution.png"
    fig, ax = plt.subplots(figsize=(7.0, 3.25), dpi=220)
    weights = np.full(vm_values.shape, 100.0 / vm_values.size)
    ax.axvspan(0.95, 1.05, color="#4F9135", alpha=0.10)
    ax.hist(
        vm_values,
        bins=26,
        weights=weights,
        color="#1F66B1",
        edgecolor="white",
        linewidth=0.5,
        alpha=0.90,
    )
    vm_min = float(vm_values.min())
    vm_max = float(vm_values.max())
    ax.axvline(vm_min, color="#E9700A", linestyle="--", linewidth=1.5)
    ax.axvline(vm_max, color="#007873", linestyle="--", linewidth=1.5)
    ymax = ax.get_ylim()[1]
    ax.text(0.951, ymax * 0.93, "参考范围 0.95–1.05 p.u.", color="#4F9135", fontsize=9)
    ax.text(vm_min + 0.001, ymax * 0.79, f"最小值 {vm_min:.4f}", color="#E9700A", fontsize=9)
    ax.text(vm_max - 0.001, ymax * 0.80, f"最大值 {vm_max:.4f}", color="#007873", fontsize=9, ha="right")
    ax.set_xlim(0.932, 1.040)
    ax.set_xlabel("电压幅值 / p.u.", fontsize=10)
    ax.set_ylabel("占比", fontsize=10)
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=100, decimals=0))
    ax.grid(axis="y", color="#D9E1E4", linewidth=0.7, alpha=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(labelsize=9)
    fig.tight_layout(pad=0.8)
    fig.savefig(voltage_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)

    error_path = ASSET_DIR / "shijiazhuang_partition_errors.png"
    fig, axes = plt.subplots(1, 2, figsize=(7.1, 3.25), dpi=220)
    rank_vm = np.linspace(0, 100, vm_errors.size)
    rank_va = np.linspace(0, 100, va_errors.size)
    configs = [
        (axes[0], rank_vm, vm_errors, "#1F66B1", "电压幅值误差", "p.u.", float(vm_errors.max())),
        (axes[1], rank_va, va_errors, "#E9700A", "相角误差", "°", float(va_errors.max())),
    ]
    for ax, rank, values, color, title, unit, maximum in configs:
        ax.plot(rank, values, color=color, linewidth=1.9)
        ax.fill_between(rank, values, np.maximum(values.min() * 0.75, 1e-16), color=color, alpha=0.10)
        ax.scatter([100], [maximum], color=color, s=18, zorder=3)
        ax.set_yscale("log")
        ax.set_title(title, fontsize=10.5, fontweight="bold", color=color, pad=3)
        ax.set_xlabel("误差排序分位 / %", fontsize=9)
        ax.set_ylabel(unit, fontsize=9)
        ax.grid(True, which="both", color="#D9E1E4", linewidth=0.55, alpha=0.75)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=8)
        exponent = int(np.floor(np.log10(abs(maximum))))
        coefficient = maximum / (10**exponent)
        ax.text(
            0.97,
            0.10,
            f"最大值\n${coefficient:.2f}\\times10^{{{exponent}}}$ {unit}",
            transform=ax.transAxes,
            ha="right",
            va="bottom",
            fontsize=8.7,
            color=color,
            bbox={"facecolor": "white", "edgecolor": color, "alpha": 0.88, "pad": 2.5},
        )
    fig.suptitle(
        r"$e_V=\max_i|V_i^{part}-V_i^{full}|$    "
        r"$e_\theta=\max_i|\theta_i^{part}-\theta_i^{full}|$",
        fontsize=10.5,
        y=0.995,
        color="#162024",
    )
    fig.tight_layout(rect=(0, 0, 1, 0.92), pad=0.8, w_pad=1.2)
    fig.savefig(error_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return voltage_path, error_path


def add_picture_contain(slide, path, x, y, w, h):
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio > box_ratio:
        target_w = w
        target_h = w / ratio
        target_x = x
        target_y = y + (h - target_h) / 2
    else:
        target_h = h
        target_w = h * ratio
        target_x = x + (w - target_w) / 2
        target_y = y
    return slide.shapes.add_picture(
        str(path),
        Inches(target_x),
        Inches(target_y),
        width=Inches(target_w),
        height=Inches(target_h),
    )


def add_chart_panel(slide, x, y, w, h, title, color, chart_path):
    add_box(slide, x, y, w, h, fill=WHITE, line=LIGHT_LINE, line_width=0.9)
    add_panel_header(slide, x, y, w, title, color)
    add_picture_contain(slide, chart_path, x + 0.10, y + 0.44, w - 0.20, h - 0.51)


def add_metric_cards(slide, summary, partitions):
    pf = summary["pandapower_powerflow"]
    cards = [
        ("整体潮流", f"{pf['iterations']} 次收敛", BLUE),
        ("运行电压范围", f"{pf['vm_pu_min']:.4f}–{pf['vm_pu_max']:.4f} p.u.", TEAL),
        (
            "最大幅值偏差",
            f"{format_sci_unicode(partitions['max_vm_delta_vs_full_pu'])} p.u.",
            GREEN,
        ),
        (
            "最大相角偏差",
            f"{format_sci_unicode(partitions['max_va_delta_vs_full_deg'])}°",
            ORANGE,
        ),
        ("失败分区", str(partitions["failed_partitions"]), PURPLE),
    ]
    for i, (label, value, color) in enumerate(cards):
        x = 0.35 + i * 2.52
        card = add_box(
            slide,
            x,
            5.82,
            2.36,
            0.76,
            fill=WHITE,
            line=color,
            rounded=True,
            line_width=1.05,
        )
        set_text(card, "", size=1, margins=(0, 0, 0, 0))
        add_text(
            slide,
            x + 0.08,
            5.90,
            2.20,
            0.20,
            label,
            size=10.8,
            bold=True,
            color=GRAY,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )
        value_size = 15.8 if len(value) < 15 else 13.0
        add_text(
            slide,
            x + 0.06,
            6.14,
            2.24,
            0.27,
            value,
            size=value_size,
            bold=True,
            color=color,
            align=PP_ALIGN.CENTER,
            margins=(0, 0, 0, 0),
        )


def build_case_slide_1(slide):
    clear_slide(slide)
    add_common_header(slide, 18)
    add_slide_title(slide, 1, "石家庄实际电网模型与分区结果")
    add_summary_band(
        slide,
        "验证对象",
        "面向石家庄实际电网模型与典型运行断面，完成模型标准化、边界自动辨识、加权图构建、多级分区与局部潮流模型生成。",
    )
    add_topology_schematic(slide)
    add_case_settings(slide)
    add_status_cards(slide)


def build_case_slide_2(slide, voltage_chart, error_chart, summary, partitions):
    clear_slide(slide)
    add_common_header(slide, 19)
    add_slide_title(slide, 2, "石家庄算例解耦精度与并行求解结果")
    add_summary_band(
        slide,
        "验证方法",
        "将各分区局部潮流结果按原始母线映射回填，并与同一运行断面的整体潮流逐点对齐，检查电压分布、幅值偏差、相角偏差及失败分区。",
    )
    add_chart_panel(slide, 0.35, 2.31, 6.10, 3.32, "整体潮流电压分布", BLUE, voltage_chart)
    add_chart_panel(slide, 6.70, 2.31, 6.23, 3.32, "分区解耦误差（逐分区排序）", ORANGE, error_chart)
    add_metric_cards(slide, summary, partitions)
    conclusion = add_box(
        slide,
        0.35,
        6.70,
        12.58,
        0.43,
        fill=RGBColor(0xF0, 0xF7, 0xF6),
        line=TEAL,
        rounded=True,
        line_width=0.95,
    )
    set_text(
        conclusion,
        "结果表明：分区解耦未引入显著潮流偏差，各局部模型均稳定收敛，具备层内任务并行执行与后续主配微协调迭代的计算基础。",
        size=13.6,
        bold=True,
        color=DEEP_TEAL,
        align=PP_ALIGN.CENTER,
        margins=(0.06, 0.06, 0.02, 0.02),
    )


def insert_after(prs, slide, index):
    sld_id_list = prs.slides._sldIdLst
    sld_id = sld_id_list[-1]
    sld_id_list.remove(sld_id)
    sld_id_list.insert(index + 1, sld_id)
    return slide


def restore_untouched_slide_parts(source, output):
    # Existing slides retain their original XML; slide 18 and the newly-created
    # slide are the only content slides intentionally changed.
    members = set()
    for number in list(range(1, 18)) + list(range(19, 22)):
        members.add(f"ppt/slides/slide{number}.xml")
        members.add(f"ppt/slides/_rels/slide{number}.xml.rels")
    replacement = output.with_suffix(".restored.pptx")
    with ZipFile(source) as source_zip, ZipFile(output) as output_zip, ZipFile(
        replacement, "w", ZIP_DEFLATED
    ) as replacement_zip:
        source_names = set(source_zip.namelist())
        for item in output_zip.infolist():
            data = output_zip.read(item.filename)
            if item.filename in members and item.filename in source_names:
                data = source_zip.read(item.filename)
            replacement_zip.writestr(item, data)
    replacement.replace(output)


def validate_no_scale_disclosure(output):
    forbidden = ("总节点", "节点总数", "母线总数", "2122", "2121", "25k", "ACTIVSg")
    prs = Presentation(output)
    if len(prs.slides) != 22:
        raise RuntimeError(f"Expected 22 slides after insertion, found {len(prs.slides)}")
    for page in (18, 19):
        text = "\n".join(
            shape.text for shape in prs.slides[page - 1].shapes if hasattr(shape, "text")
        )
        found = [token for token in forbidden if token in text]
        if found:
            raise RuntimeError(f"Slide {page} contains forbidden scale text: {found}")


def main():
    for path in (SOURCE, SUMMARY_JSON, PARTITION_JSON, NETWORK_JSON):
        if not path.exists():
            raise FileNotFoundError(path)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    summary = json.loads(SUMMARY_JSON.read_text(encoding="utf-8"))
    partitions = json.loads(PARTITION_JSON.read_text(encoding="utf-8"))
    voltage_chart, error_chart = make_charts()

    prs = Presentation(SOURCE)
    if len(prs.slides) != 21:
        raise RuntimeError(f"Expected 21 source slides, found {len(prs.slides)}")

    first_case_slide = prs.slides[17]
    new_slide = prs.slides.add_slide(first_case_slide.slide_layout)
    insert_after(prs, new_slide, 17)
    build_case_slide_1(first_case_slide)
    build_case_slide_2(new_slide, voltage_chart, error_chart, summary, partitions)

    prs.save(OUTPUT)
    restore_untouched_slide_parts(SOURCE, OUTPUT)
    validate_no_scale_disclosure(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
