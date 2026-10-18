from __future__ import annotations

import shutil
from pathlib import Path
from tempfile import TemporaryDirectory

from matplotlib import rcParams
from matplotlib.font_manager import FontProperties
from matplotlib.mathtext import math_to_image
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
DEMO = ROOT / "demo.pptx"
BACKUP = ROOT / "demo_原始备份.pptx"
IMG_DIR = ROOT / "img_ppt"

TEAL = RGBColor(0x00, 0x7A, 0x72)
TEAL_DARK = RGBColor(0x00, 0x5D, 0x58)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x18, 0x1B, 0x1D)
GRAY = RGBColor(0x63, 0x69, 0x6D)
LIGHT_GRAY = RGBColor(0xD7, 0xDD, 0xDF)
BLUE = RGBColor(0x1E, 0x5A, 0xA8)
GREEN = RGBColor(0x4E, 0x8C, 0x35)
ORANGE = RGBColor(0xDC, 0x6B, 0x12)


PAGES = [
    {
        "title": "研究进展2.1.1：主配微异构模型统一建模方法",
        "image": "ChatGPT Image 2026年8月4日 14_00_34.png",
        "statement": (
            "面向主网、配网和微网模型来源异构、参数体系不一致及运行断面难对齐问题，建立设备语义映射与统一拓扑建模方法。\n"
            "通过设备标识、参数基准和时间状态标准化，形成供边界识别、网络分区及潮流求解共享的计算对象。"
        ),
    },
    {
        "title": "研究进展2.1.2：主配微边界与微网PCC自动辨识方法",
        "image": "ChatGPT Image 2026年8月4日 14_08_26.png",
        "statement": (
            "依据母线电压等级形成主网与配网集合，通过变压器端点追踪识别跨层连接设备，并结合微网接入关系定位PCC。\n"
            "自动生成主侧母线、边界设备、配侧母线及微网并网点之间的多层接口映射。"
        ),
    },
    {
        "title": "研究进展2.2.1：接口锚点约束的主网区域生长方法",
        "image": "ChatGPT Image 2026年8月4日 14_12_29.png",
        "statement": (
            "以主配接口母线作为优先锚点，采用拓扑最远距离原则选择离散种子，并从多个种子同步开展区域生长。\n"
            "在保证子区连通和规模均衡的同时复制跨区邻接节点，为后续切口状态协调构造共享边界。"
        ),
    },
    {
        "title": "研究进展2.2.2：电气耦合约束的配网智能分区方法",
        "image": "ChatGPT Image 2026年8月4日 14_17_27.png",
        "statement": (
            "将配电网转换为带节点权和支路权的图模型，以负荷与注入表征计算权重，以阻抗和潮流表征电气耦合。\n"
            "通过多级缩约、k路划分、逐级投影及连通修复，形成区内强耦合、区间弱耦合且规模受控的子区。"
        ),
    },
    {
        "title": "研究进展2.3.1：边界节点复制与接口等值解耦方法",
        "image": "ChatGPT Image 2026年8月4日 14_22_51.png",
        "statement": (
            "通过局部网络截取和切口节点复制，将全网模型转换为多个拓扑完整的局部潮流子问题。\n"
            "上游子区提供边界电压，下游子区返回等值功率，微网通过PCC电压与等值注入接入配网局部模型。"
        ),
    },
    {
        "title": "研究进展2.3.2：边界电压与功率协调迭代方法",
        "image": "ChatGPT Image 2026年8月4日 14_35_34.png",
        "statement": (
            "按照主网边界电压下发、配网交换功率计算和微网等值注入更新的顺序交替求解。\n"
            "采用松弛更新抑制边界变量振荡，并以主配交换功率残差和边界电压残差联合判断协调收敛。"
        ),
    },
    {
        "title": "研究进展2.4.1：主配微分层并行潮流求解方法",
        "image": "ChatGPT Image 2026年8月4日 14_40_21.png",
        "statement": (
            "将主网子区、配网实例和微网实例分别封装为独立潮流任务，按任务数量与可用计算资源配置工作进程。\n"
            "同层任务并行求解，层间按照边界变量依赖关系顺序协调，并在每层末端同步汇总计算结果。"
        ),
    },
]


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]  # noqa: SLF001
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]  # noqa: SLF001


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def textbox(slide, x, y, w, h, text, size, color=BLACK, *, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font="Microsoft YaHei"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        p.line_spacing = 1.04
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, *, radius=False, line=None, width=0.8):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return shape


def add_original_header(slide, title: str, page: int) -> None:
    rect(slide, 0.0, 0.0, 0.16, 0.73, TEAL)
    textbox(slide, 0.22, 0.05, 6.2, 0.50, "二、任务进展情况", 20.5, BLACK, bold=True)
    rect(slide, 0.0, 0.72, 13.333, 0.025, TEAL)

    rect(slide, 0.28, 0.86, 11.05, 0.47, TEAL)
    textbox(slide, 0.56, 0.86, 10.60, 0.47, title, 18.5, WHITE, bold=True)
    rect(slide, 11.53, 0.86, 1.30, 0.47, TEAL_DARK, radius=True)
    textbox(slide, 11.53, 0.86, 1.30, 0.47, f"{page}/7", 13.5, WHITE, bold=True, align=PP_ALIGN.CENTER)

    rect(slide, 0.43, 1.48, 0.82, 0.92, TEAL, line=WHITE)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(1.19), Inches(1.74), Inches(0.24), Inches(0.38)
    )
    tri.fill.solid()
    tri.fill.fore_color.rgb = TEAL
    tri.line.fill.background()
    textbox(slide, 0.43, 1.48, 0.82, 0.92, "技术\n方法", 16.0, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.55, 1.42, 11.15, 0.92, PAGES[page - 1]["statement"], 13.5, BLACK)

    textbox(slide, 12.55, 7.14, 0.45, 0.20, str(page), 8.5, GRAY, align=PP_ALIGN.RIGHT)


def crop(source: Path, target: Path, box: tuple[int, int, int, int]) -> Path:
    with Image.open(source) as image:
        image.convert("RGB").crop(box).save(target, "PNG", optimize=True)
    return target


def add_picture_contain(slide, path: Path, x, y, w, h, *, border=True):
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if image_ratio >= box_ratio:
        pw = w
        ph = w / image_ratio
        px = x
        py = y + (h - ph) / 2
    else:
        ph = h
        pw = h * image_ratio
        px = x + (w - pw) / 2
        py = y
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    if border:
        pic.line.color.rgb = LIGHT_GRAY
        pic.line.width = Pt(0.6)
    return pic


def render_formula(target: Path, expression: str, *, size=28) -> Path:
    rcParams["mathtext.fontset"] = "stix"
    math_to_image(
        f"${expression}$",
        str(target),
        prop=FontProperties(size=size),
        dpi=300,
        format="png",
        color="#181B1D",
    )
    with Image.open(target).convert("RGBA") as rendered:
        alpha = ImageOps.invert(rendered.convert("L"))
        transparent = Image.new("RGBA", rendered.size, (0x18, 0x1B, 0x1D, 0))
        transparent.putalpha(alpha)
        transparent.save(target, "PNG", optimize=True)
    return target


def arrow(slide, x, y, w, h, color=TEAL, direction="right"):
    kinds = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left": MSO_SHAPE.LEFT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
        "up": MSO_SHAPE.UP_ARROW,
    }
    kind = kinds[direction]
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=BLACK, *, width=1.2, dashed=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def circle(slide, cx, cy, diameter, fill, *, line_color=None, width=0.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(cx - diameter / 2),
        Inches(cy - diameter / 2),
        Inches(diameter),
        Inches(diameter),
    )
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(width)
    return shape


def page_one(slide, source: Path, tmp: Path) -> None:
    blue = RGBColor(0x0B, 0x55, 0xA6)
    cyan = RGBColor(0x00, 0x7C, 0x82)
    cyan_dark = RGBColor(0x00, 0x5C, 0x68)
    panel_fill = RGBColor(0xF8, 0xFB, 0xFD)
    card_blue = RGBColor(0xEE, 0xF5, 0xFB)
    card_cyan = RGBColor(0xED, 0xF8, 0xF7)
    table_head = RGBColor(0xD9, 0xF0, 0xF2)

    panels = [
        (0.70, 2.15, "1. 多源模型输入", blue),
        (3.27, 2.15, "2. 参数与标识标准化", cyan),
        (5.83, 2.15, "3. 设备语义映射", cyan_dark),
        (8.38, 4.18, "4. 统一计算模型", blue),
    ]
    for x, w, title, color in panels:
        rect(slide, x, 2.46, w, 0.36, color, radius=True)
        textbox(slide, x, 2.46, w, 0.36, title, 13.0, WHITE, bold=True, align=PP_ALIGN.CENTER)
        rect(slide, x, 2.88, w, 3.52, panel_fill, radius=True, line=color, width=1.1)

    # 1. Multi-source inputs: all labels are editable PowerPoint text.
    input_cards = [
        ("主网模型", "母线、线路、发电机、主变"),
        ("配网模型", "馈线、配变、负荷、\n分布式电源"),
        ("微网模型", "PCC、光伏、储能、\n可控电源"),
        ("运行断面", "负荷曲线、开关状态、\n设备投运状态"),
    ]
    for idx, (title, detail) in enumerate(input_cards):
        y = 3.04 + idx * 0.80
        rect(slide, 0.84, y, 1.87, 0.66, card_blue, radius=True, line=blue, width=0.8)
        textbox(slide, 0.91, y + 0.04, 1.73, 0.22, title, 10.8, BLACK, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, 0.91, y + 0.26, 1.73, 0.34, detail, 9.2, BLACK, align=PP_ALIGN.CENTER)

    arrow(slide, 2.91, 4.18, 0.30, 0.34, blue, "right")

    # 2. Standardization chain.
    standard_steps = [
        (3.14, "设备名称和\n索引规范化"),
        (4.25, "电压、功率和\n阻抗基准统一"),
        (5.36, "时间断面与\n设备状态对齐"),
    ]
    for y, label in standard_steps:
        rect(slide, 3.48, y, 1.73, 0.72, card_cyan, radius=True, line=cyan, width=0.9)
        textbox(slide, 3.57, y + 0.03, 1.55, 0.66, label, 11.0, BLACK, bold=True, align=PP_ALIGN.CENTER)
    arrow(slide, 4.16, 3.91, 0.28, 0.26, cyan, "down")
    arrow(slide, 4.16, 5.02, 0.28, 0.26, cyan, "down")

    arrow(slide, 5.48, 4.18, 0.29, 0.34, blue, "right")

    # 3. Semantic mapping table redrawn as PowerPoint shapes.
    textbox(slide, 6.04, 3.01, 1.74, 0.34, "映射关系表", 12.7, BLACK, bold=True, align=PP_ALIGN.CENTER)
    mapping_rows = [
        ("原始设备", "统一对象"),
        ("母线节点", "Bus"),
        ("线路/电缆", "Branch"),
        ("主变/配变", "Transformer"),
        ("负荷/电源", "Injection"),
        ("并网点", "Boundary/PCC"),
    ]
    row_h = 0.43
    for row, (left, right) in enumerate(mapping_rows):
        y = 3.49 + row * row_h
        fill = table_head if row == 0 else WHITE
        rect(slide, 5.96, y, 1.03, row_h, fill, line=cyan_dark, width=0.65)
        rect(slide, 6.99, y, 0.86, row_h, fill, line=cyan_dark, width=0.65)
        textbox(slide, 5.98, y, 0.99, row_h, left, 9.2, BLACK, bold=(row == 0), align=PP_ALIGN.CENTER)
        textbox(slide, 7.01, y, 0.82, row_h, right, 8.8, BLACK, bold=(row == 0), align=PP_ALIGN.CENTER)

    arrow(slide, 8.03, 4.18, 0.29, 0.34, blue, "right")

    # 4. Retain only the topology illustration from the old artwork.
    topology = crop(source, tmp / "p1_topology_only.png", (1195, 90, 1645, 605))
    add_picture_contain(slide, topology, 8.58, 2.98, 3.78, 2.63, border=False)

    formula_graph = render_formula(tmp / "p1_formula_graph.png", r"G=(V,E)", size=30)
    formula_state = render_formula(
        tmp / "p1_formula_state.png",
        r"\mathbf{x}=[V,\theta,P,Q]^{\mathsf{T}}",
        size=30,
    )
    add_picture_contain(slide, formula_graph, 8.63, 5.68, 1.32, 0.28, border=False)
    textbox(slide, 9.96, 5.65, 2.28, 0.34, "：统一拓扑对象", 10.1, BLACK)
    add_picture_contain(slide, formula_state, 8.63, 6.01, 1.86, 0.28, border=False)
    textbox(slide, 10.49, 5.98, 1.78, 0.34, "：统一运行状态变量", 9.5, BLACK)

    textbox(
        slide, 1.05, 6.55, 11.25, 0.34,
        "模型校核：设备端点完整性    |    电压等级一致性    |    拓扑连通性与孤岛检测",
        11.2, BLUE, bold=True, align=PP_ALIGN.CENTER,
    )


def page_two(slide, source: Path, tmp: Path) -> None:
    blue = RGBColor(0x0B, 0x55, 0xA6)
    blue_light = RGBColor(0x78, 0xA9, 0xE2)
    green = RGBColor(0x3E, 0x8B, 0x2C)
    green_light = RGBColor(0x93, 0xC8, 0x5B)
    orange = RGBColor(0xE9, 0x67, 0x00)
    red = RGBColor(0xDF, 0x20, 0x22)
    gray_line = RGBColor(0xD5, 0xDB, 0xDE)

    panel_specs = [
        (0.72, 3.40, "1", "电压等级分层"),
        (4.47, 3.34, "2", "跨域变压器识别"),
        (8.13, 4.25, "3", "PCC映射"),
    ]
    for x, w, number, title in panel_specs:
        rect(slide, x, 2.43, w, 4.43, WHITE, line=gray_line, width=0.8)
        circle(slide, x + 0.34, 2.72, 0.42, blue)
        textbox(slide, x + 0.13, 2.52, 0.42, 0.40, number, 15.2, WHITE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.72, 2.51, w - 0.92, 0.42, title, 16.0, BLACK, bold=True, align=PP_ALIGN.CENTER)

    # Step 1: retain the original voltage-level topology artwork.
    voltage_topology = crop(source, tmp / "p2_voltage_topology.png", (45, 105, 505, 605))
    add_picture_contain(slide, voltage_topology, 0.94, 3.03, 2.95, 2.47, border=False)
    circle(slide, 1.05, 5.66, 0.14, blue_light, line_color=blue, width=0.7)
    textbox(slide, 1.18, 5.53, 2.48, 0.26, "蓝色节点：高电压集合", 9.4, BLACK)
    circle(slide, 1.05, 5.92, 0.14, green_light, line_color=green, width=0.7)
    textbox(slide, 1.18, 5.79, 2.48, 0.26, "绿色节点：低电压集合", 9.4, BLACK)

    voltage_formula = render_formula(
        tmp / "p2_voltage_formula.png",
        r"V_T=\{i\mid V_{n,i}\geq V_{\mathrm{th}}\},\quad V_D=\{i\mid V_{n,i}<V_{\mathrm{th}}\}",
        size=25,
    )
    rect(slide, 0.90, 6.17, 3.04, 0.50, WHITE, radius=True, line=blue, width=0.8)
    add_picture_contain(slide, voltage_formula, 1.00, 6.25, 2.84, 0.31, border=False)

    # Step 2: retain the original main/distribution bus and transformer topology.
    transformer_topology = crop(source, tmp / "p2_transformer_topology.png", (580, 135, 1005, 470))
    add_picture_contain(slide, transformer_topology, 4.70, 3.13, 2.88, 1.66, border=False)

    boundary_points = [
        (4.93, "高压侧母线属于主网集合"),
        (5.28, "低压侧母线属于配网集合"),
        (5.63, "跨集合连接确定主配边界设备"),
    ]
    for y, label in boundary_points:
        tri = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(4.84), Inches(y), Inches(0.16), Inches(0.17))
        tri.fill.solid()
        tri.fill.fore_color.rgb = red
        tri.line.fill.background()
        textbox(slide, 5.05, y - 0.03, 2.50, 0.23, label, 9.6, red, bold=True)

    boundary_formula = render_formula(
        tmp / "p2_boundary_formula.png",
        r"B_{TD}=\{(i,j)\in E_{\mathrm{tr}}\mid i\in V_T,\ j\in V_D\}",
        size=25,
    )
    rect(slide, 4.72, 6.17, 2.85, 0.50, WHITE, radius=True, line=blue, width=0.8)
    add_picture_contain(slide, boundary_formula, 4.85, 6.25, 2.59, 0.31, border=False)

    # Step 3: keep the full original PCC topology and replace only its formulas.
    pcc_topology = crop(source, tmp / "p2_pcc_topology.png", (1075, 175, 1655, 620))
    add_picture_contain(slide, pcc_topology, 8.35, 3.03, 3.84, 2.56, border=False)
    rect(slide, 9.50, 3.87, 0.82, 0.31, WHITE)
    rect(slide, 9.16, 4.83, 1.06, 0.33, WHITE)
    voltage_pcc = render_formula(tmp / "p2_vpcc.png", r"V_{\mathrm{PCC}}", size=26)
    power_mg = render_formula(tmp / "p2_pqmg.png", r"P_{\mathrm{MG}},\ Q_{\mathrm{MG}}", size=24)
    add_picture_contain(slide, voltage_pcc, 9.52, 3.90, 0.76, 0.25, border=False)
    add_picture_contain(slide, power_mg, 9.16, 4.87, 1.02, 0.25, border=False)
    rect(slide, 10.23, 4.90, 0.49, 0.27, WHITE)
    textbox(slide, 10.24, 4.90, 0.47, 0.27, "PCC", 10.0, red, bold=True, align=PP_ALIGN.CENTER)

    textbox(slide, 8.47, 5.67, 3.57, 0.28, "识别输出", 12.3, BLACK, bold=True, align=PP_ALIGN.CENTER)
    output_boxes = [
        (8.35, blue, RGBColor(0xE8, 0xF1, 0xFC), "主配接口", "主侧母线、边界变压器、配侧母线"),
        (10.36, green, RGBColor(0xEC, 0xF7, 0xE8), "配微接口", "配网分区、PCC母线、微网标识"),
    ]
    for x, color, header_fill, title, detail in output_boxes:
        rect(slide, x, 6.00, 1.84, 0.66, WHITE, radius=True, line=color, width=0.8)
        rect(slide, x, 6.00, 1.84, 0.29, header_fill, line=color, width=0.8)
        textbox(slide, x + 0.04, 6.01, 1.76, 0.27, title, 10.0, color, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.07, 6.30, 1.70, 0.32, detail, 8.2, BLACK, align=PP_ALIGN.CENTER)


def page_three(slide, source: Path, tmp: Path) -> None:
    blue = RGBColor(0x0B, 0x5B, 0xB5)
    blue_light = RGBColor(0x8E, 0xBA, 0xEA)
    green = RGBColor(0x21, 0x94, 0x55)
    green_light = RGBColor(0xA7, 0xD4, 0x92)
    orange = RGBColor(0xF0, 0x7A, 0x19)
    orange_light = RGBColor(0xF5, 0xB0, 0x73)
    red = RGBColor(0xE3, 0x22, 0x27)
    edge_gray = RGBColor(0x70, 0x76, 0x79)
    node_gray = RGBColor(0xEC, 0xEF, 0xF1)
    panel_line = RGBColor(0xB9, 0xC1, 0xC5)

    panel_xs = [0.43, 2.74, 5.05, 7.36]
    panel_w = 2.08
    stage_titles = ["原始网络", "种子选择", "区域生长", "边界复制"]
    stage_captions = [
        ("识别主配接口锚点", red),
        ("接口优先 + 最远距离选种", blue),
        ("多源同步扩展", TEAL_DARK),
        ("共享边界节点复制", blue),
    ]
    for idx, (x, title) in enumerate(zip(panel_xs, stage_titles)):
        rect(slide, x, 2.47, panel_w, 3.90, RGBColor(0xFC, 0xFD, 0xFE), radius=True, line=panel_line, width=0.8)
        textbox(slide, x + 0.10, 2.54, panel_w - 0.20, 0.33, title, 14.0, BLACK, bold=True, align=PP_ALIGN.CENTER)
        caption, color = stage_captions[idx]
        textbox(slide, x + 0.10, 5.98, panel_w - 0.20, 0.27, caption, 9.3, color, bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            arrow(slide, x + panel_w + 0.04, 4.21, 0.19, 0.27, RGBColor(0xA4, 0xA9, 0xAC), "right")

    relative_nodes = {
        "a": (0.55, 0.15), "b": (1.55, 0.15),
        "c": (0.25, 0.72), "d": (0.95, 0.72), "e": (1.76, 0.72),
        "f": (0.12, 1.35), "g": (0.86, 1.32), "h": (1.46, 1.18), "i": (1.90, 1.35),
        "j": (0.30, 1.98), "k": (1.12, 1.88), "l": (1.73, 2.02),
        "m": (0.69, 2.55), "n": (1.45, 2.55),
    }
    edges = [
        ("a", "b"), ("a", "c"), ("a", "d"), ("b", "d"),
        ("c", "f"), ("c", "g"), ("d", "g"),
        ("e", "h"), ("e", "i"), ("h", "i"),
        ("f", "g"), ("f", "j"), ("g", "h"), ("g", "k"), ("g", "l"),
        ("j", "m"), ("k", "l"), ("k", "n"), ("l", "n"), ("m", "n"),
    ]
    clusters = {
        "a": "blue", "b": "blue", "c": "blue", "d": "blue", "f": "blue", "g": "blue",
        "e": "green", "h": "green", "i": "green", "l": "green",
        "j": "orange", "k": "orange", "m": "orange", "n": "orange",
    }
    cluster_fill = {"blue": blue_light, "green": green_light, "orange": orange_light}
    cluster_line = {"blue": blue, "green": green, "orange": orange}
    seed_nodes = {"g": blue, "e": green, "m": orange}

    def draw_seed_ring(cx, cy, color, *, waves=False):
        circle(slide, cx, cy, 0.24, None, line_color=color, width=1.2)
        if waves:
            for diameter in (0.36, 0.50):
                wave = circle(slide, cx, cy, diameter, None, line_color=color, width=0.55)
                wave.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def growth_arrow(cx, cy, angle, color):
        shape = arrow(slide, cx - 0.13, cy - 0.035, 0.26, 0.07, color, "right")
        shape.rotation = angle
        return shape

    def draw_stage_network(panel_x, stage):
        points = {
            name: (panel_x + 0.06 + dx, 3.05 + 0.78 * dy)
            for name, (dx, dy) in relative_nodes.items()
        }
        for left, right in edges:
            color = edge_gray
            width = 0.65
            if stage in (3, 4) and clusters[left] == clusters[right]:
                color = cluster_line[clusters[left]]
                width = 0.85
            line(slide, *points[left], *points[right], color, width=width)

        for name, (cx, cy) in points.items():
            if stage in (1, 2):
                fill = node_gray
                outline = edge_gray
            else:
                fill = cluster_fill[clusters[name]]
                outline = cluster_line[clusters[name]]

            if stage == 1 and name == "g":
                fill = red
                outline = red
            if stage == 2 and name in seed_nodes:
                fill = seed_nodes[name]
                outline = seed_nodes[name]

            if stage == 4 and name in ("h", "k"):
                left_color = blue
                right_color = green if name == "h" else orange
                for start, end in [(-0.095, -0.055), (-0.020, 0.020), (0.055, 0.095)]:
                    line(slide, cx + start, cy, cx + end, cy, blue, width=1.15)
                circle(slide, cx - 0.145, cy, 0.10, left_color, line_color=left_color, width=0.8)
                circle(slide, cx + 0.145, cy, 0.10, right_color, line_color=right_color, width=0.8)
            else:
                circle(slide, cx, cy, 0.135, fill, line_color=outline, width=0.65)

        if stage == 1:
            for name in ("g", "j", "m"):
                cx, cy = points[name]
                draw_seed_ring(cx, cy, red)
            circle(slide, panel_x + 0.20, 5.56, 0.16, None, line_color=red, width=0.9)
            circle(slide, panel_x + 0.20, 5.56, 0.08, node_gray, line_color=edge_gray, width=0.5)
            textbox(slide, panel_x + 0.33, 5.45, 1.58, 0.22, "红圈：主配接口节点", 7.2, red, bold=True)
        elif stage == 2:
            for name, color in seed_nodes.items():
                cx, cy = points[name]
                draw_seed_ring(cx, cy, color)
            cx, cy = points["j"]
            draw_seed_ring(cx, cy, red)
            legend_rows = [
                (5.17, red, node_gray, "红圈：接口锚点候选"),
                (5.36, blue, blue, "蓝圈：首个种子，度数最大"),
                (5.55, green, green, "绿圈：后续种子，拓扑最远"),
                (5.74, orange, orange, "橙圈：后续种子，拓扑最远"),
            ]
            for y, color, fill, label in legend_rows:
                circle(slide, panel_x + 0.17, y + 0.075, 0.14, None, line_color=color, width=0.75)
                circle(slide, panel_x + 0.17, y + 0.075, 0.075, fill, line_color=color, width=0.45)
                textbox(slide, panel_x + 0.29, y, 1.67, 0.17, label, 6.2, color, bold=True)
        elif stage == 3:
            for name, color in seed_nodes.items():
                cx, cy = points[name]
                draw_seed_ring(cx, cy, color, waves=True)
            gx, gy = points["g"]
            ex, ey = points["e"]
            mx, my = points["m"]
            for dx, dy, angle in [(-0.34, 0.00, 180), (0.34, 0.00, 0), (0.00, -0.34, 270), (0.00, 0.34, 90)]:
                growth_arrow(gx + dx, gy + dy, angle, blue)
            growth_arrow(ex - 0.30, ey + 0.08, 180, green)
            growth_arrow(ex + 0.02, ey + 0.32, 90, green)
            growth_arrow(mx + 0.30, my - 0.05, 0, orange)
            growth_arrow(mx + 0.02, my - 0.30, 270, orange)
            circle(slide, panel_x + 0.16, 5.47, 0.09, blue, line_color=blue, width=0.4)
            circle(slide, panel_x + 0.26, 5.47, 0.09, green, line_color=green, width=0.4)
            circle(slide, panel_x + 0.36, 5.47, 0.09, orange, line_color=orange, width=0.4)
            textbox(slide, panel_x + 0.45, 5.36, 1.49, 0.20, "同色节点：种子归属子区", 6.6, TEAL_DARK, bold=True)
            textbox(slide, panel_x + 0.14, 5.62, 1.80, 0.20, "箭头/虚线环：同步扩展方向与距离层", 6.3, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
        else:
            for name in ("g", "j", "m"):
                cx, cy = points[name]
                draw_seed_ring(cx, cy, red)
            circle(slide, panel_x + 0.16, 5.48, 0.08, blue, line_color=blue, width=0.4)
            circle(slide, panel_x + 0.23, 5.48, 0.08, green, line_color=green, width=0.4)
            textbox(slide, panel_x + 0.33, 5.37, 1.60, 0.20, "双色节点：跨区共享边界", 6.5, blue, bold=True)
            line(slide, panel_x + 0.14, 5.72, panel_x + 0.30, 5.72, blue, width=0.65, dashed=True)
            textbox(slide, panel_x + 0.34, 5.61, 1.58, 0.20, "虚线：边界节点复制关系", 6.5, blue, bold=True)

    for stage, x in enumerate(panel_xs, start=1):
        draw_stage_network(x, stage)

    formula_one = render_formula(
        tmp / "p3_seed_formula_one.png",
        r"s_1=\underset{b\in B_{TD}}{\mathrm{arg\,max}}\,\deg(b)",
        size=23,
    )
    formula_next = render_formula(
        tmp / "p3_seed_formula_next.png",
        r"s_{r+1}=\underset{b\in V\setminus S}{\mathrm{arg\,max}}\;\underset{s\in S}{\min}\;d_G(b,s)",
        size=21,
    )
    rect(slide, 9.62, 2.50, 3.06, 1.30, RGBColor(0xF8, 0xFB, 0xFE), radius=True, line=blue, width=0.8)
    textbox(slide, 9.82, 2.57, 2.66, 0.30, "种子选择准则", 13.2, blue, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, formula_one, 9.83, 2.94, 2.64, 0.30, border=False)
    add_picture_contain(slide, formula_next, 9.78, 3.34, 2.74, 0.30, border=False)

    rect(slide, 9.62, 4.00, 3.06, 1.06, RGBColor(0xF3, 0xF7, 0xFB), radius=True, line=RGBColor(0x8F, 0xB1, 0xD2), width=0.8)
    textbox(slide, 9.82, 4.09, 2.66, 0.25, "选种策略", 12.0, blue, bold=True, align=PP_ALIGN.CENTER)
    textbox(
        slide, 9.84, 4.39, 2.62, 0.54,
        "接口锚点优先\n按拓扑最远距离递推选择后续种子",
        10.0, BLACK, align=PP_ALIGN.CENTER,
    )

    rect(slide, 9.62, 5.22, 3.06, 1.12, RGBColor(0xF3, 0xF9, 0xF7), radius=True, line=RGBColor(0x8F, 0xBD, 0xA7), width=0.8)
    textbox(slide, 9.82, 5.31, 2.66, 0.25, "区域生长与边界处理", 12.0, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    textbox(
        slide, 9.82, 5.61, 2.66, 0.58,
        "多源同步扩展形成连通子区\n跨区邻接节点在相邻子区共同保留",
        9.8, BLACK, align=PP_ALIGN.CENTER,
    )

    rect(slide, 9.68, 6.50, 2.94, 0.36, RGBColor(0xEC, 0xF3, 0xFA), radius=True, line=RGBColor(0x93, 0xB2, 0xD0), width=0.7)
    textbox(slide, 9.76, 6.51, 2.78, 0.33, "输出：连通、均衡且边界可协调的主网子区", 9.2, blue, bold=True, align=PP_ALIGN.CENTER)


def page_four(slide, source: Path, tmp: Path) -> None:
    blue = RGBColor(0x0B, 0x5B, 0xB5)
    blue_fill = RGBColor(0xD9, 0xE9, 0xFA)
    green = RGBColor(0x36, 0x8F, 0x45)
    green_fill = RGBColor(0xE2, 0xF1, 0xDF)
    orange = RGBColor(0xE5, 0x78, 0x16)
    orange_fill = RGBColor(0xFA, 0xEA, 0xDA)
    purple = RGBColor(0x7A, 0x4B, 0xB1)
    purple_fill = RGBColor(0xEC, 0xE4, 0xF5)
    red = RGBColor(0xE2, 0x22, 0x27)
    edge = RGBColor(0x66, 0x6C, 0x70)
    panel_line = RGBColor(0xB9, 0xC1, 0xC5)
    panel_fill = RGBColor(0xFC, 0xFD, 0xFE)

    def stage_panel(x, y, w, h, number, title):
        rect(slide, x, y, w, h, panel_fill, radius=True, line=panel_line, width=0.75)
        textbox(slide, x + 0.10, y + 0.04, w - 0.20, 0.28, f"{number}  {title}", 12.2, BLUE, bold=True, align=PP_ALIGN.CENTER)

    def labeled_node(cx, cy, diameter, label, fill=WHITE, outline=edge, *, font_size=7.2, bold=False):
        circle(slide, cx, cy, diameter, fill, line_color=outline, width=0.75)
        textbox(
            slide, cx - diameter / 2, cy - diameter / 2,
            diameter, diameter, label, font_size, BLACK, bold=bold, align=PP_ALIGN.CENTER,
        )

    def dashed_group(x, y, w, h, color, *, radius=True):
        shape = rect(slide, x, y, w, h, WHITE, radius=radius, line=color, width=0.70)
        shape.fill.background()
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return shape

    stage_panel(0.42, 2.42, 3.02, 2.16, "①", "加权图构建")
    stage_panel(3.56, 2.42, 3.02, 2.16, "②", "多级缩约")
    stage_panel(6.70, 2.42, 3.05, 2.16, "③", "粗图 k 路初始划分")
    stage_panel(9.89, 2.42, 2.98, 2.16, "", "权重定义")
    stage_panel(0.42, 4.72, 6.20, 2.16, "④", "逐级投影与局部优化")
    stage_panel(6.74, 4.72, 6.13, 2.16, "⑤", "连通修复与过小子区合并")

    # 1. Weighted graph construction.
    fine_pos = {
        "1": (0.70, 2.95), "2": (1.55, 2.91), "3": (2.56, 2.99),
        "4": (0.58, 3.48), "5": (1.25, 3.46), "6": (2.02, 3.48), "7": (2.88, 3.48),
        "8": (0.79, 3.97), "9": (1.45, 3.98), "10": (2.20, 3.97), "11": (2.85, 3.97),
    }
    weighted_edges = [
        ("1", "2", 2.0), ("2", "3", 1.8), ("1", "4", 2.1), ("1", "5", 0.8),
        ("2", "5", 2.5), ("2", "6", 1.2), ("3", "6", 2.0), ("3", "7", 0.8),
        ("4", "5", 2.2), ("4", "8", 1.6), ("5", "6", 2.0), ("5", "9", 1.0),
        ("6", "7", 1.0), ("6", "10", 1.5), ("7", "11", 1.1),
        ("8", "9", 1.8), ("9", "10", 1.0), ("10", "11", 1.8),
    ]
    for left, right, width in weighted_edges:
        line(slide, *fine_pos[left], *fine_pos[right], edge, width=width)
    for left, right in [("2", "6"), ("5", "9"), ("6", "7")]:
        line(slide, *fine_pos[left], *fine_pos[right], red, width=1.05, dashed=True)
    node_sizes = {"1": 0.23, "2": 0.34, "3": 0.23, "4": 0.34, "5": 0.28, "6": 0.35, "7": 0.25,
                  "8": 0.21, "9": 0.22, "10": 0.31, "11": 0.23}
    for label, (cx, cy) in fine_pos.items():
        labeled_node(cx, cy, node_sizes[label], label, RGBColor(0xF5, 0xF6, 0xF7), edge, font_size=6.8)
    circle(slide, 0.66, 4.26, 0.12, WHITE, line_color=edge, width=0.65)
    textbox(slide, 0.78, 4.16, 1.03, 0.18, "节点越大：计算权重越大", 6.3, BLACK)
    line(slide, 1.92, 4.25, 2.17, 4.25, edge, width=2.1)
    textbox(slide, 2.20, 4.16, 0.95, 0.18, "粗边：电气耦合强", 6.3, BLACK)
    line(slide, 0.66, 4.45, 0.94, 4.45, red, width=1.0, dashed=True)
    textbox(slide, 0.99, 4.36, 1.60, 0.18, "红虚线：候选切割边", 6.3, red, bold=True)

    # 2. Multilevel coarsening.
    textbox(slide, 3.76, 2.76, 2.62, 0.25, "仅合并相邻强耦合节点，未匹配节点直接继承", 7.7, BLACK, align=PP_ALIGN.CENTER)
    original_nodes = {
        "1": (3.71, 3.25), "2": (3.98, 3.22), "3": (4.28, 3.27),
        "4": (3.67, 3.52), "5": (3.89, 3.51), "6": (4.13, 3.52), "7": (4.38, 3.52),
        "8": (3.73, 3.79), "9": (3.94, 3.80), "10": (4.18, 3.79), "11": (4.38, 3.79),
    }
    for x, y, w, h in [
        (3.62, 3.13, 0.46, 0.20), (3.58, 3.43, 0.41, 0.18), (3.64, 3.70, 0.40, 0.19),
    ]:
        dashed_group(x, y, w, h, RGBColor(0x6F, 0x99, 0xC2), radius=False)
    for left, right, _ in weighted_edges:
        line(slide, *original_nodes[left], *original_nodes[right], edge, width=0.38)
    for label, pos in original_nodes.items():
        labeled_node(*pos, 0.095, label, RGBColor(0xF6, 0xF7, 0xF8), edge, font_size=3.2)
    arrow(slide, 4.45, 3.43, 0.18, 0.15, edge, "right")

    # Level 1: a={1,2}, b={3}, c={4,5}, d={6}, e={7},
    # f={8,9}, g={10}, h={11}. Every aggregate is connected in G0.
    level_one_nodes = {
        "a": (4.72, 3.24), "b": (4.98, 3.23),
        "c": (4.68, 3.53), "d": (4.96, 3.50), "e": (5.25, 3.52),
        "f": (4.72, 3.79), "g": (5.02, 3.78), "h": (5.28, 3.78),
    }
    dashed_group(4.62, 3.14, 0.46, 0.19, RGBColor(0x6F, 0x99, 0xC2), radius=False)
    dashed_group(4.94, 3.68, 0.44, 0.20, RGBColor(0x6F, 0x99, 0xC2), radius=False)
    level_one_edges = [
        ("a", "b"), ("a", "c"), ("a", "d"), ("b", "d"), ("b", "e"),
        ("c", "d"), ("c", "f"), ("d", "e"), ("d", "g"), ("e", "h"),
        ("f", "g"), ("g", "h"),
    ]
    for a, b in level_one_edges:
        line(slide, *level_one_nodes[a], *level_one_nodes[b], edge, width=0.52)
    for label, pos in level_one_nodes.items():
        labeled_node(*pos, 0.115, label, RGBColor(0xF6, 0xF7, 0xF8), edge, font_size=4.0)
    arrow(slide, 5.38, 3.43, 0.18, 0.15, edge, "right")

    # Level 2: A={a,b}, B={c}, C={d}, D={e}, E={f}, F={g,h}.
    # The six-node graph is the actual input to the k-way initializer.
    c2_nodes = {
        "A": (5.72, 3.23), "B": (5.68, 3.52), "C": (5.98, 3.48),
        "D": (6.30, 3.48), "E": (5.72, 3.79), "F": (6.18, 3.78),
    }
    coarse_edges = [
        ("A", "B"), ("A", "C"), ("A", "D"), ("B", "C"), ("B", "E"),
        ("C", "D"), ("C", "F"), ("D", "F"), ("E", "F"),
    ]
    for a, b in coarse_edges:
        line(slide, *c2_nodes[a], *c2_nodes[b], edge, width=0.75)
    for label, pos in c2_nodes.items():
        labeled_node(*pos, 0.135, label, RGBColor(0xF6, 0xF7, 0xF8), edge, font_size=4.5)
    textbox(slide, 3.61, 4.02, 0.88, 0.18, "原图  n=11", 5.9, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 4.59, 4.02, 0.84, 0.18, "一级缩约  n=8", 5.9, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 5.56, 4.02, 0.91, 0.18, "二级缩约  n=6", 5.9, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(
        slide, 3.63, 4.20, 2.88, 0.31,
        "一级：(1,2)→a，(4,5)→c，(8,9)→f；其余继承\n二级：(a,b)→A，(g,h)→F；其余继承",
        4.4, BLUE, bold=True, align=PP_ALIGN.CENTER,
    )

    # 3. Coarse graph partitioning and objective.
    textbox(slide, 6.92, 2.76, 2.60, 0.22, "在 6 节点粗图上求解 k=4 的初始划分", 7.7, BLACK, align=PP_ALIGN.CENTER)
    k_nodes = {
        "A": (6.98, 3.13), "B": (6.94, 3.46), "C": (7.26, 3.42),
        "D": (7.59, 3.42), "E": (6.99, 3.73), "F": (7.48, 3.73),
    }
    for a, b in coarse_edges:
        line(slide, *k_nodes[a], *k_nodes[b], edge, width=0.80)
    for label, pos in k_nodes.items():
        labeled_node(*pos, 0.155, label, RGBColor(0xF6, 0xF7, 0xF8), edge, font_size=5.0)
    arrow(slide, 7.73, 3.37, 0.25, 0.19, BLACK, "right")

    partition_cards = [
        (8.08, 2.98, blue, blue_fill, "P1", ("A",)),
        (8.82, 2.98, green, green_fill, "P2", ("B",)),
        (8.08, 3.47, orange, orange_fill, "P3", ("C", "E")),
        (8.82, 3.47, purple, purple_fill, "P4", ("D", "F")),
    ]
    for x, y, color, fill, part, node_labels in partition_cards:
        dashed_group(x, y, 0.62, 0.42, color)
        textbox(slide, x + 0.03, y + 0.01, 0.56, 0.14, part, 6.9, BLACK, bold=True, align=PP_ALIGN.CENTER)
        if len(node_labels) == 1:
            labeled_node(x + 0.31, y + 0.29, 0.145, node_labels[0], fill, color, font_size=4.8)
        else:
            left_x, right_x = x + 0.20, x + 0.43
            if part == "P4":
                line(slide, left_x, y + 0.29, right_x, y + 0.29, color, width=0.75)
            labeled_node(left_x, y + 0.29, 0.125, node_labels[0], fill, color, font_size=4.1)
            labeled_node(right_x, y + 0.29, 0.125, node_labels[1], fill, color, font_size=4.1)

    objective = render_formula(
        tmp / "p4_partition_objective.png",
        r"\min_{\pi}\;\sum_{(i,j)\in E}w_{ij}\,\mathbf{1}(\pi_i\ne\pi_j)",
        size=22,
    )
    rect(slide, 6.86, 3.96, 2.73, 0.32, WHITE, radius=True, line=blue, width=0.7)
    add_picture_contain(slide, objective, 6.98, 4.01, 2.49, 0.22, border=False)
    textbox(slide, 6.86, 4.31, 2.73, 0.20, "唯一归属  |  权重均衡  |  连通性后验检测", 6.0, BLACK, bold=True, align=PP_ALIGN.CENTER)

    # Weight definitions rendered by the formula engine.
    formula_node = render_formula(
        tmp / "p4_weight_node.png",
        r"a_i=1+\lambda_v\sqrt{P_i^2+Q_i^2}",
        size=21,
    )
    formula_line = render_formula(
        tmp / "p4_weight_line.png",
        r"w_{ij}^{\mathrm{line}}=c_0+\frac{\lambda_z}{|Z_{ij}|+\epsilon}+\lambda_s|S_{ij}|",
        size=18,
    )
    formula_transformer = render_formula(
        tmp / "p4_weight_transformer.png",
        r"w_{ij}^{\mathrm{tr}}=c_t+\lambda_c\frac{S_{n,ij}}{v_{k,ij}}+\lambda_lL_{ij}",
        size=18,
    )
    textbox(slide, 10.10, 2.77, 0.62, 0.19, "节点权", 7.2, blue, bold=True)
    add_picture_contain(slide, formula_node, 10.68, 2.74, 1.92, 0.25, border=False)
    textbox(slide, 10.10, 3.18, 0.62, 0.19, "线路权", 7.2, blue, bold=True)
    add_picture_contain(slide, formula_line, 10.66, 3.13, 2.00, 0.28, border=False)
    textbox(slide, 10.10, 3.61, 0.62, 0.19, "变压器权", 7.2, blue, bold=True)
    add_picture_contain(slide, formula_transformer, 10.66, 3.56, 2.00, 0.28, border=False)
    line(slide, 10.06, 4.02, 12.70, 4.02, blue, width=0.65, dashed=True)
    textbox(slide, 10.14, 4.10, 2.48, 0.36, "权重经归一化与离散化后\n输入多级图划分器", 7.5, BLUE, bold=True, align=PP_ALIGN.CENTER)

    # 4. Projection and local refinement.
    textbox(slide, 0.70, 5.03, 5.65, 0.21, "依据聚合映射逐层展开标签，并在每层搜索边界节点的迁移增益", 7.6, BLACK, align=PP_ALIGN.CENTER)
    coarse = {
        "A": (0.76, 5.39), "B": (0.71, 5.77), "C": (1.07, 5.73),
        "D": (1.44, 5.73), "E": (0.77, 6.12), "F": (1.34, 6.12),
    }
    for x, y, color in [
        (0.60, 5.23, blue), (0.55, 5.61, green), (0.91, 5.56, orange),
        (1.28, 5.56, purple), (0.61, 5.95, orange), (1.18, 5.95, purple),
    ]:
        dashed_group(x, y, 0.32, 0.33, color)
    for a, b in coarse_edges:
        line(slide, *coarse[a], *coarse[b], edge, width=0.62)
    coarse_style = {
        "A": (blue_fill, blue), "B": (green_fill, green),
        "C": (orange_fill, orange), "D": (purple_fill, purple),
        "E": (orange_fill, orange), "F": (purple_fill, purple),
    }
    for label, pos in coarse.items():
        fill, color = coarse_style[label]
        labeled_node(*pos, 0.145, label, fill, color, font_size=4.7)
    arrow(slide, 1.66, 5.76, 0.27, 0.19, BLACK, "right")

    middle = {
        "a": (2.24, 5.38), "b": (2.60, 5.38),
        "c": (2.13, 5.77), "d": (2.59, 5.73), "e": (3.10, 5.73),
        "f": (2.28, 6.14), "g": (2.74, 6.14), "h": (3.16, 6.14),
    }
    dashed_group(2.06, 5.21, 0.72, 0.34, blue)
    dashed_group(1.96, 5.60, 0.34, 0.35, green)
    dashed_group(2.42, 5.56, 0.34, 0.35, orange)
    dashed_group(2.93, 5.56, 0.34, 0.35, purple)
    dashed_group(2.11, 5.97, 0.34, 0.35, orange)
    dashed_group(2.57, 5.97, 0.76, 0.35, purple)
    for a, b in level_one_edges:
        line(slide, *middle[a], *middle[b], edge, width=0.57)
    middle_style = {
        "a": (blue_fill, blue), "b": (blue_fill, blue), "c": (green_fill, green),
        "d": (orange_fill, orange), "e": (purple_fill, purple), "f": (orange_fill, orange),
        "g": (purple_fill, purple), "h": (purple_fill, purple),
    }
    for label, pos in middle.items():
        fill, color = middle_style[label]
        labeled_node(*pos, 0.135, label, fill, color, font_size=4.4)
    arrow(slide, 3.44, 5.76, 0.27, 0.19, BLACK, "right")

    projected = {
        "1": (4.05, 5.43), "2": (4.58, 5.43), "3": (5.12, 5.43),
        "4": (3.93, 5.84), "5": (4.48, 5.82), "6": (5.08, 5.82), "7": (5.68, 5.82),
        "8": (4.02, 6.23), "9": (4.57, 6.23), "10": (5.26, 6.23), "11": (5.78, 6.23),
    }
    dashed_group(3.83, 5.26, 1.52, 0.35, blue)
    dashed_group(3.72, 5.66, 0.99, 0.34, green)
    dashed_group(4.80, 5.66, 0.53, 0.34, orange)
    dashed_group(5.43, 5.65, 0.50, 0.35, purple)
    dashed_group(3.80, 6.05, 1.00, 0.35, orange)
    dashed_group(5.02, 6.05, 1.02, 0.35, purple)
    projected_edges = [("1", "2"), ("2", "3"), ("1", "4"), ("1", "5"), ("2", "5"),
                       ("2", "6"), ("3", "6"), ("3", "7"), ("4", "5"), ("4", "8"),
                       ("5", "6"), ("5", "9"), ("6", "7"), ("6", "10"), ("7", "11"),
                       ("8", "9"), ("9", "10"), ("10", "11")]
    for a, b in projected_edges:
        line(slide, *projected[a], *projected[b], edge, width=0.58)
    projected_style = {
        "1": (blue_fill, blue), "2": (blue_fill, blue), "3": (blue_fill, blue),
        "4": (green_fill, green), "5": (green_fill, green), "6": (orange_fill, orange),
        "7": (purple_fill, purple), "8": (orange_fill, orange), "9": (orange_fill, orange),
        "10": (purple_fill, purple), "11": (purple_fill, purple),
    }
    for label, pos in projected.items():
        fill, color = projected_style[label]
        labeled_node(*pos, 0.15, label, fill, color, font_size=5.0)
    textbox(slide, 0.50, 6.49, 1.38, 0.19, "二级粗图初始划分", 6.1, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.91, 6.49, 1.51, 0.19, "投影至一级缩约图", 6.1, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 3.82, 6.49, 2.18, 0.19, "投影至原图并局部优化", 6.1, BLACK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 0.82, 6.70, 5.38, 0.15, "标签按聚合映射逐层展开，仅对边界节点执行增益搜索", 6.2, BLUE, bold=True, align=PP_ALIGN.CENTER)

    # 5. Connectivity repair and small-partition merging.
    textbox(slide, 6.98, 5.03, 5.55, 0.21, "优先恢复子区连通性，再将规模过小的子区并入耦合最强的相邻区域", 7.9, BLACK, align=PP_ALIGN.CENTER)
    repair = {
        "1": (7.12, 5.43), "2": (7.68, 5.43), "3": (8.24, 5.43),
        "4": (7.02, 5.84), "5": (7.58, 5.82), "6": (8.18, 5.82), "7": (8.80, 5.82),
        "8": (7.18, 6.23), "9": (7.72, 6.23), "10": (8.42, 6.23), "11": (8.95, 6.23),
    }
    dashed_group(6.91, 5.26, 1.55, 0.35, blue)
    dashed_group(6.80, 5.66, 1.02, 0.34, green)
    dashed_group(7.91, 5.66, 0.52, 0.34, orange)
    dashed_group(8.54, 5.65, 0.51, 0.35, purple)
    dashed_group(6.96, 6.05, 1.00, 0.35, orange)
    dashed_group(8.18, 6.05, 1.03, 0.35, purple)
    for a, b in projected_edges:
        line(slide, *repair[a], *repair[b], edge, width=0.58)
    for label, pos in repair.items():
        fill, color = projected_style[label]
        labeled_node(*pos, 0.15, label, fill, color, font_size=5.0)
    # Node 6 is the isolated component of the orange partition. Reassigning it
    # to the adjacent purple partition restores connectivity on both sides.
    circle(slide, *repair["6"], 0.27, None, line_color=red, width=1.1)
    circle(slide, *repair["6"], 0.34, None, line_color=red, width=0.7)
    arrow(slide, 8.36, 5.73, 0.28, 0.12, blue, "right")
    textbox(slide, 7.02, 6.47, 2.08, 0.17, "节点 6：P3→P4，使 P3、P4 均连通", 6.0, BLUE, bold=True, align=PP_ALIGN.CENTER)

    # Small-partition merge: S has two candidate neighbors. The stronger
    # boundary coupling w_SB selects B; A remains a separate partition.
    dashed_group(9.30, 5.39, 0.30, 0.35, green)
    dashed_group(9.61, 5.39, 0.28, 0.35, orange)
    dashed_group(9.91, 5.39, 0.31, 0.35, purple)
    line(slide, 9.45, 5.57, 9.75, 5.57, edge, width=0.48)
    line(slide, 9.75, 5.57, 10.06, 5.57, edge, width=1.35)
    labeled_node(9.45, 5.57, 0.145, "A", green_fill, green, font_size=4.7)
    labeled_node(9.75, 5.57, 0.145, "S", orange_fill, orange, font_size=4.7)
    labeled_node(10.06, 5.57, 0.145, "B", purple_fill, purple, font_size=4.7)
    merge_rule = render_formula(tmp / "p4_merge_rule.png", r"w_{SB}>w_{SA}", size=16)
    add_picture_contain(slide, merge_rule, 9.42, 5.77, 0.68, 0.14, border=False)
    arrow(slide, 9.66, 5.94, 0.18, 0.20, red, "down")
    dashed_group(9.29, 6.20, 0.30, 0.35, green)
    dashed_group(9.61, 6.20, 0.61, 0.35, purple)
    line(slide, 9.76, 6.38, 10.06, 6.38, purple, width=0.82)
    labeled_node(9.44, 6.38, 0.145, "A", green_fill, green, font_size=4.7)
    labeled_node(9.76, 6.38, 0.145, "S", purple_fill, purple, font_size=4.7)
    labeled_node(10.06, 6.38, 0.145, "B", purple_fill, purple, font_size=4.7)
    textbox(slide, 9.20, 6.62, 1.18, 0.16, "S 并入耦合更强的 B 所在分区", 5.4, purple, bold=True, align=PP_ALIGN.CENTER)

    textbox(
        slide, 10.42, 5.34, 2.12, 0.98,
        "连通修复\n移动孤立分量中的边界节点，使相关子区均保持连通。\n\n过小子区合并\n在满足规模上限的相邻区中，选择边界耦合权重最大的目标区。",
        7.4, BLACK,
    )
    rect(slide, 10.45, 6.46, 2.05, 0.30, RGBColor(0xEE, 0xF4, 0xFA), radius=True, line=RGBColor(0x98, 0xB6, 0xD2), width=0.7)
    textbox(slide, 10.51, 6.47, 1.93, 0.28, "减少割边  |  保证连通  |  约束规模", 6.7, BLUE, bold=True, align=PP_ALIGN.CENTER)


def page_five(slide, source: Path, tmp: Path) -> None:
    interface = crop(source, tmp / "p5_interface.png", (0, 0, 1672, 500))
    microgrid = crop(source, tmp / "p5_mg.png", (20, 500, 950, 750))
    blue_fill = RGBColor(0xF1, 0xF6, 0xFC)
    green_fill = RGBColor(0xF2, 0xF8, 0xEF)
    orange_fill = RGBColor(0xFE, 0xF5, 0xEC)
    panel_line = RGBColor(0x9B, 0xB9, 0xD8)

    # Main physical interface remains the visual anchor of the page.
    add_picture_contain(slide, interface, 0.48, 2.39, 8.34, 2.52)

    # Boundary-copy construction replaces the previous mode-selection panel.
    rect(slide, 8.96, 2.39, 3.87, 2.52, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 9.15, 2.47, 3.49, 0.28, "边界副本与局部网络生成", 11.4, BLUE, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 9.15, 2.78, 12.64, 2.78, panel_line, width=0.7)

    circle(slide, 9.25, 3.08, 0.25, BLUE, line_color=BLUE, width=0.5)
    textbox(slide, 9.13, 2.96, 0.24, 0.24, "1", 7.2, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.48, 2.88, 1.16, 0.22, "切口识别", 8.8, BLUE, bold=True)
    textbox(slide, 10.56, 2.86, 1.91, 0.38, "定位跨区线路或变压器，形成边界节点集合 B。", 7.1, BLACK)
    line(slide, 9.15, 3.36, 12.64, 3.36, LIGHT_GRAY, width=0.55)

    circle(slide, 9.25, 3.72, 0.25, GREEN, line_color=GREEN, width=0.5)
    textbox(slide, 9.13, 3.60, 0.24, 0.24, "2", 7.2, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.48, 3.51, 1.16, 0.22, "节点复制", 8.8, GREEN, bold=True)
    copy_formula = render_formula(tmp / "p5_boundary_copy.png", r"B\mapsto\left(B^{(i)},B^{(j)}\right)", size=18)
    add_picture_contain(slide, copy_formula, 10.58, 3.49, 1.73, 0.24, border=False)
    textbox(slide, 9.48, 3.78, 3.00, 0.26, "两个副本分别纳入相邻局部子图，保留各自拓扑闭合。", 6.9, BLACK)
    line(slide, 9.15, 4.08, 12.64, 4.08, LIGHT_GRAY, width=0.55)

    circle(slide, 9.25, 4.43, 0.25, ORANGE, line_color=ORANGE, width=0.5)
    textbox(slide, 9.13, 4.31, 0.24, 0.24, "3", 7.2, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.48, 4.20, 1.25, 0.22, "外部量替代", 8.8, ORANGE, bold=True)
    textbox(slide, 10.62, 4.17, 1.89, 0.48, "边界电压作为参考条件，切除网络以 P/Q 等值注入替代。", 7.0, BLACK)
    textbox(slide, 9.47, 4.65, 3.03, 0.17, "全网耦合关系转化为低维接口变量耦合", 6.7, BLUE, bold=True, align=PP_ALIGN.CENTER)

    # Bottom row: PCC equivalent, local equations, and exchanged variables.
    add_picture_contain(slide, microgrid, 0.48, 5.04, 5.86, 1.77)

    rect(slide, 6.49, 5.04, 3.11, 1.77, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 6.65, 5.12, 2.79, 0.24, "局部潮流子问题", 10.0, BLUE, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 6.65, 5.40, 9.44, 5.40, panel_line, width=0.65)
    local_t = render_formula(tmp / "p5_local_t.png", r"F_i(x_i,u_B,S_{\mathrm{ext},i})=0", size=18)
    local_d = render_formula(tmp / "p5_local_d.png", r"F_j(x_j,u_B,S_{\mathrm{MG}})=0", size=18)
    local_m = render_formula(tmp / "p5_local_m.png", r"F_m(x_m,V_{\mathrm{PCC}})=0", size=18)
    textbox(slide, 6.65, 5.50, 0.40, 0.22, "主网", 7.2, BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, local_t, 7.06, 5.50, 2.34, 0.22, border=False)
    textbox(slide, 6.65, 5.85, 0.40, 0.22, "配网", 7.2, GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, local_d, 7.06, 5.85, 2.34, 0.22, border=False)
    textbox(slide, 6.65, 6.20, 0.40, 0.22, "微网", 7.2, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, local_m, 7.06, 6.20, 2.34, 0.22, border=False)
    rect(slide, 6.70, 6.52, 2.69, 0.18, blue_fill, radius=True, line=None)
    textbox(slide, 6.72, 6.51, 2.65, 0.19, "局部拓扑完整，仅通过接口变量耦合", 5.8, BLUE, bold=True, align=PP_ALIGN.CENTER)

    rect(slide, 9.72, 5.04, 3.11, 1.77, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 9.88, 5.12, 2.79, 0.24, "接口变量映射", 10.0, BLUE, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 9.88, 5.40, 12.67, 5.40, panel_line, width=0.65)
    state_formula = render_formula(tmp / "p5_state_vector.png", r"u_B=(V_B,\theta_B)^\top", size=18)
    power_formula = render_formula(tmp / "p5_power_vector.png", r"s_B=(P_B,Q_B)^\top", size=18)
    mg_formula = render_formula(tmp / "p5_mg_equiv.png", r"V_{\mathrm{PCC}}\longleftrightarrow S_{\mathrm{MG}}", size=18)
    textbox(slide, 9.88, 5.49, 0.55, 0.24, "主→配", 7.0, BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, state_formula, 10.48, 5.49, 1.98, 0.23, border=False)
    textbox(slide, 9.88, 5.84, 0.55, 0.24, "配→主", 7.0, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, power_formula, 10.48, 5.84, 1.98, 0.23, border=False)
    textbox(slide, 9.88, 6.19, 0.55, 0.24, "配↔微", 7.0, GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, mg_formula, 10.48, 6.19, 1.98, 0.23, border=False)
    rect(slide, 9.93, 6.52, 2.69, 0.18, green_fill, radius=True, line=None)
    textbox(slide, 9.95, 6.51, 2.65, 0.19, "电压作边界条件，功率作等值反馈", 5.8, GREEN, bold=True, align=PP_ALIGN.CENTER)


def page_six(slide, source: Path, tmp: Path) -> None:
    del source
    blue_fill = RGBColor(0xF1, 0xF6, 0xFC)
    green_fill = RGBColor(0xF2, 0xF8, 0xEF)
    orange_fill = RGBColor(0xFE, 0xF5, 0xEC)
    red = RGBColor(0xC7, 0x2B, 0x2B)
    red_fill = RGBColor(0xFD, 0xF1, 0xF1)
    panel_line = RGBColor(0x9B, 0xB9, 0xD8)

    def stage_box(x, y, w, h, number, title, color, fill):
        rect(slide, x, y, w, h, WHITE, radius=True, line=color, width=0.9)
        rect(slide, x + 0.01, y + 0.01, w - 0.02, 0.31, fill, radius=True, line=None)
        circle(slide, x + 0.25, y + 0.165, 0.22, color, line_color=color, width=0.4)
        textbox(slide, x + 0.14, y + 0.055, 0.22, 0.22, number, 6.8, WHITE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + 0.43, y + 0.035, w - 0.57, 0.26, title, 10.1, color, bold=True, align=PP_ALIGN.CENTER)

    def main_icon(cx, cy, color):
        points = [(cx, cy - 0.22), (cx - 0.25, cy + 0.18), (cx + 0.25, cy + 0.18), (cx, cy + 0.05)]
        for a, b in [(0, 1), (0, 2), (1, 3), (2, 3), (1, 2)]:
            line(slide, *points[a], *points[b], color, width=0.9)
        for px, py in points:
            circle(slide, px, py, 0.09, WHITE, line_color=color, width=0.8)

    def feeder_icon(cx, cy, color):
        line(slide, cx, cy - 0.22, cx, cy + 0.22, color, width=1.0)
        line(slide, cx - 0.28, cy - 0.10, cx + 0.28, cy - 0.10, color, width=1.0)
        for dx in (-0.23, 0.0, 0.23):
            line(slide, cx + dx, cy - 0.10, cx + dx, cy + 0.16, color, width=0.8)
            rect(slide, cx + dx - 0.045, cy + 0.15, 0.09, 0.09, WHITE, line=color, width=0.7)

    def microgrid_icon(cx, cy, color):
        center = (cx, cy)
        points = [(cx - 0.25, cy - 0.18), (cx + 0.25, cy - 0.18), (cx, cy + 0.22)]
        for point in points:
            line(slide, *center, *point, color, width=0.85)
        circle(slide, cx, cy, 0.11, WHITE, line_color=color, width=0.8)
        for idx, (px, py) in enumerate(points):
            circle(slide, px, py, 0.14, WHITE, line_color=color, width=0.8)
            textbox(slide, px - 0.07, py - 0.07, 0.14, 0.14, ("PV", "ES", "L")[idx], 3.5, color, bold=True, align=PP_ALIGN.CENTER)

    stage_box(0.55, 2.45, 3.05, 1.35, "1", "主网分区潮流", BLUE, blue_fill)
    stage_box(4.02, 2.45, 4.14, 1.35, "2", "配网分区潮流", GREEN, green_fill)
    stage_box(8.58, 2.45, 4.20, 1.35, "3", "微网局部潮流", ORANGE, orange_fill)
    arrow(slide, 3.66, 2.99, 0.28, 0.20, BLUE, "right")
    arrow(slide, 8.23, 2.99, 0.28, 0.20, GREEN, "right")

    t_io = render_formula(
        tmp / "p6_t_io.png",
        r"S_{TD}^{(r-1)}\;\longrightarrow\;u_B^{(r)}=(V_B^{(r)},\theta_B^{(r)})",
        size=16,
    )
    d_io = render_formula(
        tmp / "p6_d_io.png",
        r"(u_B^{(r)},S_{MG}^{(r-1)})\;\longrightarrow\;(\widehat S_{TD}^{(r)},V_{\mathrm{PCC}}^{(r)})",
        size=15,
    )
    m_io = render_formula(
        tmp / "p6_m_io.png",
        r"V_{\mathrm{PCC}}^{(r)}\;\longrightarrow\;\widehat S_{MG}^{(r)}",
        size=17,
    )
    textbox(slide, 0.73, 2.82, 2.68, 0.19, "各主网子区并行求解并汇总共享边界状态", 6.7, BLACK, align=PP_ALIGN.CENTER)
    main_icon(1.00, 3.34, BLUE)
    add_picture_contain(slide, t_io, 1.36, 3.13, 2.02, 0.27, border=False)
    textbox(slide, 1.37, 3.45, 1.99, 0.18, "输入交换功率，输出边界电压", 6.0, BLUE, bold=True, align=PP_ALIGN.CENTER)

    textbox(slide, 4.22, 2.82, 3.72, 0.19, "读取边界电压与上一轮微网注入，求解各配网实例", 6.7, BLACK, align=PP_ALIGN.CENTER)
    feeder_icon(4.48, 3.34, GREEN)
    add_picture_contain(slide, d_io, 4.88, 3.12, 3.05, 0.28, border=False)
    textbox(slide, 4.90, 3.45, 3.00, 0.18, "输出主配交换功率与 PCC 电压", 6.0, GREEN, bold=True, align=PP_ALIGN.CENTER)

    textbox(slide, 8.78, 2.82, 3.78, 0.19, "在给定 PCC 电压下独立计算微网等值注入", 6.7, BLACK, align=PP_ALIGN.CENTER)
    microgrid_icon(9.08, 3.34, ORANGE)
    add_picture_contain(slide, m_io, 9.55, 3.13, 2.76, 0.27, border=False)
    textbox(slide, 9.57, 3.45, 2.73, 0.18, "输出等值有功、无功注入", 6.0, ORANGE, bold=True, align=PP_ALIGN.CENTER)

    stage_box(0.55, 4.18, 5.55, 1.55, "5", "功率残差与电压残差判定", BLUE, blue_fill)
    stage_box(6.48, 4.18, 6.30, 1.55, "4", "边界变量松弛更新", TEAL, RGBColor(0xEF, 0xF8, 0xF7))
    arrow(slide, 10.56, 3.86, 0.24, 0.25, ORANGE, "down")
    arrow(slide, 6.14, 4.82, 0.28, 0.20, TEAL, "left")

    update_mg = render_formula(
        tmp / "p6_update_mg.png",
        r"S_{MG}^{(r)}=(1-\alpha)S_{MG}^{(r-1)}+\alpha\widehat S_{MG}^{(r)}",
        size=17,
    )
    update_td = render_formula(
        tmp / "p6_update_td.png",
        r"S_{TD}^{(r)}=(1-\beta)S_{TD}^{(r-1)}+\beta\widehat S_{TD}^{(r)}",
        size=17,
    )
    add_picture_contain(slide, update_mg, 6.78, 4.67, 3.65, 0.24, border=False)
    add_picture_contain(slide, update_td, 6.78, 5.05, 3.65, 0.24, border=False)
    rect(slide, 10.62, 4.62, 1.86, 0.74, RGBColor(0xF3, 0xF8, 0xF8), radius=True, line=panel_line, width=0.6)
    textbox(slide, 10.70, 4.68, 1.70, 0.61, "α：微网等值注入\nβ：主配交换功率\n0 < α, β ≤ 1", 7.1, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 6.86, 5.43, 5.45, 0.18, "松弛更新抑制相邻轮次边界变量振荡", 6.4, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)

    residual_s = render_formula(
        tmp / "p6_residual_s.png",
        r"r_S^{(r)}=\|\widehat S_{TD}^{(r)}-S_{TD}^{(r-1)}\|_2",
        size=16,
    )
    residual_v = render_formula(
        tmp / "p6_residual_v.png",
        r"r_V^{(r)}=\|V_B^{(r)}-V_B^{(r-1)}\|_2",
        size=16,
    )
    criterion = render_formula(
        tmp / "p6_criterion.png",
        r"r_S^{(r)}\leq\varepsilon_S,\qquad r_V^{(r)}\leq\varepsilon_V",
        size=15,
    )
    add_picture_contain(slide, residual_s, 0.82, 4.66, 2.80, 0.23, border=False)
    add_picture_contain(slide, residual_v, 0.82, 5.01, 2.80, 0.23, border=False)
    line(slide, 3.72, 4.60, 3.72, 5.50, panel_line, width=0.6)
    add_picture_contain(slide, criterion, 3.88, 4.70, 1.90, 0.23, border=False)
    rect(slide, 3.94, 5.05, 0.82, 0.27, green_fill, radius=True, line=GREEN, width=0.5)
    textbox(slide, 3.96, 5.05, 0.78, 0.26, "满足：输出", 6.1, GREEN, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 4.83, 5.05, 0.86, 0.27, red_fill, radius=True, line=red, width=0.5)
    textbox(slide, 4.85, 5.05, 0.82, 0.26, "否则：返回 1", 6.0, red, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 0.82, 5.43, 2.80, 0.17, "联合判据同时约束功率映射与电压变化", 6.0, BLUE, bold=True, align=PP_ALIGN.CENTER)

    # Explicit feedback loop for the next outer iteration.
    line(slide, 0.78, 5.73, 0.78, 6.09, red, width=1.0)
    line(slide, 0.78, 6.09, 0.30, 6.09, red, width=1.0)
    line(slide, 0.30, 6.09, 0.30, 3.78, red, width=1.0)
    arrow(slide, 0.20, 3.55, 0.20, 0.24, red, "up")
    textbox(slide, 1.00, 5.91, 7.20, 0.28, "未满足阈值：使用更新后的主配交换功率与微网等值注入进入下一轮协调", 7.6, red, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 8.72, 5.88, 3.78, 0.42, green_fill, radius=True, line=GREEN, width=0.6)
    textbox(slide, 8.84, 5.91, 3.54, 0.34, "收敛后汇总主网、配网和微网状态量", 7.4, GREEN, bold=True, align=PP_ALIGN.CENTER)


def page_seven(slide, source: Path, tmp: Path) -> None:
    del source
    blue_fill = RGBColor(0xF1, 0xF6, 0xFC)
    green_fill = RGBColor(0xF2, 0xF8, 0xEF)
    orange_fill = RGBColor(0xFE, 0xF5, 0xEC)
    purple = RGBColor(0x72, 0x43, 0xA8)
    panel_line = RGBColor(0x9B, 0xB9, 0xD8)

    out_t = render_formula(tmp / "p7_out_t.png", r"u_B", size=19)
    out_d = render_formula(tmp / "p7_out_d.png", r"(\widehat S_{TD},V_{\mathrm{PCC}})", size=15)
    out_m = render_formula(tmp / "p7_out_m.png", r"\widehat S_{MG}", size=19)

    def task_chip(x, y, label, color):
        rect(slide, x, y, 0.42, 0.27, WHITE, radius=True, line=color, width=0.65)
        textbox(slide, x, y, 0.42, 0.27, label, 5.5, color, bold=True, align=PP_ALIGN.CENTER)

    def lane(y, title, prefix, color, fill, output_formula, pool_title):
        rect(slide, 0.48, y, 8.76, 1.12, fill, radius=True, line=color, width=0.85)
        textbox(slide, 0.67, y + 0.06, 1.72, 0.25, title, 9.4, color, bold=True)

        labels = (f"{prefix}1", f"{prefix}2", f"{prefix}3", f"{prefix}N")
        for x, label in zip((0.72, 1.23, 1.74, 2.43), labels):
            task_chip(x, y + 0.43, label, color)
        textbox(slide, 2.15, y + 0.43, 0.22, 0.27, "…", 7.0, color, bold=True, align=PP_ALIGN.CENTER)
        line(slide, 0.78, y + 0.83, 2.79, y + 0.83, color, width=0.65, dashed=True)
        for x in (0.93, 1.44, 1.95, 2.64):
            line(slide, x, y + 0.70, x, y + 0.83, color, width=0.55, dashed=True)
        textbox(slide, 1.18, y + 0.84, 1.26, 0.17, "独立任务队列", 5.8, color, bold=True, align=PP_ALIGN.CENTER)
        arrow(slide, 2.93, y + 0.53, 0.25, 0.18, color, "right")

        rect(slide, 3.25, y + 0.31, 1.88, 0.67, WHITE, radius=True, line=color, width=0.7)
        textbox(slide, 3.37, y + 0.34, 1.64, 0.20, pool_title, 6.4, color, bold=True, align=PP_ALIGN.CENTER)
        for idx, x in enumerate((3.46, 3.92, 4.38), start=1):
            rect(slide, x, y + 0.61, 0.34, 0.23, fill, radius=True, line=color, width=0.5)
            textbox(slide, x, y + 0.61, 0.34, 0.23, f"w{idx}", 4.9, color, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, 4.75, y + 0.62, 0.20, 0.21, "…", 6.0, color, bold=True, align=PP_ALIGN.CENTER)
        arrow(slide, 5.18, y + 0.53, 0.25, 0.18, color, "right")

        rect(slide, 5.50, y + 0.31, 1.70, 0.67, WHITE, radius=True, line=color, width=0.7)
        textbox(slide, 5.61, y + 0.34, 1.48, 0.18, "按任务顺序汇总", 6.3, color, bold=True, align=PP_ALIGN.CENTER)
        add_picture_contain(slide, output_formula, 5.70, y + 0.60, 1.30, 0.23, border=False)
        arrow(slide, 7.25, y + 0.53, 0.23, 0.18, color, "right")

        line(slide, 7.72, y + 0.31, 7.72, y + 0.91, purple, width=1.3)
        circle(slide, 7.72, y + 0.31, 0.10, purple, line_color=purple, width=0.4)
        circle(slide, 7.72, y + 0.91, 0.10, purple, line_color=purple, width=0.4)
        textbox(slide, 7.87, y + 0.39, 1.05, 0.43, "层末同步\n等待全部完成", 6.8, purple, bold=True, align=PP_ALIGN.CENTER)

    lane(2.43, "主网分区任务", "T", BLUE, blue_fill, out_t, "主网工作进程池")
    lane(3.78, "配网潮流任务", "D", GREEN, green_fill, out_d, "配网工作进程池")
    lane(5.13, "微网潮流任务", "M", ORANGE, orange_fill, out_m, "微网工作进程池")

    arrow(slide, 6.25, 3.57, 0.22, 0.17, BLUE, "down")
    textbox(slide, 6.49, 3.56, 0.88, 0.17, "下发边界电压", 5.7, BLUE, bold=True)
    arrow(slide, 6.25, 4.92, 0.22, 0.17, GREEN, "down")
    textbox(slide, 6.49, 4.91, 0.98, 0.17, "下发 PCC 电压", 5.7, GREEN, bold=True)

    rect(slide, 0.48, 6.42, 8.76, 0.38, RGBColor(0xF1, 0xF7, 0xF6), radius=True, line=RGBColor(0x8F, 0xB5, 0xAF), width=0.65)
    textbox(
        slide, 0.67, 6.44, 8.38, 0.31,
        "层内批量提交并行求解  |  层末按任务顺序汇总  |  层间依据边界变量顺序推进",
        7.5, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER,
    )

    # Scheduling rule and performance interpretation.
    rect(slide, 9.45, 2.43, 3.38, 1.45, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 9.64, 2.51, 3.00, 0.25, "动态工作进程配置", 10.2, BLUE, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 9.64, 2.80, 12.64, 2.80, panel_line, width=0.65)
    worker_rule = render_formula(
        tmp / "p7_worker_rule.png",
        r"W_\ell=\min(N_\ell,C),\qquad \ell\in\{T,D,M\}",
        size=18,
    )
    add_picture_contain(slide, worker_rule, 9.75, 2.93, 2.77, 0.27, border=False)
    textbox(slide, 9.70, 3.26, 2.87, 0.43, "Nℓ：本层任务数    C：可用处理核心数\n任务少时不空占进程，任务多时受资源上限约束", 6.8, BLACK, align=PP_ALIGN.CENTER)

    rect(slide, 9.45, 4.01, 3.38, 1.18, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 9.64, 4.09, 3.00, 0.24, "层内耗时与负载均衡", 9.8, TEAL_DARK, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 9.64, 4.37, 12.64, 4.37, panel_line, width=0.65)
    time_rule = render_formula(
        tmp / "p7_time_rule.png",
        r"T_\ell^{\mathrm{par}}\approx\max_q t_{\ell,q}+t_\ell^{\mathrm{gather}}",
        size=17,
    )
    add_picture_contain(slide, time_rule, 9.73, 4.48, 2.80, 0.25, border=False)
    textbox(slide, 9.73, 4.79, 2.80, 0.25, "最慢任务决定层内等待时间，分区均衡直接影响并行效率", 6.5, BLACK, align=PP_ALIGN.CENTER)

    rect(slide, 9.45, 5.32, 3.38, 1.48, RGBColor(0xFA, 0xFB, 0xFC), radius=True, line=panel_line, width=0.8)
    textbox(slide, 9.64, 5.40, 3.00, 0.24, "并行启用、回退与复用", 9.8, purple, bold=True, align=PP_ALIGN.CENTER)
    line(slide, 9.64, 5.68, 12.64, 5.68, panel_line, width=0.65)
    circle(slide, 9.78, 5.93, 0.18, GREEN, line_color=GREEN, width=0.5)
    textbox(slide, 9.69, 5.84, 0.18, 0.18, "✓", 6.5, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.96, 5.79, 2.55, 0.27, "可并行且进程数大于 1：进入工作进程池", 6.5, GREEN, bold=True)
    circle(slide, 9.78, 6.25, 0.18, RGBColor(0xB5, 0x37, 0x37), line_color=RGBColor(0xB5, 0x37, 0x37), width=0.5)
    textbox(slide, 9.69, 6.16, 0.18, 0.18, "×", 6.5, WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.96, 6.11, 2.55, 0.27, "不可序列化或并行关闭：自动串行求解", 6.7, RGBColor(0xA0, 0x32, 0x32), bold=True)
    textbox(slide, 9.70, 6.47, 2.87, 0.20, "配网与微网进程池跨外层迭代复用，减少重复创建开销", 6.3, purple, bold=True, align=PP_ALIGN.CENTER)


LAYOUTS = [page_one, page_two, page_three, page_four, page_five, page_six, page_seven]


def build() -> Path:
    if not BACKUP.exists():
        shutil.copy2(DEMO, BACKUP)
    prs = Presentation(str(BACKUP))

    # Keep seven original progress slides and remove the directory/background lead-in pages.
    delete_slide(prs, 1)
    delete_slide(prs, 0)
    if len(prs.slides) != 7:
        raise RuntimeError(f"Expected 7 retained slides, got {len(prs.slides)}")

    with TemporaryDirectory(prefix="demo-ppt-crops-") as temp:
        tmp = Path(temp)
        for idx, (slide, page, layout_fn) in enumerate(zip(prs.slides, PAGES, LAYOUTS), start=1):
            clear_slide(slide)
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = WHITE
            add_original_header(slide, page["title"], idx)
            source = IMG_DIR / page["image"]
            if not source.exists():
                raise FileNotFoundError(source)
            layout_fn(slide, source, tmp)

        prs.save(str(DEMO))
    return DEMO


if __name__ == "__main__":
    print(build())
