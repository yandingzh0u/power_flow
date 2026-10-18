"""Create a Chinese DOCX/Markdown quality report from audited local evidence."""

from __future__ import annotations

from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/石家庄数据与模型质量核查"
TITLE = "石家庄示范区电网潮流数据与模型质量核查报告"


def read(path):
    return json.loads((ROOT / path).read_text())


def table(headers, rows):
    def cell(value):
        return str(value).replace("|", "／").replace("\n", "；")
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"] + ["| " + " | ".join(cell(v) for v in row) + " |" for row in rows])


def line_of(path, needle):
    for number, line in enumerate((ROOT / path).read_text().splitlines(), 1):
        if needle in line:
            return f"{path}:{number}"
    return path


def exported_records(reports, inventory):
    records = []
    # Bus order is preserved by the BPA converter and is independently keyed by
    # source metadata. Read saved cases here using their serialized bus tables.
    for report in reports:
        label = report["scenario"]
        raw = read(f"data/demo_area/converted/demo_area_{label}_pandapower.json")
        obj = raw["_object"]
        def frame(key):
            data = json.loads(obj[key]["_object"])
            return {index: dict(zip(data["columns"], values)) for index, values in zip(data["index"], data["data"])}
        bus, trafos, lines = frame("bus"), frame("trafo"), frame("line")
        for stage in ("before", "after"):
            for kind, key, model in (("变压器", "overloaded_trafos", trafos), ("线路", "overloaded_lines", lines)):
                for r in report[stage][key]:
                    m = model[r["index"]]
                    u, v = (int(m["hv_bus"]), int(m["lv_bus"])) if kind == "变压器" else (int(m["from_bus"]), int(m["to_bus"]))
                    records.append({"断面": label + "点", "模型": "原始控制模型" if stage == "before" else "假设校正模型", "设备类型": kind, "显示编号": r["name"], "模型索引": r["index"], "起点母线": str(bus[u]["bpa_name"]), "终点母线": str(bus[v]["bpa_name"]), "负载率_%": round(r["loading_percent"], 6), "限值_%": r["limit_percent"], "额定容量_MVA": float(m["sn_mva"]) if kind == "变压器" else "", "额定电流_kA": float(m["max_i_ka"]) if kind == "线路" else "", "起点有功_MW": r.get("p_hv_mw", r.get("p_from_mw")), "起点无功_Mvar": r.get("q_hv_mvar", r.get("q_from_mvar")), "现场已核验": "否"})
    return records


def write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build_markdown(audit, inventory, diagnostics, repairs, records):
    parts = []
    def p(text):
        parts.append(text.strip())
    def heading(text):
        p("## " + text)
    p("# " + TITLE)
    p("编制日期：2026年10月18日\n\n核查对象：当前项目中的匿名示范区13点、20点断面及10 kV馈线曲线。原始DAT的CASEID为zhengding；实际地理边界和设备名称仍以数据提供方说明为准。")
    p("报告性质：本地数据、转换模型与算法实现的质量核查。计算结果用于定位问题；不构成现场设备状态核验或已恢复正常运行的证明。")
    heading("一、主要结论")
    p("目前的异常是数据完整性、控制模型假设、容量约束和算法实现边界共同作用的结果，不能只归因于求解算法，也不能笼统判定原始数据全部错误。")
    p("1. 原始控制模型中TRAFO-008在20点的负载率为489.26%，13点为397.90%。常规牛顿法与Iwamoto牛顿法均能复现；该现象与BE节点为维持目标电压而产生的大量等效无功强相关，尚无证据证明现场设备实际承受同等负载。\n2. 缺少BE实际补偿设备容量、投切状态、无功上下限、变压器分接头可调范围及原BPA结果。已找到的容量、阻抗和分接电压只能证明原始算例怎么填写，不能替代设备台账。\n3. 按明确假设撤去部分BE额外无功控制后，目标变压器负载率下降，但全网仍过载，并出现更明显的电压筛查问题。不能以单台设备低于100%认定模型或电网已正常。\n4. 新核查发现102条有效曲线记录未匹配到模型母线、2个空行被转换成零曲线、目录有2个馈线名称未找到同名文件；还缺少GIS、详细微网与开关控制资料。\n5. 当前示范区分区流程使用全网已解电压作为边界进行复算验证，未形成独立的主配微边界协调求解。已修复部分数值重试和结果校验问题，但该结构性限制仍存在。")
    p("原始目录明确选取了“过载”“重载”和反送场景，因此某些过载可能是研究工况的一部分。整改目标应是得到有依据的物理模型与可信计算结果，而不是人为把每个数值调进正常区间。")
    heading("二、核查依据与数据范围")
    p("证据分为三类：可直接核对的原始文件字段；当前环境可复现的计算结果；尚待台账或原软件结果确认的工程解释。报告分别表述，避免把推断当作实测。")
    p(table(["资料", "位置／数量", "核查用途"], [
        ["原始潮流断面", "data/demo_area/raw/13点sim2合并anony.dat；20点sim2合并anony.dat", "母线、支路、控制类型、1000 MVA基准"],
        ["原始馈线表", "raw/各10kV馈线：18份XLS；目录XLS：1份", "全部21张工作表、线路属性与96点P/Q"],
        ["转换后模型", "converted/demo_area_13/20_pandapower.json", "模型规模、约束、设备编号和静态连通性"],
        ["BPA手册", "raw/潮流手册.pdf，95页", "BE、T、R、X卡及未安排无功语义"],
        ["其他项目材料", "24份DOCX/PPTX、2份WPS、3份进度／成果表、121页任务书及项目ZIP", "查找设备参数与来源；包含Git HEAD中3份已删除文档的只读检查"],
        ["诊断与校正证据", "data/diagnostics/；data/demo_area/corrected/summary.json", "数值对照、假设变更、剩余异常及源文件哈希"],
    ]))
    p("未对报告内所有图片完整OCR；未进行现场量测采集或原BPA软件对照。项目ZIP只有现有论文、专利与报告，没有发现额外设备台账。IEEE和合成网络参数未作为示范区实际参数使用。")
    scale = []
    for k, c in audit["cases"].items():
        n = c["counts"]
        scale.append([k + "点", n["bus"], n["line"], n["trafo"], n["impedance"], n["gen"], c["connected_components"]])
    p(table(["断面", "母线", "线路", "两绕组变压器", "阻抗元件", "BE等效gen", "连通分量"], scale))
    p("每个断面原始DAT均有91张T卡，转换后为73个变压器和18个阻抗元件；不能直接把91与73的差额当作丢失设备。当前两模型各为1个连通分量，未发现重复的“BPA节点名＋基准电压”键，已供给的线路额定电流和变压器容量未出现非正值；这些检查不代表所有物理参数已经正确。")
    heading("三、已确认的计算异常")
    rows = []
    for c in repairs:
        for label, stage in (("原始控制", "before"), ("假设校正", "after")):
            d = c[stage]
            rows.append([c["scenario"] + "点", label, f"{d['trafo_008']['loading_percent']:.2f}%", len(d["overloaded_trafos"]), len(d["overloaded_lines"]), f"{d['vm_pu_min']:.5f}–{d['vm_pu_max']:.5f}"])
    p(table(["断面", "模型", "TRAFO-008", "过载变压器", "过载线路", "电压范围 pu"], rows))
    p("统计口径：本报告引用的诊断和假设校正结果均从原始转换模型起算，不叠加馈线曲线回写或等值微网；全网交流潮流采用NR、flat初值、T型变压器、开启相角、容差1e−8 MVA。过载按模型额定值及100%最大负载率判断，全部73个变压器、2036／2038条线路参与相应断面统计。曲线回写或微网叠加会改变工况，不应混用这些数值。")
    p("3.1 目标变压器及BE无功异常")
    detail = []
    for c in diagnostics["cases"]:
        original = next(t for t in repairs if t["scenario"] == c["scenario"])["before"]
        detail.append([c["scenario"] + "点", f"{original['trafo_008']['p_hv_mw']:.4f}", f"{original['trafo_008']['q_hv_mvar']:.4f}", "−38.0549" if c["scenario"] == "13" else "−47.0887", f"{c['raw_voltage_target_pu']:.3f}"])
    p(table(["断面", "高压侧P MW", "高压侧Q Mvar", "BE求解Q Mvar", "目标电压 pu"], detail))
    p("TRAFO-008是显示编号，对应pandapower变压器索引7；高压母线14（GXVS2）、低压母线16（GXVY2），原始容量10 MVA、37/10.5 kV。20点时约48.07 Mvar的高压侧无功远大于约5.60 MW的有功，是高负载率的重要原因。负载率采用额定电流口径，因此489.26%表示模型电流约为额定值的4.89倍，不能把这个比例直接作为现场实测结论。")
    p("3.2 最严重设备及有功输送瓶颈")
    worst = []
    for c in repairs:
        for stage in ("before", "after"):
            for kind in ("变压器", "线路"):
                subset = [r for r in records if r["断面"] == c["scenario"] + "点" and r["模型"] == ("原始控制模型" if stage == "before" else "假设校正模型") and r["设备类型"] == kind]
                r = max(subset, key=lambda x: x["负载率_%"])
                worst.append([r["断面"], "原始" if stage == "before" else "假设", r["显示编号"], r["起点母线"] + "→" + r["终点母线"], f"{r['负载率_%']:.2f}%"])
    p(table(["断面", "模型", "设备", "两端母线", "最大负载率"], worst))
    bottlenecks = [[c["scenario"] + "点", b["from_name"] + "→" + b["to_name"], f"{b['downstream_net_demand_mw']:.3f}", f"{b['active_transfer_upper_bound_mw']:.3f}", f"{b['support_or_transfer_reduction_lower_bound_mw']:.3f}"] for c in repairs for b in c["active_transfer_bottlenecks"]]
    p(table(["断面", "径向断面", "下游净P MW", "可传P上界 MW", "必要缺口下界 MW"], bottlenecks))
    p("这些容量下界在忽略损耗、允许电压达到模型上限1.10 pu的乐观条件下计算，仍存在正缺口，说明只调无功不能消除所有过载。数值以当前负荷、拓扑和额定电流正确为前提；缺口不等于推荐切负荷量。应先核对线路／并联回路额定值及场景定义，再研究转供、本地支撑或真实设备改造。")
    p("3.3 电压筛查不能以默认限值代替现场标准")
    p(table(["断面", "模型", "超出模型0.9–1.1 pu", "超出筛查0.95–1.05 pu"], [[c["scenario"] + "点", "原始" if stage == "before" else "假设校正", c[stage]["voltage_violations_supplied_limits"], c[stage]["voltage_screening_095_105"]] for c in repairs for stage in ("before", "after")]))
    p("转换器在build_ppc中统一写入VMAX=1.1、VMIN=0.9。这是模型默认值，不是已确认的现场运行标准；0.95–1.05 pu仅作附加筛查，也未被认定为所有电压等级的统一标准。20点假设校正后GXVY2达到1.06510 pu，虽然在默认上限内，仍需工程核查；13点筛查点数由210增至464，说明目标设备负载率改善伴随其他指标变化。")
    heading("四、数据集缺失与一致性问题")
    p("4.1 馈线目录、有效记录及母线映射")
    pr = audit["profiles"]
    p(table(["项目", "实查结果", "影响／判断"], [
        ["目录与文件", "21条目录项，20个唯一名称；DD线跨分类重复；实际18份馈线XLS", "重复分类本身不代表重复设备"],
        ["目录未发现同名文件", "UF现线、YE线", "仅按去空格、去末尾句点比较；可能缺文件或命名未对齐，不擅自视为同一馈线"],
        ["原始总行数", "1663行，其中QL大学线第7、8行为全空记录", "有效记录1661行；转换后保留2个空行"],
        ["成功映射", "两断面均1559条、1559个唯一目标母线；未发现重复目标母线", "有效记录匹配率93.86%"],
        ["有效但未匹配", "VG线100条、CW线1条（CW36）、ZF线1条（ZFA128），共102条", "不能直接将曲线写回现有母线；需对照拓扑版本和匿名映射表"],
        ["JSON的mapped_row_count", "1661仅统计具有生成名称的行", "不是实际匹配成功数，必须与运行时matched_rows区分"],
    ]))
    misses = []
    for k, c in audit["cases"].items():
        missing = c["feeder_matching"]["unmatched_rows"]
        misses.append([k + ":00", f"{sum(r['p_mw_at_named_hour'] for r in missing):.6f}", f"{sum(r['q_mvar_at_named_hour'] for r in missing):.6f}"])
    p(table(["曲线时刻", "未匹配记录P合计 MW", "未匹配记录Q合计 Mvar"], misses))
    p("上述合计是尚未映射的曲线数据，不是全网漏负荷量。当前代码跳过未匹配曲线，并保留未回写母线原DAT负荷，可能形成来源混合的工况；需要按母线核对，不能简单追加上述P/Q。全量逐行位置见附件《馈线未匹配记录.csv》。")
    p("4.2 曲线质量、空白与静默补零")
    p(f"18份曲线均覆盖00:00–23:45的96个15分钟时刻，未发现重复时刻标签、转换后长度错误或非有限数。逐单元格比对318,912个原始数值P/Q单元格与转换JSON，未发现超过1e−12数值容差的不一致。384个空白P/Q单元格全部来自QL大学线两个全空行，不应描述成384个有效测点缺失。")
    p(f"转换结果有{pr['missing_line_type_rows']}行线路型号为空、{pr['nonpositive_length_rows']}行长度非正；这些计数包含上述2个空行。零长度可能是连接／母线段等效，型号缺失也可能对应非线路记录，现有资料不足以逐项判错。曲线有功负值{pr['negative_p_values']:,}个、无功负值{pr['negative_q_values']:,}个，可能来自反送或容性负荷，不能批量改为正数或零。")
    p("转换器float_or_zero把空白或不能解析的数值文本都返回0；运行时曲线长度不足也回填0。这是数据质量风险：缺失、解析失败和实际零值会失去区分。此次有效P/Q原始单元格未发现文本或Excel错误值，但上述风险应在后续输入校验中修复。")
    p("4.3 DAT断面与XLS曲线不一致，且缺少来源日期")
    diffs = []
    for k, c in audit["cases"].items():
        dd = next(d for d in c["dat_vs_curve_load_totals_at_named_hour"] if d["feeder"] == "DD线")
        diffs.append([k + "点", f"{dd['dat_model_load_p_mw']:.6f}", f"{dd['xls_curve_load_p_mw']:.6f}", f"{dd['difference_mw']:+.6f}"])
    p(table(["DD线", "DAT对应母线负荷P MW", "XLS同名时刻P MW", "曲线−DAT MW"], diffs))
    p("存在时刻名相同但负荷不同的情况，可能涉及日期、量测口径、负荷分配、场景加工或网架版本差异；当前证据不能确定哪个来源错误。96点标签仅有时刻，缺少可供关联的明确采样日期、数据质量标识、场景编号和同步量测来源。不应在没有场景清单的情况下把DAT、XLS曲线和微网等值同时叠加后仍称为原始13点／20点结果。")
    p("4.4 缺少的关键设备与工程参数")
    p(table(["优先级", "缺失内容", "需要的最小字段", "对计算的影响"], [
        ["P0", "BE实际补偿与控制台账", "节点／设备ID、类型、额定及分组Mvar、Qmin/Qmax、控制模式、U目标、13/20点投切状态", "决定恒压无功是否实际可用及越限处理"],
        ["P0", "变压器调压边界", "设备ID、调压侧、最小／最大分接、档数、步长、当前档、死区和控制方式", "决定能否用真实调压替代虚拟无功"],
        ["P0", "原BPA结果与场景说明", "PFO／未安排无功表、软件版本、控制选项、断面日期及工况目的", "分离原算例现象与转换差异"],
        ["P0", "严重过载设备额定值核验", "铭牌MVA／允许电流、并联回路、降容系数、温度与运行方式", "验证有功瓶颈是否代表真实约束"],
        ["P1", "曲线—母线完整映射", "原／匿名节点对照、拓扑版本、102条未匹配记录的接入位置", "避免部分回写形成混合断面"],
        ["P1", "电压及电源能力标准", "分电压等级运行限值、发电机／逆变器P-Q能力及场景状态", "避免默认限值被误当工程合格标准"],
        ["P1", "真实微网、储能和DG资料", "PCC、内部网架、设备容量、P/Q控制、SOC与实测曲线", "替换当前研究用等值微网"],
        ["P1", "开关及联络控制资料", "开关ID、两端、常开／常闭、断面状态、联络关系", "验证网架及转供可行性"],
        ["P2", "地理与线路空间资料", "站址／杆塔坐标、坐标系、线路折线、设备GIS对应关系", "支撑真实位置展示与空间拓扑核对"],
    ]))
    p("P0表示影响当前结果可信度、应优先补齐；P1表示影响完整工况和工程核验；P2为后续地理校核资料。优先级仅用于本次整改排序，不是电网安全等级。")
    heading("五、参数语义与建模风险")
    p("5.1 BE节点：未安排无功不等于现场补偿能力")
    p("BPA手册PDF第74–75页说明，BE为无功不受限的PV节点，超出已安排能力的部分可作为未安排无功保持电压；第95页要求检查未安排无功并结合补偿、变比、发电机电压调整。当前转换器将BE等效为PV gen，Q边界统一给到±9999 Mvar，服务入口还显式设置enforce_q_lims=False。仅打开限值开关、仍保留±9999边界，不能得到真实设备能力约束。")
    p("13点49个BE的Qmin均为空，Qmax中40个为0、9个为负；20点48个BE的Qmin均为空、Qmax均为0。并联导纳字段均为空。这里必须区分“空白”“数值0”和“已确认能力为0”。旧诊断JSON中的raw_qmin_mvar=0是解析器的缺省值，不是原卡显式填写；本报告以actual_parameter_inventory.json保存的原始字段为准。负Qmax与缺失Qmin也不能通过交换上下限或擅自解释为固定Q来解决。")
    p("5.2 变比和阻抗：有断面输入值，没有可调范围")
    p(table(["参数", "13点", "20点", "证据性质"], [
        ["GXVS2→GXVY2额定容量", "10 MVA", "10 MVA", "原T卡第34–37列"],
        ["基准电压", "37/10.5 kV", "37/10.5 kV", "原T卡母线基准"],
        ["输入分接电压", "36.75/10.5 kV", "35.87/10.5 kV", "原T卡第63–72列，不能当上下限"],
        ["系统基准", "1000 MVA", "1000 MVA", "DAT第7行"],
        ["该T卡R／X", "0.0436／0.6802 pu", "0.0436／0.6802 pu", "系统基准标幺值，非设备基准百分数"],
        ["R／X控制卡数量", "0／0", "0／0", "未提供调压范围或投切分组数据"],
    ]))
    p("变压器阻抗和系统基准应继续对照设备试验参数，但现有证据不足以认定为单位错位，不能为降低负载率擅自乘10、修改额定容量或调节分接头。两个断面各有21张T卡含负电抗，转换模型相应有21个负vk_percent变压器；三绕组星形等值可能出现负支路电抗，应按原设备整体等值核验，不能见负值就取绝对值。手册PDF第84页也明确三绕组通过多个两绕组等值表示。")
    p("5.3 电气连接、数值等效与实际地理位置")
    p("DAT提供母线—支路电气连接，但当前模型的bus/line geo均没有有效记录。18份XLS提供了馈线线路长度，却没有杆塔坐标和线路折线。转换后的所有net.line.length_km均为1 km，是PPC阻抗换算的等效表达；不得用于统计真实线路总长。当前布局文件来自spring_layout算法，文件名中的actual不构成地理真实性证据。")
    p("两个断面均有26条零R、零X的L卡在转换中赋予X=1e−6 pu，用于数值正则化。这是已发生的参数变更，应单独记录并评估敏感性；可能对应母联或连接段，不能直接认定为损坏数据。当前模型switch表为空，仅包含静态连通关系，不能据此宣称现场没有开关。")
    p("5.4 微网与发电元件的解释边界")
    p("真实微网内部网架、PCC和设备台账未找到。代码的model_type为equivalent_microgrid_overlay，并注明用于接口及并行计算验证，不能将其规模和设备参数作为实测数据。13点与20点sgen数量分别为121和6；负荷／发电符号与转换等值会影响表中元件数量，不能据此断言现场少了115台发电设备。")
    heading("六、算法审查、已修复事项与未完成能力")
    p("6.1 与489.26%直接相关的数值证据")
    p("原始模型用NR及Iwamoto NR从flat初值求解均得到同一目标负载率；按现有±9999 Mvar边界开启无功限值结果也不变。仅把GXVY2的额外Q固定为0时，13点约54.77%、20点约55.16%，但同时释放了原恒压要求。这证明结果对控制模型敏感，不证明Q=0就是正确现场模型，也不是原BPA软件独立验证。")
    p("6.2 分区验证与独立协调求解的区别")
    p("run_demo_area_service先完成全网NR；_compute_demo_partition_boundary再读取全网res_bus中的电压幅值和相角固定各分区边界；局部结果用于与全网解比较，不回写全网res_line或res_trafo。因此极小的分区电压误差只能证明已知边界下局部复算一致，不能证明独立分区法已计算出全网潮流，也不能将该误差作为独立协调算法精度或加速效果的充分证据。")
    p(table(["断面", "当前诊断分区／收敛", "最大电压差 pu", "最大相角差 °", "全网方程残差 MVA"], [[c["scenario"] + "点", f"{c['n_parts']}/{c['converged_partitions']}", f"{c['max_vm_delta_vs_full_pu']:.3e}", f"{c['max_va_delta_vs_full_deg']:.3e}", f"{c['power_mismatch_max_mva']:.3e}"] for c in diagnostics["partition_checks"]]))
    p("当前诊断20点为67分区，旧转换产物保存64分区。划分可能受版本、参数和图划分环境影响，现有证据未锁定差异的唯一原因。引用结果必须带输入、算法配置和环境，不能为对齐旧报告手工合并分区并宣称算法改善。")
    p(table(["事项", "已采取处理", "当前状态"], [
        ["数值重试改变物理约束", "重试保持参考模型enforce_q_lims策略", "已修复并有回归测试"],
        ["退化／不支持的Iwamoto重试", "需要Q限制或无PQ变量时使用NR", "已修复并有回归测试"],
        ["缺失／非有限电压被当零误差", "缺失或非有限结果标为失败，误差不伪造为0", "已修复并有回归测试"],
        ["分区角色不明确", "返回solution_role及independent_partition_solve=False", "已明确；独立边界协调未实现"],
        ["源数据被覆盖风险", "假设模型隔离输出，检查负荷／出力／额定值／拓扑不变及哈希", "已实现；假设未被现场资料确认"],
        ["正式服务仍用旧控制模型", "当前入口读取converted并显式关闭Q限制", "尚未接入按真实设备能力校核的正式模型"],
    ]))
    p("假设校正JSON位于corrected目录，并不会自动替代正式服务或页面的数据。获得真实能力后，需要同时核对服务的模型路径和无功限值策略；只生成新文件不足以让在线结果改变。")
    heading("七、已完成的假设校正及验证边界")
    p("默认假设：仅对原解析字段P、Pmax、Qmax、Qmin均按零值处理的BE，按没有额外未记载无功源的PQ节点求解；实际原卡中部分字段为空，该归零源自解析缺省。零有功不能证明没有SVC／SVG或补偿设备，因此此方案明确属于对照假设。13点处理40个BE、保留9个能力未定控制器；20点处理48个BE。")
    unresolved_names = "、".join(r["name"] for r in repairs[0]["unresolved_controls"])
    p("13点保留待核的BE节点为：" + unresolved_names + "。这些节点原有非零有功或Q字段，没有在缺少依据时整体改为零无功；保留也不意味着其无功能力已验证。")
    p("校正保持全部负荷P/Q、各母线计划有功出力、拓扑、线路／变压器额定值、阻抗、分接输入和外网电压不变；未通过缩放负荷、放大容量或伪造补偿能力消除告警。允许通过带evidence字段的fixed_q或limited_pv配置，在资料齐备后替换假设。")
    p(table(["验证项目", "结果", "说明"], [
        ["全网求解及有限电压", "两个假设模型均收敛", "数值收敛不代表工程可行"],
        ["13点最大方程残差", f"{repairs[0]['after']['power_mismatch_max_mva']:.3e} MVA", "独立重算的方程残差，不等同于输入tolerance_mva或现场误差"],
        ["20点最大方程残差", f"{repairs[1]['after']['power_mismatch_max_mva']:.3e} MVA", "原始输入及结果哈希可追溯"],
        ["JSON重新读取与求解", "此前已验证一致复现", "保存NR、T模型和Q限值策略"],
        ["回归测试", "本次重新运行11项测试全部通过", "覆盖约束重试、缺失结果、假设保护及真实设置入口；不替代现场验证"],
        ["运行状态标记", "field_validated=False；violations_remaining", "不将单台设备改善标记为全网正常"],
    ]))
    heading("八、建议整改顺序与验收条件")
    p(table(["阶段", "具体工作", "验收证据"], [
        ["1. 固化基线", "锁定DAT／XLS／转换器／环境版本，分别保存原断面、曲线回写、微网叠加工况", "文件哈希、工况清单和可重复计算结果"],
        ["2. 补齐控制资料", "优先核查GXVY2及BE极端无功节点，录入补偿能力、状态和调压范围", "设备台账或原软件设置出处，不以猜测值替代"],
        ["3. 修复数据接入", "清理2个空行、逐条处置102个未匹配记录及2个目录缺口；保留缺失标志", "有效记录100%有去向：成功映射或有依据地排除；不要求强行全部匹配"],
        ["4. 重建实际约束", "按真实能力处理PV/PQ转换、补偿投切、调压及电源能力；核验严重瓶颈额定值", "输入物理一致、全网异常清单及各项变化原因"],
        ["5. 对照与接入服务", "与相同设置的BPA结果和可用量测对照，再同步模型路径及约束策略", "母线／支路／外网P/Q对照及工程判限；原始与修订结果均保留"],
        ["6. 独立分区求解", "实现边界电压—功率协调并设置残差停止准则，计算不依赖预先全网解", "独立初值、边界残差、失败场景、与集中式同模型结果对比及总耗时"],
    ]))
    p("验收时应分别回答：数据来源是否完整，模型约束是否代表实际设备，数值方程是否收敛，运行工况是否满足核定约束。若真实过载仍存在，应如实报告并提出有依据的运行或工程方案，不以修改数据消除告警作为验收标准。")
    heading("附录A：假设校正后剩余过载设备")
    p("下表仅列假设校正模型的剩余过载。原始及假设两套模型的全部187条过载记录见同目录《异常设备明细.csv》，包含设备索引、两端母线、负载率、P/Q和额定值。显示编号由索引加1生成，不是现场资产编码。")
    for label in ("13点", "20点"):
        for kind in ("变压器", "线路"):
            subset = [r for r in records if r["断面"] == label and r["模型"] == "假设校正模型" and r["设备类型"] == kind]
            p(f"{label}剩余过载{kind}：{len(subset)}个。")
            p(table(["编号", "起点→终点", "负载率 %", "限值 %"], [[r["显示编号"], r["起点母线"] + "→" + r["终点母线"], f"{r['负载率_%']:.2f}", f"{r['限值_%']:.0f}"] for r in subset]))
    heading("附录B：来源定位、复现和附件")
    refs = [
        ["S1", "原始DAT", "13点：第47行BE、第49行目标T卡；20点：第46行BE、第48行目标T卡；两者第7行MVA_BASE"],
        ["S2", "原始BPA手册", "PDF第74–75页（BE与字段）；84–86页（T/R）；95页（未安排无功）"],
        ["S3", "BE／默认电压约束转换", line_of("tools/convert_bpa_to_pandapower.py", "def build_ppc")],
        ["S4", "零阻抗正则化", line_of("tools/convert_bpa_to_pandapower.py", "if abs(r) <")],
        ["S5", "曲线静默补零", line_of("tools/convert_demo_feeders.py", "def float_or_zero")],
        ["S6", "曲线回写与匹配", line_of("gmss/core.py", "def _apply_demo_feeder_profiles")],
        ["S7", "服务入口", line_of("gmss/core.py", "def run_demo_area_service")],
        ["S8", "分区边界／求解", line_of("gmss/core.py", "def _compute_demo_partition_boundary") + "；" + line_of("gmss/core.py", "def _run_demo_distribution_partitions")],
        ["S9", "微网等值", line_of("gmss/core.py", "def _run_demo_microgrid_equivalents")],
        ["S10", "布局来源", line_of("rebuild_case_study_slides_v2.py", "positions = nx.spring_layout")],
    ]
    p(table(["索引", "证据", "定位"], refs))
    p("随报告提供：异常设备明细.csv（187条）；馈线未匹配记录.csv（两个断面共208条，包括各2个空行，便于完整核对）；核查证据.json（数据统计、参数清单、输入文件完整哈希与结论来源）。")
    p("只读质量检查：在安装pandapower、xlrd的环境运行 python tools/audit_demo_dataset.py。报告生成：在安装python-docx的环境运行 python tools/build_dataset_quality_report.py。回归测试：python -m unittest discover -s tests -v。两个计算脚本为tools/diagnose_demo_transformer.py和tools/repair_demo_model.py；后者输出明确标注假设的独立模型，不覆盖原始输入。")
    p("源文件完整哈希见核查证据.json。相同输入的不同实验结果仅在明确其控制假设、曲线设置和微网设置后比较；本报告没有修改原始模型、核心算法或网页。")
    return "\n\n".join(parts) + "\n"


def clean(text):
    return re.sub(r"`([^`]+)`", r"\1", text).replace("**", "")


def font(run, size=None, bold=None, heading=False):
    face = "Noto Sans CJK SC" if heading else "Noto Serif CJK SC"
    run.font.name = face
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), face)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def docx_from_markdown(markdown, destination, display_title=None,
                       subtitle="原始断面 / 假设校正 / 缺失参数 / 整改依据"):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Noto Serif CJK SC", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Noto Sans CJK SC", Pt(size)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.color.rgb = RGBColor.from_string("203A4F")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("石家庄示范区 · 数据与模型质量核查"), 8, heading=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("第 "), 8)
    for field_name, suffix in (("PAGE", " 页 / 共 "), ("NUMPAGES", " 页")):
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        footer._p.append(field)
        font(footer.add_run(suffix), 8)
    doc.core_properties.title = display_title or TITLE
    doc.core_properties.subject = "数据异常、缺失参数、控制模型和算法边界"
    doc.core_properties.author = "项目资料核查"
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            t = doc.add_paragraph(style="Title")
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(t.add_run(display_title or "石家庄示范区电网潮流\n数据与模型质量核查报告"), 22, True, True)
            doc.add_paragraph(subtitle, style="Subtitle")
        elif line.startswith("## "):
            title = line[3:]
            h = doc.add_heading(title, level=1)
            if title.startswith(("附录A", "附录B")):
                h.paragraph_format.page_break_before = True
        elif re.match(r"\d\.\d ", line):
            doc.add_heading(line, level=2)
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [clean(v.strip()) for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in row):
                    block.append(row)
                i += 1
            t = doc.add_table(rows=0, cols=len(block[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for n, row in enumerate(block):
                cells = t.add_row().cells
                for c, value in zip(cells, row):
                    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    c.text = value
                    for para in c.paragraphs:
                        para.paragraph_format.space_after = Pt(3)
                        para.paragraph_format.space_before = Pt(3)
                        para.paragraph_format.line_spacing = 1.05
                        for run in para.runs:
                            font(run, 8.5, n == 0, n == 0)
                    if n == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EAF0F4")
                        c._tc.get_or_add_tcPr().append(shade)
                # Keep rows intact, and repeat header if a table spans pages.
                pr = t.rows[n]._tr.get_or_add_trPr()
                pr.append(OxmlElement("w:cantSplit"))
                if n == 0:
                    pr.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        else:
            para = doc.add_paragraph()
            font(para.add_run(clean(line)))
        i += 1
    doc.save(destination)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    audit = read("data/diagnostics/dataset_quality_audit.json")
    inventory = read("data/diagnostics/actual_parameter_inventory.json")
    diagnostics = read("data/diagnostics/trafo_008.json")
    repairs = sorted(read("data/demo_area/corrected/summary.json"), key=lambda r: r["scenario"])
    # Refuse to publish figures whose original model or DAT has since changed.
    for c in repairs:
        for name, digest in c["input_sha256"].items():
            if hashlib.sha256((ROOT / name).read_bytes()).hexdigest() != digest:
                raise RuntimeError(f"Input changed since the cited experiment: {name}")
    records = exported_records(repairs, inventory)
    assert len(records) == 187
    write_csv(OUT / "异常设备明细.csv", records)
    unmapped = []
    for label, c in audit["cases"].items():
        for row in c["feeder_matching"]["unmatched_rows"]:
            unmapped.append({"断面": label + "点", "馈线": row["feeder"], "源文件": row["source"], "Excel行号": row["excel_row"], "候选母线名": "/".join(row["candidates"]), "是否空行": "是" if row["empty_bus_name"] else "否", "曲线时刻P_MW": row["p_mw_at_named_hour"], "曲线时刻Q_Mvar": row["q_mvar_at_named_hour"]})
    write_csv(OUT / "馈线未匹配记录.csv", unmapped)
    evidence = {"date": "2026-10-18", "field_validated": False, "audit": audit, "original_parameters": inventory, "experiments": repairs, "diagnosis": diagnostics, "source_sha256": {}}
    inputs = ["data/diagnostics/dataset_quality_audit.json", "data/diagnostics/actual_parameter_inventory.json", "data/diagnostics/trafo_008.json", "data/demo_area/corrected/summary.json", "data/demo_area/converted/demo_area_feeders.json", "gmss/core.py", "tools/convert_bpa_to_pandapower.py", "tools/convert_demo_feeders.py"]
    inputs += [str(p.relative_to(ROOT)) for p in (ROOT / "data/demo_area/raw").rglob("*") if p.is_file()]
    for name in inputs:
        evidence["source_sha256"][name] = hashlib.sha256((ROOT / name).read_bytes()).hexdigest()
    test_log = Path("/tmp/powerflow_quality_report_tests.txt")
    if test_log.exists():
        evidence["test_log"] = test_log.read_text()
        shutil.copyfile(test_log, OUT / "回归测试记录.txt")
    (OUT / "核查证据.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
    markdown = build_markdown(audit, inventory, diagnostics, repairs, records)
    (OUT / (TITLE + ".md")).write_text(markdown)
    docx_from_markdown(markdown, OUT / (TITLE + ".docx"))
    print(OUT / (TITLE + ".docx"))
    print("Overload records:", len(records), "Unmapped records (two cases, incl. blanks):", len(unmapped))


if __name__ == "__main__":
    main()
