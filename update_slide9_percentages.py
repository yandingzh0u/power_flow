from pathlib import Path
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "课题1-推进会-任务2-0302(1)_第7-8页修改版.pptx"
BACKUP = ROOT / "课题1-推进会-任务2-0302(1)_第9页进度修改前备份.pptx"


def set_east_asian_font(run, typeface="微软雅黑"):
    run.font.name = typeface
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        from lxml import etree

        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", typeface)


def set_progress(cell, text):
    text_frame = cell.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(15)
    run.font.bold = True
    set_east_asian_font(run)


def restore_other_slides(source, output):
    replacement = output.with_suffix(".restored.pptx")
    members = set()
    for slide_number in range(1, 22):
        if slide_number == 9:
            continue
        members.add(f"ppt/slides/slide{slide_number}.xml")
        members.add(f"ppt/slides/_rels/slide{slide_number}.xml.rels")

    with ZipFile(source) as source_zip, ZipFile(output) as output_zip, ZipFile(
        replacement, "w", ZIP_DEFLATED
    ) as replacement_zip:
        source_names = set(source_zip.namelist())
        for item in output_zip.infolist():
            data = output_zip.read(item.filename)
            if item.filename in members and item.filename in source_names:
                data = source_zip.read(item.filename)
            replacement_zip.writestr(item, data)
    replacement.replace(output)


def main():
    if not PPTX.exists():
        raise FileNotFoundError(PPTX)
    if not BACKUP.exists():
        shutil.copy2(PPTX, BACKUP)

    presentation = Presentation(PPTX)
    slide = presentation.slides[8]
    table = next(shape.table for shape in slide.shapes if getattr(shape, "has_table", False))
    for row, progress in enumerate(("完成70%", "完成65%", "完成60%", "完成50%"), start=1):
        set_progress(table.cell(row, 4), progress)

    presentation.save(PPTX)
    restore_other_slides(BACKUP, PPTX)
    print(PPTX)


if __name__ == "__main__":
    main()
