from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "demo.pptx"
IMAGE_DIR = ROOT / "img_ppt"
OUTPUT = ROOT / "demo_七页研究进展.pptx"

SLIDES = [
    {
        "title": "研究进展一：主配微异构模型统一解析与拓扑建模方法",
        "image": "ChatGPT Image 2026年8月4日 14_00_34.png",
        "step": "统一建模",
    },
    {
        "title": "研究进展二：主配微边界与微网PCC自动辨识方法",
        "image": "ChatGPT Image 2026年8月4日 14_08_26.png",
        "step": "边界辨识",
    },
    {
        "title": "研究进展三：接口锚点约束下的主网多源区域生长方法",
        "image": "ChatGPT Image 2026年8月4日 14_12_29.png",
        "step": "主网分区",
    },
    {
        "title": "研究进展四：融合电气耦合与连通约束的配网智能分区方法",
        "image": "ChatGPT Image 2026年8月4日 14_17_27.png",
        "step": "配网分区",
    },
    {
        "title": "研究进展五：边界节点复制与接口等值的分区解耦建模方法",
        "image": "ChatGPT Image 2026年8月4日 14_22_51.png",
        "step": "解耦建模",
    },
    {
        "title": "研究进展六：边界电压与功率交替映射及松弛协调迭代方法",
        "image": "ChatGPT Image 2026年8月4日 14_35_34.png",
        "step": "协调迭代",
    },
    {
        "title": "研究进展七：主网子区、配网群与微网群分层并行求解方法",
        "image": "ChatGPT Image 2026年8月4日 14_40_21.png",
        "step": "并行求解",
    },
]

TEAL = RGBColor(0x00, 0x78, 0x70)
TEAL_DARK = RGBColor(0x00, 0x5F, 0x5A)
TEAL_LIGHT = RGBColor(0xE8, 0xF3, 0xF1)
TEXT = RGBColor(0x1D, 0x25, 0x2B)
MUTED = RGBColor(0x7A, 0x83, 0x89)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = prs.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public delete API
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def remove_placeholders(slide) -> None:
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            shape._element.getparent().remove(shape._element)


def add_textbox(slide, x, y, w, h, text, size, color, *, bold=False,
                align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.text_frame.clear()
    box.text_frame.margin_left = Pt(2)
    box.text_frame.margin_right = Pt(2)
    box.text_frame.margin_top = Pt(0)
    box.text_frame.margin_bottom = Pt(0)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def trim_image(source: Path, target: Path, inset: int = 8) -> None:
    with Image.open(source) as image:
        rgb = image.convert("RGB")
        width, height = rgb.size
        crop = (inset, inset, width - inset, height - inset)
        rgb.crop(crop).save(target, format="PNG", optimize=True)


def add_contained_picture(slide, image_path: Path, x, y, w, h):
    with Image.open(image_path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if image_ratio >= box_ratio:
        pic_w = w
        pic_h = int(w / image_ratio)
        pic_x = x
        pic_y = y + int((h - pic_h) / 2)
    else:
        pic_h = h
        pic_w = int(h * image_ratio)
        pic_x = x + int((w - pic_w) / 2)
        pic_y = y
    pic = slide.shapes.add_picture(str(image_path), pic_x, pic_y, pic_w, pic_h)
    pic.line.color.rgb = RGBColor(0xD5, 0xDC, 0xDF)
    pic.line.width = Pt(0.7)
    return pic


def add_header(slide, title: str, number: int) -> None:
    # Compact version of the original demo header, leaving maximum room for method figures.
    add_textbox(
        slide,
        Inches(0.30), Inches(0.03), Inches(3.10), Inches(0.34),
        "二、任务进展情况", 15.5, TEXT, bold=True,
    )
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.42), Inches(13.333), Inches(0.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.28), Inches(0.52), Inches(11.75), Inches(0.58)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    add_textbox(
        slide,
        Inches(0.55), Inches(0.52), Inches(11.10), Inches(0.58),
        title, 20.5, WHITE, bold=True,
    )

    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.15), Inches(0.52), Inches(0.88), Inches(0.58)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = TEAL_DARK
    badge.line.fill.background()
    add_textbox(
        slide,
        Inches(12.15), Inches(0.52), Inches(0.88), Inches(0.58),
        f"{number}/7", 15, WHITE, bold=True, align=PP_ALIGN.CENTER,
    )


def add_progress(slide, active: int) -> None:
    labels = [item["step"] for item in SLIDES]
    total_w = 12.45
    x0 = 0.44
    gap = 0.07
    cell_w = (total_w - gap * 6) / 7
    y = 7.25
    for idx, label in enumerate(labels, start=1):
        x = x0 + (idx - 1) * (cell_w + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y), Inches(cell_w), Inches(0.17),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = TEAL if idx == active else TEAL_LIGHT
        shape.line.fill.background()
        add_textbox(
            slide,
            Inches(x), Inches(y - 0.005), Inches(cell_w), Inches(0.18),
            label, 7.8, WHITE if idx == active else MUTED,
            bold=(idx == active), align=PP_ALIGN.CENTER,
        )


def build() -> Path:
    prs = Presentation(str(TEMPLATE))
    remove_all_slides(prs)
    layout = prs.slide_layouts[6]

    with TemporaryDirectory(prefix="research-progress-ppt-") as tmp:
        tmp_dir = Path(tmp)
        for number, item in enumerate(SLIDES, start=1):
            source = IMAGE_DIR / item["image"]
            if not source.exists():
                raise FileNotFoundError(source)
            trimmed = tmp_dir / f"slide_{number}.png"
            trim_image(source, trimmed, inset=8)

            slide = prs.slides.add_slide(layout)
            remove_placeholders(slide)
            background = slide.background.fill
            background.solid()
            background.fore_color.rgb = WHITE

            add_header(slide, item["title"], number)
            add_contained_picture(
                slide,
                trimmed,
                Inches(0.22), Inches(1.16), Inches(12.89), Inches(6.00),
            )
            add_progress(slide, number)

        prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(build())
